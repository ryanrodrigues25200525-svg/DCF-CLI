"""Review fix: zero-base pre-revenue biotech must survive the workbook guard (#57)."""

from __future__ import annotations

from openpyxl.workbook import Workbook

from app.services.excel_export.mappers.biotech_model import apply_biotech_model


def _payload(base):
    asset = {
        "asset_id": "A1",
        "asset_name": "A1",
        "stage": "Phase 2",
        "development_status": "disclosed",
        "accession_number": "0000000000-26-000001",
        "filing_date": "2026-02-01",
        "report_date": "2025-12-31",
        "form": "10-K",
        "source_statement": "pipeline table",
    }
    return {
        "valuationModel": "biotech_pipeline_rnpv",
        "company": {"ticker": "TST", "name": "Test Bio", "unitsScale": "millions"},
        "market": {"marketCap": 1e9, "debt": 0, "sharesDiluted": 1e8, "currentPrice": 10.0},
        "biotechModel": {
            "baseYear": 2025,
            "forecastYears": 10,
            "commercialRevenueBase": base,
            "pipelineAssets": [asset],
            "assumptions": {
                "baseYear": 2025,
                "commercialRevenueBase": base,
                "commercialRevenueGrowth": [0.1] * 10,
                "commercialFcfMargin": 0.2,
                "otherPipelineRnpv": 0.0,
                "terminalGrowthRate": 0.02,
                "riskFreeRate": 0.04,
                "equityRiskPremium": 0.05,
                "beta": 1.2,
                "costOfDebt": 0.05,
                "normalizedTaxRate": 0.21,
                "marketCapitalization": 1e9,
                "dilutedShares": 1e8,
                "currentPrice": 10.0,
                "cash": 5e8,
                "marketableSecurities": 0.0,
                "debt": 0.0,
                "preferredEquity": 0.0,
                "nonControllingInterest": 0.0,
                "assets": [{
                    "assetId": "A1",
                    "include": 1,
                    "launchYear": 2026,
                    "peakSales": 5e8,
                    "yearsToPeak": 5,
                    "exclusivityYear": 2035,
                    "postLoeErosion": 0.5,
                    "probabilityOfSuccess": 0.3,
                    "retainedShare": 1.0,
                    "contributionMargin": 0.5,
                    "developmentCostPv": 1e7,
                }],
            },
        },
    }


def test_zero_base_pre_revenue_workbook_builds():
    workbook = Workbook()
    apply_biotech_model(workbook, _payload(0))
    assert "Biotech Model" in workbook.sheetnames


def test_negative_base_still_rejected():
    workbook = Workbook()
    try:
        apply_biotech_model(workbook, _payload(-5))
    except ValueError as exc:
        assert "consolidated revenue" in str(exc)
    else:
        raise AssertionError("negative commercial base must raise")
