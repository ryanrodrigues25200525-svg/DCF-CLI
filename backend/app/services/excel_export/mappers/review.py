from __future__ import annotations

import math
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook import Workbook


def _text_items(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [item.strip()[:2000] for item in value if isinstance(item, str) and item.strip()]


def _safe_excel_text(value: str) -> str:
    if value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _review_area(note: str) -> str:
    prefix, separator, _ = note.partition(":")
    if separator and 0 < len(prefix.strip()) <= 40:
        return prefix.strip()
    return "Model"


def _map_data_review_sheet(workbook: Workbook, payload: dict[str, Any]) -> None:
    if "Data Review" in workbook.sheetnames:
        workbook.remove(workbook["Data Review"])
    review = workbook.create_sheet("Data Review")
    review.sheet_view.showGridLines = False

    ui_meta = payload.get("uiMeta")
    ui_meta = ui_meta if isinstance(ui_meta, dict) else {}
    company = payload.get("company")
    company = company if isinstance(company, dict) else {}
    ticker = str(company.get("ticker") or "").strip().upper()
    name = str(company.get("name") or ticker or "Company").strip()

    review["A1"] = _safe_excel_text(f"Draft review — {name} ({ticker})".strip())
    review["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    review["A1"].fill = PatternFill("solid", fgColor="17365D")
    review["A1"].alignment = Alignment(vertical="center")
    review.merge_cells("A1:C1")
    review.row_dimensions[1].height = 26

    review["A2"] = "Editable starting model. Check REVIEW rows, then adjust the input cells and assumptions to fit your view."
    review.merge_cells("A2:C2")
    review["A2"].font = Font(name="Arial", size=10, italic=True, color="404040")
    review["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    review.row_dimensions[2].height = 32

    confidence_label = str(ui_meta.get("confidenceLabel") or "Not scored")
    confidence_score = ui_meta.get("confidenceScore")
    if isinstance(confidence_score, (int, float)) and not isinstance(confidence_score, bool) and math.isfinite(confidence_score):
        confidence_label = f"{confidence_label} ({confidence_score:.0%})"
    review["A3"] = "Model confidence"
    review["B3"] = _safe_excel_text(confidence_label)
    review["A3"].font = Font(name="Arial", size=10, bold=True)
    review["B3"].font = Font(name="Arial", size=10)

    headers = ("Status", "Area", "Review note")
    for column, value in enumerate(headers, start=1):
        cell = review.cell(row=5, column=column, value=value)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="365F91")
        cell.alignment = Alignment(vertical="center")
    review.row_dimensions[5].height = 20

    warnings = _text_items(ui_meta.get("warnings"))
    source_notes = _text_items(ui_meta.get("sourceNotes"))
    items = [("REVIEW", _review_area(note), note) for note in warnings]
    items.extend(("INFO", "Source", note) for note in source_notes)
    if not items:
        items.append(("OK", "Data quality", "No source or model warnings were reported."))

    for row_index, (status, area, note) in enumerate(items, start=6):
        status_cell = review.cell(row=row_index, column=1, value=status)
        status_cell.font = Font(name="Arial", size=10, bold=True)
        status_color = {"REVIEW": "FCE4D6", "INFO": "DDEBF7", "OK": "E2F0D9"}.get(status, "FFFFFF")
        status_cell.fill = PatternFill("solid", fgColor=status_color)
        review.cell(row=row_index, column=2, value=_safe_excel_text(area))
        review.cell(row=row_index, column=3, value=_safe_excel_text(note))
        for column in range(1, 4):
            review.cell(row=row_index, column=column).alignment = Alignment(wrap_text=True, vertical="top")
            review.cell(row=row_index, column=column).font = Font(name="Arial", size=10, bold=column == 1)
        review.row_dimensions[row_index].height = 32

    review.column_dimensions["A"].width = 14
    review.column_dimensions["B"].width = 24
    review.column_dimensions["C"].width = 100
    review.freeze_panes = "A6"
    review.auto_filter.ref = f"A5:C{5 + len(items)}"
