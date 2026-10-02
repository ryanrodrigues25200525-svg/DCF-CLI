from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook import Workbook


_INPUT_BLUE = "0000FF"
_INPUT_FILL = "DDEBF7"
_FORMULA_GREEN = "008000"


def _safe_text(value: Any) -> str | None:
    if not isinstance(value, str):
        return None if value is None else str(value)
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _period(requirement: dict[str, Any]) -> str:
    year = requirement.get("fiscalYear")
    if isinstance(year, int):
        return f"FY{year}"
    return str(requirement.get("asOfDate") or "")


def _identity(requirement: dict[str, Any]) -> str:
    return f"{requirement.get('key')}:{requirement.get('fiscalYear') or requirement.get('asOfDate') or ''}"


def _input_error_condition(value_ref: str, requirement: dict[str, Any]) -> str:
    clauses = [f'AND({value_ref}<>"",NOT(ISNUMBER({value_ref})))']
    allowed_values = requirement.get("allowedValues")
    if isinstance(allowed_values, list) and allowed_values:
        allowed_condition = f"OR({','.join(f'{value_ref}={value}' for value in allowed_values)})"
        clauses.append(f"AND(ISNUMBER({value_ref}),NOT({allowed_condition}))")
    minimum = requirement.get("minimumValue")
    maximum = requirement.get("maximumValue")
    if isinstance(minimum, (int, float)):
        clauses.append(f"AND(ISNUMBER({value_ref}),{value_ref}<{minimum})")
    if isinstance(maximum, (int, float)):
        clauses.append(f"AND(ISNUMBER({value_ref}),{value_ref}>{maximum})")
    return f"OR({','.join(clauses)})"


def _validation_for(cell: str, requirement: dict[str, Any]) -> DataValidation:
    minimum = requirement.get("minimumValue")
    maximum = requirement.get("maximumValue")
    allowed_values = requirement.get("allowedValues")
    if isinstance(allowed_values, list) and allowed_values:
        validation = DataValidation(type="list", formula1=f'"{",".join(str(value) for value in allowed_values)}"')
    elif isinstance(minimum, (int, float)) and isinstance(maximum, (int, float)):
        validation = DataValidation(type="decimal", operator="between", formula1=str(minimum), formula2=str(maximum))
    elif isinstance(minimum, (int, float)):
        validation = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=str(minimum))
    elif isinstance(maximum, (int, float)):
        validation = DataValidation(type="decimal", operator="lessThanOrEqual", formula1=str(maximum))
    else:
        validation = DataValidation(type="decimal", operator="between", formula1="-1E+100", formula2="1E+100")
    validation.allow_blank = True
    validation.showErrorMessage = True
    validation.errorTitle = "Invalid model input"
    validation.error = "Enter a numeric value within the stated model bounds."
    validation.add(cell)
    return validation


def apply_incomplete_input_register(
    workbook: Workbook,
    requirements: list[dict[str, Any]],
    input_cells: dict[str, dict[str, str]],
) -> None:
    """Add editable input references and formulas that withhold value until ready."""
    if "Data Review" in workbook.sheetnames:
        workbook.remove(workbook["Data Review"])
    review = workbook.create_sheet("Data Review")
    review.sheet_view.showGridLines = False
    review["A1"] = "Required input review"
    review["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    review["A1"].fill = PatternFill("solid", fgColor="17365D")
    review.merge_cells("A1:I1")
    review["A2"] = "Fill blue input cells on Input Required. Reported facts require a source reference; missing values are not zero."
    review["A2"].font = Font(name="Arial", size=10, italic=True, color="404040")
    review.merge_cells("A2:I2")

    review_headers = ("Input key", "Required input", "Status", "Source status", "Period", "Unit", "Reason", "Workbook cell", "Source / reference")
    for column, label in enumerate(review_headers, start=1):
        cell = review.cell(row=4, column=column, value=label)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="365F91")
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    input_sheet = workbook["Input Required"]
    for index, requirement in enumerate(requirements):
        identity = _identity(requirement)
        destination = input_cells.get(identity)
        if not destination or destination.get("sheet") != "Input Required":
            raise ValueError(f"Incomplete workbook input {identity} has no Input Required cell destination.")
        value_ref = f"'Input Required'!{destination['cell']}"
        source_ref = f"'Input Required'!G{destination['cell'][1:]}"
        source_required_ref = f"'Input Required'!H{destination['cell'][1:]}"

        review_row = 5 + index
        review_values = (
            requirement.get("key"),
            requirement.get("label"),
            f'=IF({value_ref}="","INPUT REQUIRED",IF(AND({source_required_ref}="Yes",{source_ref}=""),"SOURCE REQUIRED",IF({_input_error_condition(value_ref, requirement)},"INPUT ERROR","READY")))',
            requirement.get("sourceStatus"),
            _period(requirement),
            requirement.get("unit") or "",
            requirement.get("reason"),
            f"Input Required!{destination['cell']}",
            f"={source_ref}",
        )
        for column, value in enumerate(review_values, start=1):
            cell = review.cell(row=review_row, column=column, value=_safe_text(value) if column not in {3, 9} else value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(name="Arial", size=10, color=_FORMULA_GREEN if column in {3, 9} else "000000")

    first_review_row = 5
    last_review_row = 4 + len(requirements)
    review_statuses = f"'Data Review'!$C${first_review_row}:$C${last_review_row}"
    has_errors = f'COUNTIF({review_statuses},"INPUT ERROR")>0'
    has_gaps = f'COUNTIF({review_statuses},"INPUT REQUIRED")+COUNTIF({review_statuses},"SOURCE REQUIRED")>0'
    input_sheet["A3"] = "Model status"
    input_sheet["A3"].font = Font(name="Arial", size=10, bold=True)
    input_sheet["B3"] = f'=IF({has_errors},"INPUT ERROR — check required inputs",IF({has_gaps},"INCOMPLETE — fill required inputs","READY"))'
    input_sheet["B3"].font = Font(name="Arial", size=10, bold=True, color=_FORMULA_GREEN)
    input_sheet["B3"].alignment = Alignment(wrap_text=True)

    input_sheet.sheet_view.showGridLines = False
    review.freeze_panes = "A5"
    review.auto_filter.ref = f"A4:I{max(4, 4 + len(requirements))}"
    for column, width in {"A": 34, "B": 30, "C": 20, "D": 18, "E": 14, "F": 16, "G": 64, "H": 28, "I": 36}.items():
        review.column_dimensions[column].width = width


def apply_incomplete_workbook(
    workbook: Workbook,
    payload: dict[str, Any],
    *,
    preserve_existing: bool = False,
) -> dict[str, dict[str, str]]:
    """Create a no-valuation input shell until the family schedule is mapped."""
    if "Input Required" in workbook.sheetnames:
        workbook.remove(workbook["Input Required"])
    sheet = workbook.create_sheet("Input Required")
    if not preserve_existing:
        for existing in list(workbook.worksheets):
            if existing is not sheet:
                workbook.remove(existing)

    company = payload.get("company") if isinstance(payload.get("company"), dict) else {}
    name = str(company.get("name") or "Company")
    ticker = str(company.get("ticker") or "").upper()
    model = str(payload.get("valuationModel") or "unselected model")
    requirements = payload.get("requiredInputs")
    requirements = requirements if isinstance(requirements, list) else []

    sheet.sheet_view.showGridLines = False
    sheet["A1"] = f"INCOMPLETE MODEL — {name} ({ticker})"
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="17365D")
    sheet.merge_cells("A1:H1")
    sheet["A2"] = f"Model family: {model}. No valuation has been calculated because required inputs are missing."
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="404040")
    sheet.merge_cells("A2:H2")

    headers = ("Input key", "Required input", "Period", "Unit", "Why required", "Enter value", "Source / reference", "Source required")
    for column, label in enumerate(headers, start=1):
        cell = sheet.cell(row=5, column=column, value=label)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="365F91")
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    input_cells: dict[str, dict[str, str]] = {}
    for row, requirement in enumerate(requirements, start=6):
        if not isinstance(requirement, dict):
            raise ValueError(f"Incomplete workbook requirement at index {row - 6} must be an object.")
        identity = _identity(requirement)
        input_cells[identity] = {"sheet": "Input Required", "cell": f"F{row}"}
        values = (
            requirement.get("key"),
            requirement.get("label"),
            _period(requirement),
            requirement.get("unit") or "",
            requirement.get("reason"),
            None,
            None,
            "Yes" if requirement.get("sourceReferenceRequired") is True else "No",
        )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=_safe_text(value))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(name="Arial", size=10, color=_INPUT_BLUE if column in {6, 7} else "000000")
            if column in {6, 7}:
                cell.fill = PatternFill("solid", fgColor=_INPUT_FILL)
                cell.protection = Protection(locked=False)
        sheet.add_data_validation(_validation_for(f"F{row}", requirement))

    for column, width in {"A": 34, "B": 30, "C": 14, "D": 16, "E": 64, "F": 18, "G": 36, "H": 16}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A5:H{max(5, 5 + len(requirements))}"
    apply_incomplete_input_register(workbook, requirements, input_cells)
    return input_cells
