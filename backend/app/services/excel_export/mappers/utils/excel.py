from __future__ import annotations

from copy import copy
from typing import Any

from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .constants import (
    DCF_HELPER_CASH_CELL,
    DCF_HELPER_CASH_CELL_ABS,
    DCF_HELPER_DEBT_CELL,
    DCF_HELPER_DEBT_CELL_ABS,
    DCF_HELPER_NON_OP_CELL,
    DCF_HELPER_NON_OP_CELL_ABS,
    DCF_TIMELINE_COLUMNS,
    RECALC_COLUMNS,
    SHEET_DATA_RECALCULATED,
    SHEET_DCF_BASE,
    SHEET_DCF_BEAR,
    SHEET_DCF_BULL,
)
from .core import _first_float, _scale, _to_float
from .financial import (
    _forecast_map,
    _infer_revenue_growth_rate,
    _sanitize_terminal_growth_rate,
    _sanitize_wacc_rate,
    _scenario_assumptions,
    _scenario_first_projection_forecast,
)


def _set_comment(worksheet: Worksheet, cell_ref: str, text: str, *, author: str = "DCF Builder") -> None:
    cell = worksheet[cell_ref]
    if isinstance(cell, MergedCell):
        return
    if not text.strip():
        return
    cell.comment = Comment(text, author)

def _safe_set(worksheet: Worksheet, cell_ref: str, value: Any) -> None:
    _safe_set_with_options(worksheet, cell_ref, value, clear_if_none=False)

def _safe_set_or_clear(worksheet: Worksheet, cell_ref: str, value: Any) -> None:
    _safe_set_with_options(worksheet, cell_ref, value, clear_if_none=True)

def _safe_set_with_options(
    worksheet: Worksheet,
    cell_ref: str,
    value: Any,
    *,
    clear_if_none: bool,
) -> None:
    cell = worksheet[cell_ref]
    if isinstance(cell, MergedCell):
        return
    if cell.data_type == "f":
        return
    if value is None:
        if not clear_if_none:
            return
        cell.value = None
        cell.hyperlink = None
        return
    cell.value = value

def _force_set(worksheet: Worksheet, cell_ref: str, value: Any) -> None:
    cell = worksheet[cell_ref]
    if isinstance(cell, MergedCell):
        return
    cell.value = value

def _axis_from_bounds(
    values: list[float],
    *,
    center: float,
    step: float,
    min_value: float,
    max_value: float,
) -> list[float]:
    if len(values) >= 2:
        low = max(min_value, min(values))
        high = min(max_value, max(values))
        if high <= low:
            high = min(max_value, low + step * 4)
        if high > low:
            return [round(low + (high - low) * idx / 4, 4) for idx in range(5)]

    seed = center if center == center else (min_value + max_value) / 2
    axis = [seed + (idx - 2) * step for idx in range(5)]
    clamped = [max(min_value, min(max_value, value)) for value in axis]
    for idx in range(1, len(clamped)):
        if clamped[idx] <= clamped[idx - 1]:
            clamped[idx] = min(max_value, clamped[idx - 1] + max(step / 2, 0.0005))
    return [round(value, 4) for value in clamped]

def _uniform_axis_step(values: list[float], *, tolerance: float = 1e-6) -> float | None:
    if len(values) != 5:
        return None
    deltas = [values[idx + 1] - values[idx] for idx in range(len(values) - 1)]
    if any(delta <= 0 for delta in deltas):
        return None
    first = deltas[0]
    if all(abs(delta - first) <= tolerance for delta in deltas[1:]):
        return first
    return None

def _set_percent_axis_row(
    sheet: Worksheet,
    *,
    cells: tuple[str, str, str, str, str],
    values: list[float],
) -> None:
    step = _uniform_axis_step(values)
    if step is not None:
        center = cells[2]
        _force_set(sheet, center, values[2])
        _force_set(sheet, cells[1], f"={center}-{step:.4f}")
        _force_set(sheet, cells[0], f"={cells[1]}-{step:.4f}")
        _force_set(sheet, cells[3], f"={center}+{step:.4f}")
        _force_set(sheet, cells[4], f"={cells[3]}+{step:.4f}")
    else:
        for cell_ref, value in zip(cells, values, strict=False):
            _force_set(sheet, cell_ref, value)

    for cell_ref in cells:
        sheet[cell_ref].number_format = "0.0%"

def _set_percent_axis_column(
    sheet: Worksheet,
    *,
    cells: tuple[str, str, str, str, str],
    values: list[float],
) -> None:
    step = _uniform_axis_step(values)
    if step is not None:
        center = cells[2]
        _force_set(sheet, center, values[2])
        _force_set(sheet, cells[1], f"={center}-{step:.4f}")
        _force_set(sheet, cells[0], f"={cells[1]}-{step:.4f}")
        _force_set(sheet, cells[3], f"={center}+{step:.4f}")
        _force_set(sheet, cells[4], f"={cells[3]}+{step:.4f}")
    else:
        for cell_ref, value in zip(cells, values, strict=False):
            _force_set(sheet, cell_ref, value)

    for cell_ref in cells:
        sheet[cell_ref].number_format = "0.0%"

def _clear_sensitivity_blocks(*scenario_sheets: Worksheet) -> None:
    for sheet in scenario_sheets:
        for row in range(117, 131):
            for col in "BCDEFGHIJKLMNO":
                _force_set(sheet, f"{col}{row}", None)

def _map_scenario_forecasts_to_sheet(
    sheet: Worksheet,
    forecast_by_year: dict[int, dict[str, Any]],
    timeline_years: list[int | None],
    divisor: float,
) -> None:
    projection_columns = DCF_TIMELINE_COLUMNS[2:]  # J..Q
    projection_years = timeline_years[2:] if len(timeline_years) >= 3 else []
    first_projection_year = next((year for year in projection_years if year in forecast_by_year), None)

    for idx, col in enumerate(projection_columns):
        if idx >= len(projection_years):
            break
        projection_year = projection_years[idx]
        forecast = forecast_by_year.get(projection_year)
        if not isinstance(forecast, dict):
            continue

        if projection_year == first_projection_year:
            dcf_index = DCF_TIMELINE_COLUMNS.index(col)
            previous_col = DCF_TIMELINE_COLUMNS[dcf_index - 1] if dcf_index > 0 else None
            if previous_col is not None:
                _force_set(sheet, f"{col}20", f"={previous_col}20*(1+$F$14)")

def _apply_scenario_snapshot_to_sheet(
    sheet: Worksheet,
    snapshot: dict[str, Any],
    timeline_years: list[int | None],
    divisor: float,
) -> None:
    assumptions = _scenario_assumptions(snapshot)
    forecast_by_year = _forecast_map(snapshot)

    tax_rate = _to_float(assumptions.get("taxRate"))
    wacc_rate = _sanitize_wacc_rate(assumptions.get("waccRate"))
    da_pct_revenue = _to_float(assumptions.get("daPctRevenue"))
    ebit_margin = _to_float(assumptions.get("ebitMargin"))
    steady_state_ebit_margin = _to_float(assumptions.get("ebitMarginSteadyState"))
    gross_margin = _to_float(assumptions.get("grossMargin"))
    revenue_growth_rate = _to_float(assumptions.get("revenueGrowthStage1"))
    if revenue_growth_rate is None:
        revenue_growth_rate = _to_float(assumptions.get("revenueGrowthRate"))
    if revenue_growth_rate is None:
        revenue_growth_rate = _to_float(assumptions.get("revenueGrowth"))
    if revenue_growth_rate is None:
        revenue_growth_rate = _infer_revenue_growth_rate(snapshot.get("forecasts") if isinstance(snapshot.get("forecasts"), list) else [])
    terminal_growth = _sanitize_terminal_growth_rate(
        assumptions.get("terminalGrowthRate"),
        reference_wacc=wacc_rate,
    )
    exit_multiple = _to_float(assumptions.get("terminalExitMultiple"))

    if tax_rate is not None:
        _force_set(sheet, "F11", tax_rate)
    if da_pct_revenue is not None:
        _force_set(sheet, "F13", da_pct_revenue)
    if ebit_margin is not None:
        _force_set(sheet, "F15", ebit_margin)
        sheet["F15"].number_format = "0.0%"
    if steady_state_ebit_margin is not None:
        _force_set(sheet, "F17", steady_state_ebit_margin)
        sheet["F17"].number_format = "0.0%"
    if gross_margin is not None:
        _force_set(sheet, "F16", gross_margin)
        sheet["F16"].number_format = "0.0%"
    if revenue_growth_rate is not None:
        _force_set(sheet, "F14", revenue_growth_rate)
        sheet["F14"].number_format = "0.0%"
    stage2_growth = _to_float(assumptions.get("revenueGrowthStage2"))
    if stage2_growth is not None:
        _force_set(sheet, "F19", stage2_growth)
        sheet["F19"].number_format = "0.0%"
    for cell_ref, key, label, default in (
        ("I13", "ebitMarginConvergenceYears", "EBIT-margin convergence period", 5),
        ("I14", "revenueGrowthStage1Years", "Revenue-growth Stage 1 duration", 3),
        ("I15", "revenueGrowthFadeYears", "Revenue-growth fade duration", 4),
    ):
        value = assumptions.get(key)
        if value is not None:
            parsed = _to_float(value)
            if parsed is None:
                parsed = float(default)
            if parsed < 1 or not parsed.is_integer():
                raise ValueError(f"{label} must be a positive whole number.")
            _force_set(sheet, cell_ref, int(parsed))
    if terminal_growth is not None:
        _force_set(sheet, "Q103", terminal_growth)
    if exit_multiple is not None:
        _force_set(sheet, "C16", exit_multiple)

    first_projection = _scenario_first_projection_forecast(forecast_by_year, timeline_years)
    capex_ratio = _to_float(assumptions.get("capexPctRevenue"))
    if capex_ratio is None:
        capex_ratio = _to_float(assumptions.get("capexRatio"))
    if capex_ratio is None and isinstance(first_projection, dict):
        capex = _to_float(first_projection.get("capex"))
        revenue = _to_float(first_projection.get("revenue"))
        capex_ratio = abs(capex / revenue) if capex is not None and revenue else None
    nwc_ratio = _to_float(assumptions.get("nwcPctRevenue"))
    if nwc_ratio is None:
        nwc_ratio = _to_float(assumptions.get("nwcChangeRatio"))
    if capex_ratio is not None:
        _force_set(sheet, "F9", capex_ratio)
        sheet["F9"].number_format = "0.0%"
    _force_set(sheet, "F10", nwc_ratio if nwc_ratio is not None else 0.0)
    sheet["F10"].number_format = "0.0%"

    _map_scenario_forecasts_to_sheet(sheet, forecast_by_year, timeline_years, divisor)

def _link_dcf_income_statement_to_recalculated_data(
    timeline_years: list[int | None],
    historical_years: set[int],
    *scenario_sheets: Worksheet,
) -> None:
    for sheet in scenario_sheets:
        for idx, dcf_col in enumerate(DCF_TIMELINE_COLUMNS):
            if idx >= len(timeline_years):
                continue
            year = timeline_years[idx]
            if year is None:
                for row in range(18, 116):
                    _force_set(sheet, f"{dcf_col}{row}", None)
                continue
            recalc_col = RECALC_COLUMNS[idx]
            if year in historical_years:
                _force_set(sheet, f"{dcf_col}20", f"='{SHEET_DATA_RECALCULATED}'!{recalc_col}12")
                _force_set(sheet, f"{dcf_col}24", f"='{SHEET_DATA_RECALCULATED}'!{recalc_col}16")
                _force_set(sheet, f"{dcf_col}27", f"='{SHEET_DATA_RECALCULATED}'!{recalc_col}17")
                _force_set(sheet, f"{dcf_col}36", f"='{SHEET_DATA_RECALCULATED}'!{recalc_col}24")
                _force_set(sheet, f"{dcf_col}39", f"='{SHEET_DATA_RECALCULATED}'!{recalc_col}25")
                _force_set(sheet, f"{dcf_col}42", f"='{SHEET_DATA_RECALCULATED}'!{recalc_col}26")
                _force_set(sheet, f"{dcf_col}45", f"='{SHEET_DATA_RECALCULATED}'!{recalc_col}27")
                continue
            if idx == 0:
                continue
            prior = DCF_TIMELINE_COLUMNS[idx - 1]
            formulas = {
                20: f"={prior}20*(1+{dcf_col}21)",
                24: f"={dcf_col}20*(1-$F$16)-{dcf_col}27",
                25: f"=IFERROR({dcf_col}24/{dcf_col}20,0)",
                27: f"={dcf_col}20*{prior}28",
                28: f"=IFERROR({dcf_col}27/{dcf_col}20,0)",
                30: f"={dcf_col}24+{dcf_col}27",
                32: f"={dcf_col}20-{dcf_col}30",
                33: f"=IFERROR({dcf_col}32/{dcf_col}20,0)",
                36: f"={dcf_col}20*{prior}37",
                37: f"=IFERROR({dcf_col}36/{dcf_col}20,0)",
                39: f"={dcf_col}20*{prior}40",
                40: f"=IFERROR({dcf_col}39/{dcf_col}20,0)",
                42: f"={dcf_col}20*$F$13",
                43: f"=IFERROR({dcf_col}42/{dcf_col}20,0)",
                45: f"={dcf_col}32-{dcf_col}36-{dcf_col}39-{dcf_col}20*{dcf_col}53",
                46: f"=IFERROR({dcf_col}45/{dcf_col}20,0)",
                48: f"=-({dcf_col}36+{dcf_col}39+{dcf_col}45)",
                49: f"=IFERROR(-{dcf_col}48/{dcf_col}20,0)",
                51: f"={dcf_col}32+{dcf_col}48",
                52: f"=IFERROR({dcf_col}51/{dcf_col}20,0)",
                54: f"={dcf_col}51+{dcf_col}68",
                55: f"=IFERROR({dcf_col}54/{dcf_col}20,0)",
                57: f"=-{dcf_col}51*$F$11",
                58: "=$F$11",
                60: f"={dcf_col}51+{dcf_col}57",
                61: f"=IFERROR({dcf_col}60/{dcf_col}20,0)",
            }
            for row, formula in formulas.items():
                _force_set(sheet, f"{dcf_col}{row}", formula)

def _normalize_public_dcf_assumption_block(
    dcf_base: Worksheet,
    dcf_bull: Worksheet,
    dcf_bear: Worksheet,
    payload: dict[str, Any],
    divisor: float,
) -> None:
    market = payload.get("market", {})
    cash = _to_float(market.get("cash")) or 0.0
    debt = _to_float(market.get("debt")) or 0.0
    non_operating_assets = _to_float(market.get("nonOperatingAssets")) or 0.0

    for sheet, wacc_cell in (
        (dcf_base, "D34"),
        (dcf_bull, "D35"),
        (dcf_bear, "D36"),
    ):
        _safe_set(sheet, "B8", "Valuation Summary")
        _safe_set(sheet, "B9", "Implied Enterprise Value")
        _safe_set(sheet, "B10", "Implied EV / EBITDA")
        _safe_set(sheet, "B11", "Current Equity Value")
        _safe_set(sheet, "B12", "Implied Equity Value")
        _safe_set(sheet, "B13", "Implied Share Price (Gordon Growth)")
        _safe_set(sheet, "B14", "Implied Share Price (Exit Multiple)")
        sheet["B13"]._style = copy(sheet["B12"]._style)
        sheet["B14"]._style = copy(sheet["B12"]._style)
        sheet["C13"]._style = copy(sheet["C12"]._style)
        sheet["C14"]._style = copy(sheet["C12"]._style)
        sheet["C13"].number_format = "$0.00"
        sheet["C14"].number_format = "$0.00"
        _safe_set(sheet, "E14", "Stage 1 Revenue Growth")
        _safe_set(sheet, "B15", "Terminal Assumptions")
        _safe_set(sheet, "B16", "Exit EBITDA Multiple")
        _safe_set(sheet, "B17", "Cash and Cash Equivalents")
        _safe_set(sheet, "B18", "Income Statement")
        _safe_set_or_clear(sheet, "B19", None)
        sheet["F14"]._style = copy(sheet["F11"]._style)
        sheet["F13"]._style = copy(sheet["F11"]._style)
        sheet["F15"]._style = copy(sheet["F11"]._style)
        sheet["F16"]._style = copy(sheet["F11"]._style)
        sheet["F17"]._style = copy(sheet["F15"]._style)
        sheet["F19"]._style = copy(sheet["F14"]._style)
        _safe_set(sheet, "E15", "Starting EBIT Margin")
        _safe_set(sheet, "E17", "Steady-State EBIT Margin")
        _safe_set(sheet, "E19", "Stage 2 Revenue Growth")
        for row, label in (
            (13, "EBIT Margin Convergence (Years)"),
            (14, "Revenue Growth Stage 1 (Years)"),
            (15, "Revenue Growth Fade (Years)"),
            (16, "Diluted Shares (millions)"),
            (17, "Current Share Price"),
        ):
            _safe_set(sheet, f"H{row}", label)
            sheet[f"H{row}"]._style = copy(sheet["H9"]._style)
            sheet[f"I{row}"]._style = copy(sheet["F14"]._style)
        sheet["F14"].number_format = "0.0%"
        sheet["F17"].number_format = "0.0%"
        sheet["F19"].number_format = "0.0%"
        sheet["I13"].number_format = "0"
        sheet["I14"].number_format = "0"
        sheet["I15"].number_format = "0"
        sheet["I16"].number_format = "#,##0.0;(#,##0.0);-"
        sheet["I17"].number_format = "$0.00"
        _safe_set(sheet, "E16", "Forecast Gross Margin")
        _force_set(sheet, "C17", _scale(cash, divisor))
        _force_set(sheet, "C10", "=IFERROR(C9/L54,0)")
        _force_set(sheet, "C12", f"=C9-{DCF_HELPER_DEBT_CELL}+{DCF_HELPER_CASH_CELL}+{DCF_HELPER_NON_OP_CELL}-F20-F21")
        _force_set(sheet, "C13", "=Q111/$I$16")
        _force_set(sheet, "C14", "=C12/$I$16")
        _force_set(sheet, "F12", f"=WACC!{wacc_cell}")
        _safe_set_or_clear(sheet, "F18", None)
        _force_set(sheet, DCF_HELPER_CASH_CELL, "=C17")
        _force_set(sheet, DCF_HELPER_DEBT_CELL, _scale(debt, divisor))
        _force_set(sheet, DCF_HELPER_NON_OP_CELL, _scale(non_operating_assets, divisor))
        for assumption_ref in ("F9", "F10", "F11", "F12", "F13", "C16"):
            sheet[assumption_ref].comment = None

def _enforce_core_public_dcf_formulas(*scenario_sheets: Worksheet) -> None:
    for sheet in scenario_sheets:
        first_forecast_idx = next(
            (
                idx
                for idx, col in enumerate(DCF_TIMELINE_COLUMNS)
                if isinstance(sheet[f"{col}18"].value, str)
                and str(sheet[f"{col}18"].value).endswith("E")
            ),
            None,
        )
        if first_forecast_idx is None:
            continue
        formula_projection_columns = DCF_TIMELINE_COLUMNS[first_forecast_idx:]

        for idx, col in enumerate(DCF_TIMELINE_COLUMNS):
            if not sheet[f"{col}18"].value:
                for row in range(18, 116):
                    _force_set(sheet, f"{col}{row}", None)
                continue
            previous_col = DCF_TIMELINE_COLUMNS[idx - 1] if idx > 0 else None
            _force_set(
                sheet,
                f"{col}21",
                '=""' if previous_col is None else f"=IFERROR({col}20/{previous_col}20-1,0)",
            )
            _force_set(sheet, f"{col}30", f"={col}24+{col}27")
            _force_set(sheet, f"{col}32", f"={col}20-{col}30")
            _force_set(sheet, f"{col}33", f"=IFERROR({col}32/{col}20,0)")
            _force_set(sheet, f"{col}52", f"=IFERROR({col}51/{col}20,0)")
            _force_set(sheet, f"{col}54", f"={col}51+{col}68")
            _force_set(sheet, f"{col}55", f"=IFERROR({col}54/{col}20,0)")
            _force_set(sheet, f"{col}57", f"=-{col}51*{col}58")
            _force_set(sheet, f"{col}58", "=$F$11")
            _force_set(sheet, f"{col}60", f"={col}51+{col}57")
            _force_set(sheet, f"{col}61", f"=IFERROR({col}60/{col}20,0)")
            _force_set(sheet, f"{col}70", f"=IFERROR(-{col}68/{col}65,0)")
            _force_set(sheet, f"{col}74", f"={col}60")
            _force_set(sheet, f"{col}75", f"={col}68")
            _force_set(sheet, f"{col}76", f"={col}65")
            _force_set(sheet, f"{col}78", f"=SUM({col}74:{col}77)")
            _force_set(sheet, f"{col}79", f"=IFERROR({col}78/{col}20,0)")
            if col not in formula_projection_columns:
                for row in (82, 83, 84, 85, 87):
                    _force_set(sheet, f"{col}{row}", None)
                _force_set(sheet, f"{col}25", f"=IFERROR({col}24/{col}20,0)")
                _force_set(sheet, f"{col}28", f"=IFERROR({col}27/{col}20,0)")
                _force_set(sheet, f"{col}37", f"=IFERROR({col}36/{col}20,0)")
                _force_set(sheet, f"{col}40", f"=IFERROR({col}39/{col}20,0)")
                _force_set(sheet, f"{col}43", f"=IFERROR({col}42/{col}20,0)")
                _force_set(sheet, f"{col}46", f"=IFERROR({col}45/{col}20,0)")

        for idx, col in enumerate(formula_projection_columns):
            prev_idx = first_forecast_idx + idx - 1
            if prev_idx >= 0:
                prev_col = DCF_TIMELINE_COLUMNS[prev_idx]
                _force_set(sheet, f"{col}37", f"={prev_col}37")
                _force_set(sheet, f"{col}40", f"={prev_col}40")
                _force_set(sheet, f"{col}43", f"={prev_col}43")
                _force_set(sheet, f"{col}27", f"={col}20*{prev_col}28")
            _force_set(sheet, f"{col}24", f"={col}20*(1-$F$16)-{col}27")
            _force_set(sheet, f"{col}25", f"=IFERROR({col}24/{col}20,0)")
            _force_set(sheet, f"{col}28", f"=IFERROR({col}27/{col}20,0)")
            _force_set(sheet, f"{col}36", f"={col}20*{col}37")
            _force_set(sheet, f"{col}39", f"={col}20*{col}40")
            _force_set(sheet, f"{col}42", f"={col}20*$F$13")
            _force_set(sheet, f"{col}45", f"={col}32-{col}36-{col}39-{col}20*{col}53")
            _force_set(sheet, f"{col}46", f"=IFERROR({col}45/{col}20,0)")
            _force_set(sheet, f"{col}85", "=$F$12")
            _force_set(sheet, f"{col}65", "=-$F$9")
            _force_set(sheet, f"{col}77", "=-$F$10")
            _force_set(
                sheet,
                f"{col}21",
                f"=$F$14+($F$19-$F$14)*MAX(0,MIN(1,({col}$83-$I$14)/$I$15))",
            )
            _force_set(
                sheet,
                f"{col}53",
                f"=IF({col}$83<$I$13,$F$15+($F$17-$F$15)*{col}$83/$I$13,$F$17)",
            )
            sheet[f"{col}53"]._style = copy(sheet[f"{col}52"]._style)
            sheet[f"{col}53"].number_format = "0.0%"

        _force_set(sheet, "E81", "=I9")
        for idx, col in enumerate(formula_projection_columns):
            prev_col = formula_projection_columns[idx - 1] if idx > 0 else None
            _force_set(sheet, f"{col}82", "=EOMONTH(I11,12)" if prev_col is None else f"=EOMONTH({prev_col}82,12)")
            _force_set(sheet, f"{col}83", "=1" if prev_col is None else f"={prev_col}83+1")
            _force_set(sheet, f"{col}84", f"={col}83/2" if idx == 0 else f"={col}83-0.5")
            _force_set(sheet, f"{col}87", f"={col}78/(1+{col}85)^{col}84")

        _force_set(sheet, "Q92", "=Q54*$C$16")
        _force_set(sheet, "Q93", "=Q92/(1+Q85)^Q83")
        pv_first_col = formula_projection_columns[0]
        _force_set(sheet, "Q94", f"=Q93+SUM({pv_first_col}87:Q87)")
        _force_set(sheet, "C9", "=Q94")
        _force_set(sheet, "C13", "=Q111/$I$16")
        _force_set(sheet, "C14", "=C12/$I$16")
        _force_set(sheet, "Q95", f"=-{DCF_HELPER_DEBT_CELL_ABS}")
        _force_set(sheet, "Q96", "=-$F$20-$F$21")
        _safe_set(sheet, "B96", "(-) Minority Interest and Preferred Equity")
        _force_set(sheet, "Q97", f"={DCF_HELPER_CASH_CELL_ABS}+{DCF_HELPER_NON_OP_CELL_ABS}")
        _safe_set(sheet, "B97", "(+) Cash and Non-Operating Assets")
        _force_set(sheet, "Q98", "=SUM(Q94:Q97)")
        _force_set(sheet, "Q99", "=C11")
        _force_set(sheet, "Q100", "=IFERROR(Q98/Q99 - 1,0)")
        _force_set(sheet, "Q104", "=Q78*(1+Q103)")
        _safe_set(sheet, "B104", "Final Year UFCF × (1 + g)")
        _force_set(sheet, "Q105", "=IFERROR(IF(Q85>Q103,Q104/(Q85-Q103),0),0)")
        _force_set(sheet, "Q106", "=Q105/(1+Q85)^Q83")
        _force_set(sheet, "Q107", f"=Q106+SUM({pv_first_col}87:Q87)")
        _force_set(sheet, "Q108", f"=-{DCF_HELPER_DEBT_CELL_ABS}")
        _force_set(sheet, "Q109", "=-$F$20-$F$21")
        _safe_set(sheet, "B109", "(-) Minority Interest and Preferred Equity")
        _force_set(sheet, "Q110", f"={DCF_HELPER_CASH_CELL_ABS}+{DCF_HELPER_NON_OP_CELL_ABS}")
        _safe_set(sheet, "B110", "(+) Cash and Non-Operating Assets")
        _force_set(sheet, "Q111", "=SUM(Q107:Q110)")
        _force_set(sheet, "Q112", "=C11")
        _force_set(sheet, "Q113", "=IFERROR(Q111/Q112 - 1,0)")
        _safe_set(sheet, "B53", "Forecast EBIT Margin Driver")

        anchor_col = formula_projection_columns[0]
        sheet[f"{anchor_col}20"].comment = None
        if first_forecast_idx > 0:
            previous_col = DCF_TIMELINE_COLUMNS[first_forecast_idx - 1]
            _force_set(
                sheet,
                f"{anchor_col}20",
                f"={previous_col}20*(1+{anchor_col}21)",
            )
        for idx in range(1, len(formula_projection_columns)):
            col = formula_projection_columns[idx]
            prev_col = formula_projection_columns[idx - 1]
            _force_set(sheet, f"{col}20", f"={prev_col}20*(1+{col}21)")

def _scenario_choose_formula(base_cell: str, bull_cell: str | None = None, bear_cell: str | None = None) -> str:
    bull_ref = bull_cell or base_cell
    bear_ref = bear_cell or base_cell
    return (
        "=CHOOSE(Cover!$C$12,"
        f"'{SHEET_DCF_BASE}'!{base_cell},"
        f"'{SHEET_DCF_BULL}'!{bull_ref},"
        f"'{SHEET_DCF_BEAR}'!{bear_ref})"
    )

def _enforce_outputs_bridge_formulas(outputs: Worksheet) -> None:
    _force_set(outputs, "E18", _scenario_choose_formula("I9"))
    first_forecast_idx = next(
        (
            idx
            for idx, col in enumerate(DCF_TIMELINE_COLUMNS)
            if isinstance(outputs[f"{col}6"].value, str)
            and str(outputs[f"{col}6"].value).endswith("E")
        ),
        None,
    )
    if first_forecast_idx is None:
        return
    forecast_columns = DCF_TIMELINE_COLUMNS[first_forecast_idx:]
    _force_set(outputs, "D27", _scenario_choose_formula("$F$12"))
    _force_set(outputs, "H27", "=D27")
    _force_set(outputs, "D28", _scenario_choose_formula("$C$16"))
    _force_set(outputs, "H28", _scenario_choose_formula("$Q$103"))

    for col in ("H", "I", "J", "K", "L", "M", "N", "O", "P", "Q"):
        if not outputs[f"{col}6"].value:
            for row in (8, 9, 10, 12, 13, 14, 15, 16, 19, 20, 21, 22, 23):
                _force_set(outputs, f"{col}{row}", None)
            continue
        _force_set(outputs, f"{col}8", _scenario_choose_formula(f"{col}51"))
        _force_set(outputs, f"{col}9", _scenario_choose_formula(f"{col}68"))
        _force_set(outputs, f"{col}10", _scenario_choose_formula(f"{col}54"))
        _force_set(outputs, f"{col}12", _scenario_choose_formula(f"{col}60"))
        _force_set(outputs, f"{col}13", f"={col}9")
        _force_set(outputs, f"{col}14", _scenario_choose_formula(f"{col}65"))
        _force_set(outputs, f"{col}15", _scenario_choose_formula(f"{col}77"))
        _force_set(outputs, f"{col}16", f"=SUM({col}12:{col}15)")
        _force_set(outputs, f"{col}22", _scenario_choose_formula("$F$12"))
        if col in forecast_columns:
            _force_set(outputs, f"{col}19", _scenario_choose_formula(f"{col}82"))
            _force_set(outputs, f"{col}20", _scenario_choose_formula(f"{col}83"))
            _force_set(outputs, f"{col}21", _scenario_choose_formula(f"{col}84"))
            _force_set(outputs, f"{col}23", f"={col}16/(1+{col}22)^{col}21")
        else:
            for row in (19, 20, 21, 23):
                _force_set(outputs, f"{col}{row}", None)

    _safe_set(outputs, "B37", "(-) Minority Interest and Preferred Equity")
    _safe_set(outputs, "F37", "(-) Minority Interest and Preferred Equity")
    _safe_set(outputs, "B38", "(+) Cash and Non-Operating Assets")
    _safe_set(outputs, "F38", "(+) Cash and Non-Operating Assets")
    first_forecast_col = forecast_columns[0]
    _force_set(outputs, "D30", f"=SUM({first_forecast_col}23:Q23)")
    _force_set(outputs, "H30", "=D30")
    _force_set(outputs, "F31", "Final Year UFCF × (1 + g)")
    _force_set(outputs, "D33", "=D32/(1+D27)^Q20")
    _force_set(outputs, "H33", "=H32/(1+H27)^Q20")
    _force_set(outputs, "D36", f"=-{_scenario_choose_formula(DCF_HELPER_DEBT_CELL_ABS)[1:]}")
    _force_set(outputs, "H36", "=D36")
    _force_set(
        outputs,
        "D37",
        f"=-{_scenario_choose_formula('$F$20')[1:]}-{_scenario_choose_formula('$F$21')[1:]}",
    )
    _force_set(outputs, "H37", "=D37")
    _force_set(
        outputs,
        "D38",
        f"={_scenario_choose_formula(DCF_HELPER_CASH_CELL_ABS)[1:]}+{_scenario_choose_formula(DCF_HELPER_NON_OP_CELL_ABS)[1:]}",
    )
    _force_set(outputs, "H38", "=D38")
    _force_set(outputs, "D41", _scenario_choose_formula("$C$11"))
    _force_set(outputs, "H41", "=D41")
    _safe_set(outputs, "B43", "Implied Share Price")
    _safe_set(outputs, "F43", "Implied Share Price")
    for target, source in (("B43", "B42"), ("D43", "D42"), ("F43", "F42"), ("H43", "H42")):
        outputs[target]._style = copy(outputs[source]._style)
    _force_set(outputs, "D43", f"=D40/'{SHEET_DCF_BASE}'!$I$16")
    _force_set(outputs, "H43", f"=H40/'{SHEET_DCF_BASE}'!$I$16")
    outputs["D43"].number_format = "$0.00"
    outputs["H43"].number_format = "$0.00"

def _rewrite_formula_sheet_name_references(workbook: Workbook, *, old_name: str, new_name: str) -> None:
    old_ref = f"'{old_name}'!"
    new_ref = f"'{new_name}'!"
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("=") and old_ref in value:
                    cell.value = value.replace(old_ref, new_ref)
                hyperlink = cell.hyperlink
                if hyperlink is not None and isinstance(hyperlink.location, str) and old_ref in hyperlink.location:
                    hyperlink.location = hyperlink.location.replace(old_ref, new_ref)
