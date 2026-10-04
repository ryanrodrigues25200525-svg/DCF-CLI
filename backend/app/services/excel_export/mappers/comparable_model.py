from __future__ import annotations

import math
import time
from typing import Any

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

from app.services.valuation.model_eligibility import COMPARABLE_VALUATION_METHODS

from .review import _map_data_review_sheet
from .utils import _resolve_amount_scale_divisor, _scale, _to_float


def _latest_metric(income: dict[str, Any], years: list[Any], keys: list[str]) -> float | None:
    if not years:
        return None
    index = len(years) - 1
    for key in keys:
        values = income.get(key)
        if not isinstance(values, list) or index >= len(values):
            continue
        return _to_float(values[index])
    return None


def apply_comparable_model(workbook: Workbook, payload: dict[str, Any]) -> None:
    valuation_model = str(payload.get("valuationModel") or "")
    if valuation_model not in COMPARABLE_VALUATION_METHODS:
        raise ValueError("Comparable valuation workbook requires EV/EBITDA or EV/Revenue.")
    comparable_source = payload.get("comparableModel") if isinstance(payload.get("comparableModel"), dict) else {}
    if comparable_source.get("method") != valuation_model:
        raise ValueError("Comparable valuation export method does not match its production route.")
    peer_source = str(comparable_source.get("peerSource") or "").strip().lower()
    peer_status = str(comparable_source.get("peerStatus") or "").strip().lower()
    peer_fetched_at = _to_float(comparable_source.get("peerFetchedAtMs"))
    if (
        peer_status not in {"live", "cached"}
        or comparable_source.get("peerFallbackUsed") is not False
        or not peer_source
        or any(token in peer_source for token in ("default", "stale", "unavailable"))
        or peer_fetched_at is None
        or not math.isfinite(peer_fetched_at)
        or not 0 <= time.time() * 1000 - peer_fetched_at <= 24 * 60 * 60 * 1000
    ):
        raise ValueError("Comparable valuation requires a current curated non-fallback peer source.")

    company = payload.get("company") if isinstance(payload.get("company"), dict) else {}
    market = payload.get("market") if isinstance(payload.get("market"), dict) else {}
    historicals = payload.get("historicals") if isinstance(payload.get("historicals"), dict) else {}
    income = historicals.get("income") if isinstance(historicals.get("income"), dict) else {}
    years = historicals.get("years") if isinstance(historicals.get("years"), list) else []
    peers = payload.get("comps") if isinstance(payload.get("comps"), list) else []
    divisor = _resolve_amount_scale_divisor(payload)
    metric_name = "EBITDA" if valuation_model == "ev_ebitda" else "Revenue"
    metric_keys = ["EBITDA"] if valuation_model == "ev_ebitda" else ["Total Revenue", "Revenue"]
    target_metric = _latest_metric(income, years, metric_keys)
    if target_metric is None or target_metric <= 0:
        raise ValueError("Comparable valuation needs positive source-backed target " + metric_name + ".")

    valid_peers = []
    seen: set[str] = set()
    for peer in peers:
        if not isinstance(peer, dict):
            continue
        ticker = str(peer.get("ticker") or "").strip().upper()
        enterprise_value = _to_float(peer.get("ev"))
        revenue = _to_float(peer.get("revenue"))
        ebitda = _to_float(peer.get("ebitda"))
        denominator = ebitda if valuation_model == "ev_ebitda" else revenue
        if (
            not ticker
            or ticker in seen
            or enterprise_value is None or enterprise_value <= 0
            or denominator is None or denominator <= 0
        ):
            continue
        seen.add(ticker)
        valid_peers.append({
            "ticker": ticker,
            "name": str(peer.get("company") or ticker),
            "enterprise_value": enterprise_value,
            "revenue": revenue,
            "ebitda": ebitda,
        })
    if len(valid_peers) < 3:
        raise ValueError("Comparable valuation workbook requires at least three source-ready peers.")
    metric_key = "ebitda" if valuation_model == "ev_ebitda" else "revenue"
    peer_multiples = [peer["enterprise_value"] / peer[metric_key] for peer in valid_peers]
    sorted_multiples = sorted(peer_multiples)
    middle = len(sorted_multiples) // 2
    peer_median = sorted_multiples[middle] if len(sorted_multiples) % 2 else (sorted_multiples[middle - 1] + sorted_multiples[middle]) / 2
    selected_multiple = _to_float(comparable_source.get("selectedMultiple"))
    if selected_multiple is None or not math.isfinite(selected_multiple) or selected_multiple <= 0:
        raise ValueError("Comparable valuation selected multiple is missing or invalid.")
    if abs(peer_median - selected_multiple) > max(0.25, peer_median * 0.05):
        raise ValueError("Comparable valuation median does not reconcile to the source-ready peer set.")

    if not years:
        raise ValueError("Comparable valuation requires filed historical fiscal years.")
    current_price = _to_float(market.get("currentPrice"))
    shares = _to_float(market.get("sharesDiluted"))
    cash = _to_float(market.get("cash"))
    debt = _to_float(market.get("debt"))
    securities = _to_float(market.get("nonOperatingAssets"))
    minority_interest = _to_float(market.get("minorityInterest"))
    preferred_equity = _to_float(market.get("preferredEquity"))
    bridge_values = (current_price, shares, cash, debt, securities, minority_interest, preferred_equity)
    if any(value is None for value in bridge_values):
        raise ValueError("Comparable valuation requires a complete source-backed common-equity bridge.")
    if current_price <= 0 or shares <= 0:
        raise ValueError("Comparable valuation requires positive current price and diluted shares.")

    title = "EV/EBITDA Trading Comparables" if valuation_model == "ev_ebitda" else "EV/Revenue Trading Comparables"
    ticker = str(company.get("ticker") or "").strip().upper()
    model = workbook.create_sheet("Comparable Valuation")
    for sheet in list(workbook.worksheets):
        if sheet is not model:
            workbook.remove(sheet)
    workbook.active = workbook.index(model)
    model.sheet_view.showGridLines = False

    navy = "17365D"
    medium_blue = "365F91"
    input_blue = "0000FF"
    pale_yellow = "FFF2CC"
    light_blue = "DDEBF7"
    white = "FFFFFF"
    thin_gray = Side(style="thin", color="B7C9D6")
    amount_format = "#,##0.0;(#,##0.0);-"
    multiple_format = "0.0x"
    price_format = "$0.00;($0.00);-"

    model["A1"] = title + " — " + ticker
    model["A1"].font = Font(name="Arial", size=16, bold=True, color=white)
    model["A1"].fill = PatternFill("solid", fgColor=navy)
    model.merge_cells("A1:G1")
    model.row_dimensions[1].height = 28
    model["A2"] = str(company.get("name") or company.get("ticker") or "Company") + " (" + str(company.get("ticker") or "") + ")"
    model["A2"].font = Font(name="Arial", size=11, bold=True, color=navy)
    model["A3"] = "Amounts are shown in " + str(company.get("unitsScale") or "millions") + "; peer multiples are calculated from current enterprise value and filed LTM metrics."
    model.merge_cells("A3:G3")
    model["A3"].alignment = Alignment(wrap_text=True)

    model["A4"] = "Target FY" + str(years[-1]) + " " + metric_name
    model["B4"] = _scale(target_metric, divisor)
    model["B4"].number_format = amount_format
    model["B4"].comment = Comment("Source: latest filed canonical " + metric_name + "; see Data Review for filing concept and provenance.", "Codex")
    model["A5"] = "Peer median " + ("EV/EBITDA" if valuation_model == "ev_ebitda" else "EV/Revenue")
    peer_multiple_column = "G" if valuation_model == "ev_ebitda" else "F"
    first_peer_row = 24
    last_peer_row = first_peer_row + len(valid_peers) - 1
    model["B5"] = f"=MEDIAN({peer_multiple_column}{first_peer_row}:{peer_multiple_column}{last_peer_row})"
    model["B5"].number_format = multiple_format
    model["A6"] = "Optional analyst multiple override"
    model["B6"] = None
    model["B6"].font = Font(name="Arial", size=10, color=input_blue)
    model["B6"].fill = PatternFill("solid", fgColor=pale_yellow)
    model["B6"].number_format = multiple_format
    model["B6"].comment = Comment("Optional editable assumption. Leave blank to use the live peer median formula in B5.", "Codex")
    model["A7"] = "Selected multiple"
    model["B7"] = "=IF(B6>0,B6,B5)"
    model["B7"].number_format = multiple_format
    model["A8"] = "Enterprise Value"
    model["B8"] = "=B4*B7"
    model["B8"].number_format = amount_format

    bridge_rows = [
        (10, "Cash and cash equivalents", cash),
        (11, "Marketable securities", securities),
        (12, "Debt", debt),
        (13, "Noncontrolling interest", minority_interest),
        (14, "Preferred equity", preferred_equity),
    ]
    for row, label, value in bridge_rows:
        model.cell(row=row, column=1, value=label)
        model.cell(row=row, column=2, value=_scale(value, divisor))
        model.cell(row=row, column=2).number_format = amount_format
    model["A15"] = "Common Equity Value"
    model["B15"] = "=B8+B10+B11-B12-B13-B14"
    model["B15"].number_format = amount_format
    model["A16"] = "Diluted Shares"
    model["B16"] = shares / divisor
    model["B16"].number_format = amount_format
    model["A17"] = "Implied Share Price"
    model["B17"] = "=IFERROR(B15/B16,0)"
    model["B17"].number_format = price_format
    model["A18"] = "Current Share Price"
    model["B18"] = current_price
    model["B18"].number_format = price_format
    model["A19"] = "Implied Upside / (Downside)"
    model["B19"] = "=IFERROR(B17/B18-1,0)"
    model["B19"].number_format = "0.0%;(0.0%);-"

    model["A22"] = "Current Peer Set"
    model["A22"].font = Font(name="Arial", size=11, bold=True, color=navy)
    headers = ("Ticker", "Company", "Enterprise Value", "LTM Revenue", "LTM EBITDA", "EV/Revenue", "EV/EBITDA")
    for column, header in enumerate(headers, start=1):
        cell = model.cell(row=23, column=column, value=header)
        cell.font = Font(name="Arial", size=10, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=medium_blue)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for idx, peer in enumerate(valid_peers, start=24):
        model.cell(row=idx, column=1, value=peer["ticker"])
        model.cell(row=idx, column=2, value=peer["name"])
        model.cell(row=idx, column=3, value=_scale(peer["enterprise_value"], divisor))
        model.cell(row=idx, column=4, value=_scale(peer["revenue"], divisor) if peer["revenue"] is not None else None)
        model.cell(row=idx, column=5, value=_scale(peer["ebitda"], divisor) if peer["ebitda"] is not None else None)
        model.cell(row=idx, column=6, value=f"=IFERROR(C{idx}/D{idx},0)")
        model.cell(row=idx, column=7, value=f"=IFERROR(C{idx}/E{idx},0)")
        for column in (3, 4, 5):
            model.cell(row=idx, column=column).number_format = amount_format
        for column in (6, 7):
            model.cell(row=idx, column=column).number_format = multiple_format

    model["I4"] = "Multiple Sensitivity"
    model["I4"].font = Font(name="Arial", size=11, bold=True, color=navy)
    for column, header in enumerate(("Change", "Multiple", "Enterprise Value", "Common Equity", "Per Share"), start=9):
        cell = model.cell(row=5, column=column, value=header)
        cell.font = Font(name="Arial", size=9, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=medium_blue)
        cell.alignment = Alignment(wrap_text=True)
    for row, change in enumerate((-0.2, 0, 0.2), start=6):
        model.cell(row=row, column=9, value=change)
        model.cell(row=row, column=9).font = Font(name="Arial", size=10, color=input_blue)
        model.cell(row=row, column=9).fill = PatternFill("solid", fgColor=pale_yellow)
        model.cell(row=row, column=9).number_format = "0%;(0%);-"
        model.cell(row=row, column=10, value=f"=$B$7*(1+I{row})")
        model.cell(row=row, column=11, value=f"=$B$4*J{row}")
        model.cell(row=row, column=12, value=f"=K{row}+$B$10+$B$11-$B$12-$B$13-$B$14")
        model.cell(row=row, column=13, value=f"=IFERROR(L{row}/$B$16,0)")
        model.cell(row=row, column=10).number_format = multiple_format
        for column in (11, 12):
            model.cell(row=row, column=column).number_format = amount_format
        model.cell(row=row, column=13).number_format = price_format

    for row in (7, 8, 15, 17, 19):
        model.cell(row=row, column=1).font = Font(name="Arial", size=10, bold=True)
        model.cell(row=row, column=2).font = Font(name="Arial", size=10, bold=True)
        model.cell(row=row, column=1).fill = PatternFill("solid", fgColor=light_blue)
        model.cell(row=row, column=2).fill = PatternFill("solid", fgColor=light_blue)
    for row in range(4, 20):
        for column in (1, 2):
            model.cell(row=row, column=column).alignment = Alignment(vertical="center")
            model.cell(row=row, column=column).border = Border(bottom=thin_gray)
    for column, width in {"A": 32, "B": 28, "C": 18, "D": 16, "E": 16, "F": 14, "G": 14, "I": 12, "J": 12, "K": 20, "L": 20, "M": 16}.items():
        model.column_dimensions[column].width = width
    model.freeze_panes = "A24"
    model.auto_filter.ref = f"A23:G{23 + len(valid_peers)}"
    _map_data_review_sheet(workbook, payload)


def apply_incomplete_comparable_model(workbook: Workbook, payload: dict[str, Any]) -> None:
    """Build a formula-based peer schedule with blank, sourced denominator inputs."""
    from .incomplete import apply_incomplete_workbook

    valuation_model = str(payload.get("valuationModel") or "")
    if valuation_model not in COMPARABLE_VALUATION_METHODS:
        raise ValueError("Incomplete comparable workbook requires EV/EBITDA or EV/Revenue.")
    comparable = payload.get("comparableModel") if isinstance(payload.get("comparableModel"), dict) else {}
    if comparable.get("method") != valuation_model:
        raise ValueError("Incomplete comparable model requires a matching method.")
    peer_fallback_mode = comparable.get("peerFallbackUsed") is True
    if not peer_fallback_mode and comparable.get("peerFallbackUsed") is not False:
        raise ValueError("Incomplete comparable workbook requires a current non-fallback peer set.")
    peer_source = str(comparable.get("peerSource") or "").strip()
    peer_fetched_at = _to_float(comparable.get("peerFetchedAtMs"))
    if (not peer_source or any(token in peer_source.lower() for token in ("default", "stale", "unavailable"))
        or peer_fetched_at is None or not math.isfinite(peer_fetched_at)
        or not 0 <= time.time() * 1000 - peer_fetched_at <= 24 * 60 * 60 * 1000):
        raise ValueError("Incomplete comparable workbook requires a current peer-source record.")

    company = payload.get("company") if isinstance(payload.get("company"), dict) else {}
    market = payload.get("market") if isinstance(payload.get("market"), dict) else {}
    historicals = payload.get("historicals") if isinstance(payload.get("historicals"), dict) else {}
    years = historicals.get("years") if isinstance(historicals.get("years"), list) else []
    if not years:
        raise ValueError("Incomplete comparable workbook requires filed fiscal periods.")
    metric = "ebitda" if valuation_model == "ev_ebitda" else "revenue"
    metric_name = "EBITDA" if metric == "ebitda" else "Revenue"
    income = historicals.get("income") if isinstance(historicals.get("income"), dict) else {}
    target_metric = _latest_metric(income, years, ["EBITDA"] if metric == "ebitda" else ["Total Revenue", "Revenue"])
    if target_metric is None or target_metric <= 0:
        raise ValueError(f"Incomplete comparable workbook requires positive filed target {metric_name}.")

    current_price = _to_float(market.get("currentPrice"))
    shares = _to_float(market.get("sharesDiluted"))
    cash = _to_float(market.get("cash"))
    debt = _to_float(market.get("debt"))
    securities = _to_float(market.get("nonOperatingAssets"))
    minority_interest = _to_float(market.get("minorityInterest"))
    preferred_equity = _to_float(market.get("preferredEquity"))
    bridge_values = (current_price, shares, cash, debt, securities, minority_interest, preferred_equity)
    if any(value is None for value in bridge_values) or current_price <= 0 or shares <= 0:
        raise ValueError("Incomplete comparable schedule requires a complete common-equity bridge.")

    requirements = payload.get("requiredInputs")
    requirements = requirements if isinstance(requirements, list) else []
    requirement_by_ticker: dict[str, dict[str, Any]] = {}
    for item in requirements:
        if not isinstance(item, dict) or not str(item.get("key") or "").startswith(f"peer_{metric}:"):
            continue
        requirement_by_ticker[str(item["key"]).split(":", 1)[1].upper()] = item
    peers = payload.get("comps") if isinstance(payload.get("comps"), list) else []
    target_ticker = str(company.get("ticker") or "").strip().upper()
    peer_rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    for peer in peers:
        if not isinstance(peer, dict):
            continue
        ticker = str(peer.get("ticker") or "").strip().upper()
        enterprise_value = _to_float(peer.get("ev"))
        denominator = _to_float(peer.get(metric))
        requirement = requirement_by_ticker.get(ticker)
        revenue = _to_float(peer.get("revenue"))
        ebitda = _to_float(peer.get("ebitda"))
        if not ticker or ticker == target_ticker or ticker in seen or enterprise_value is None or enterprise_value <= 0:
            continue
        if metric == "ebitda" and revenue is not None and revenue <= 0:
            continue
        if metric == "revenue" and ebitda is not None and ebitda <= 0:
            continue
        if denominator is None or denominator <= 0:
            if requirement is None:
                continue
        else:
            multiple = enterprise_value / denominator
            if not math.isfinite(multiple) or multiple <= 0 or multiple >= 100:
                continue
        seen.add(ticker)
        if peer_fallback_mode and requirement is None:
            # Fallback market data never enters the schedule unconfirmed: every
            # usable fallback peer must carry an analyst confirmation requirement.
            continue
        peer_rows.append({
            "ticker": ticker,
            "name": str(peer.get("company") or ticker),
            "ev": enterprise_value,
            "revenue": revenue,
            "ebitda": ebitda,
            "requirement": requirement,
        })
    if len(peer_rows) < 3:
        raise ValueError("Incomplete comparable workbook requires three curated peer rows, including the missing denominator.")

    divisor = _resolve_amount_scale_divisor(payload)
    for existing in list(workbook.worksheets):
        workbook.remove(existing)
    model = workbook.create_sheet("Comparable Valuation")
    model.sheet_view.showGridLines = False
    navy, medium_blue, pale_yellow, light_blue, white = "17365D", "365F91", "FFF2CC", "DDEBF7", "FFFFFF"
    thin_gray = Side(style="thin", color="B7C9D6")
    amount_format, multiple_format, price_format = "#,##0.0;(#,##0.0);-", "0.0x", "$0.00;($0.00);-"
    status_ref = "'Input Required'!$B$3"
    title = "EV/EBITDA Trading Comparables" if valuation_model == "ev_ebitda" else "EV/Revenue Trading Comparables"
    ticker = str(company.get("ticker") or "").strip().upper()
    model["A1"] = f"{title} — {ticker}"
    model["A1"].font = Font(name="Arial", size=16, bold=True, color=white)
    model["A1"].fill = PatternFill("solid", fgColor=navy)
    model.merge_cells("A1:G1")
    model["A2"] = str(company.get("name") or ticker)
    model["A2"].font = Font(name="Arial", size=11, bold=True, color=navy)
    model["A3"] = "Missing current peer facts are entered on Input Required; peer medians and valuation remain blank until all inputs pass source and range checks."
    model.merge_cells("A3:G3")
    model["A3"].alignment = Alignment(wrap_text=True)
    model["A4"] = f"Target FY{years[-1]} {metric_name}"
    model["B4"] = _scale(target_metric, divisor)
    model["B4"].number_format = amount_format
    model["B4"].comment = Comment(f"Latest filed target {metric_name}; see Data Review for source concept and accession.", "Codex")
    model["A5"] = "Peer median EV/EBITDA" if metric == "ebitda" else "Peer median EV/Revenue"
    multiple_column = "G" if metric == "ebitda" else "F"
    first_peer_row, last_peer_row = 24, 23 + len(peer_rows)
    model["B5"] = f'=IF({status_ref}<>"READY","",MEDIAN({multiple_column}{first_peer_row}:{multiple_column}{last_peer_row}))'
    model["B5"].number_format = multiple_format
    model["B5"].comment = Comment(f"Formula: median of the current peer {multiple_column} multiple column. It remains blank until all required inputs are ready.", "Codex")
    model["A6"] = "Optional analyst multiple override"
    model["B6"] = None
    model["B6"].font = Font(name="Arial", size=10, color="0000FF")
    model["B6"].fill = PatternFill("solid", fgColor=pale_yellow)
    model["B6"].number_format = multiple_format
    model["B7"] = f'=IF({status_ref}<>"READY","",IF(B6>0,B6,B5))'
    model["A7"] = "Selected multiple"
    model["B7"].number_format = multiple_format
    model["B7"].comment = Comment("Formula: use the optional override in B6 when positive; otherwise use the peer median in B5.", "Codex")
    model["A8"] = "Enterprise value"
    model["B8"] = f'=IF({status_ref}<>"READY","",B4*B7)'
    model["B8"].number_format = amount_format
    for row, label, value in (
        (10, "Cash and cash equivalents", cash),
        (11, "Marketable securities", securities),
        (12, "Debt", debt),
        (13, "Noncontrolling interest", minority_interest),
        (14, "Preferred equity", preferred_equity),
    ):
        model.cell(row=row, column=1, value=label)
        model.cell(row=row, column=2, value=_scale(value, divisor))
        model.cell(row=row, column=2).number_format = amount_format
    model["A15"] = "Common Equity Value"
    model["B15"] = f'=IF({status_ref}<>"READY","",B8+B10+B11-B12-B13-B14)'
    model["B15"].number_format = amount_format
    model["B15"].comment = Comment("Formula: enterprise value plus cash and marketable securities, less debt, noncontrolling interest, and preferred equity.", "Codex")
    model["A16"], model["B16"] = "Diluted Shares", shares / divisor
    model["B16"].number_format = amount_format
    model["A17"] = "Implied Share Price"
    model["B17"] = f'=IF({status_ref}<>"READY","",IFERROR(B15/B16,0))'
    model["B17"].number_format = price_format
    model["A18"], model["B18"] = "Current Share Price", current_price
    model["B18"].number_format = price_format
    model["A19"] = "Implied Upside / (Downside)"
    model["B19"] = f'=IF({status_ref}<>"READY","",IFERROR(B17/B18-1,0))'
    model["B19"].number_format = "0.0%;(0.0%);-"
    model["A22"] = "Current Peer Set"
    model["A22"].font = Font(name="Arial", size=11, bold=True, color=navy)
    headers = ("Ticker", "Company", "Enterprise Value", "LTM Revenue", "LTM EBITDA", "EV/Revenue", "EV/EBITDA")
    for column, header in enumerate(headers, start=1):
        cell = model.cell(row=23, column=column, value=header)
        cell.font = Font(name="Arial", size=10, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=medium_blue)
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row, peer in enumerate(peer_rows, start=first_peer_row):
        model.cell(row=row, column=1, value=peer["ticker"])
        model.cell(row=row, column=2, value=peer["name"])
        model.cell(row=row, column=3, value=_scale(peer["ev"], divisor))
        requirement = peer["requirement"]
        revenue = peer["revenue"]
        ebitda = peer["ebitda"]
        model.cell(row=row, column=4, value=_scale(revenue, divisor) if revenue is not None and metric != "revenue" else None)
        model.cell(row=row, column=5, value=_scale(ebitda, divisor) if ebitda is not None and metric != "ebitda" else None)
        if metric == "revenue" and requirement is None and revenue is not None:
            model.cell(row=row, column=4, value=_scale(revenue, divisor))
        if metric == "ebitda" and requirement is None and ebitda is not None:
            model.cell(row=row, column=5, value=_scale(ebitda, divisor))
        for column in (3, 4, 5):
            model.cell(row=row, column=column).number_format = amount_format
        for column, numerator, denominator in ((6, 3, 4), (7, 3, 5)):
            cell = model.cell(row=row, column=column)
            cell.value = f'=IF(OR({status_ref}<>"READY",{get_column_letter(denominator)}{row}<=0),"",{get_column_letter(numerator)}{row}/{get_column_letter(denominator)}{row})'
            cell.number_format = multiple_format
            cell.font = Font(name="Arial", size=10, color="000000")

    model["I4"] = "Multiple Sensitivity"
    for column, header in enumerate(("Change", "Multiple", "Enterprise Value", "Common Equity", "Per Share"), start=9):
        cell = model.cell(row=5, column=column, value=header)
        cell.font = Font(name="Arial", size=9, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=medium_blue)
        cell.alignment = Alignment(wrap_text=True)
    for row, change in enumerate((-0.2, 0, 0.2), start=6):
        model.cell(row=row, column=9, value=change)
        model.cell(row=row, column=9).font = Font(name="Arial", size=10, color="0000FF")
        model.cell(row=row, column=9).fill = PatternFill("solid", fgColor=pale_yellow)
        model.cell(row=row, column=9).number_format = "0%;(0%);-"
        formulas = {
            10: f'=$B$7*(1+I{row})', 11: f'=$B$4*J{row}',
            12: f'=K{row}+$B$10+$B$11-$B$12-$B$13-$B$14',
            13: f'=IFERROR(L{row}/$B$16,0)',
        }
        for column, formula in formulas.items():
            cell = model.cell(row=row, column=column, value=f'=IF({status_ref}<>"READY","",{formula[1:]})')
            cell.font = Font(name="Arial", size=10, color="000000")
        model.cell(row=row, column=10).number_format = multiple_format
        for column in (11, 12):
            model.cell(row=row, column=column).number_format = amount_format
        model.cell(row=row, column=13).number_format = price_format
    # Add the generic register, then link each missing denominator to its peer's row.
    input_cells = apply_incomplete_workbook(workbook, payload, preserve_existing=True)
    for row, peer in enumerate(peer_rows, start=first_peer_row):
        requirement = peer["requirement"]
        if requirement is None:
            continue
        identity = f"{requirement['key']}:{requirement.get('asOfDate') or ''}"
        destination = input_cells.get(identity)
        if not destination:
            raise ValueError(f"Missing comparable peer input cell for {peer['ticker']}.")
        ref = f"'Input Required'!{destination['cell']}"
        denominator_column = "E" if metric == "ebitda" else "D"
        model[f"{denominator_column}{row}"] = f'=IF(AND({status_ref}="READY",ISNUMBER({ref}),{ref}>0),{ref}/{divisor},"")'
        model[f"{denominator_column}{row}"].font = Font(name="Arial", size=10, color="008000")
        model[f"{denominator_column}{row}"].number_format = amount_format

    for row in (7, 8, 15, 17, 19):
        for column in (1, 2):
            cell = model.cell(row=row, column=column)
            cell.font = Font(name="Arial", size=10, bold=True, color="000000")
            cell.fill = PatternFill("solid", fgColor=light_blue)
    for row in range(4, 20):
        for column in (1, 2):
            cell = model.cell(row=row, column=column)
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=thin_gray)
    for column, width in {"A": 32, "B": 28, "C": 18, "D": 16, "E": 16, "F": 14, "G": 14, "I": 12, "J": 12, "K": 20, "L": 20, "M": 16}.items():
        model.column_dimensions[column].width = width
    model.freeze_panes = "A24"
    model.auto_filter.ref = f"A23:G{23 + len(peer_rows)}"
