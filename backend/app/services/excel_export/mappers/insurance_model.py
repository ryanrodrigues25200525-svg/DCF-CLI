from __future__ import annotations

import math
from typing import Any

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

_INPUT_FONT = Font(name="Arial", size=10, color="0000FF")
_FORMULA_FONT = Font(name="Arial", size=10, color="000000")
_BODY_FONT = Font(name="Arial", size=10, color="000000")
_HEADER_FONT = Font(name="Arial", size=9, bold=True, color="FFFFFF")
_SECTION_FONT = Font(name="Arial", size=10, bold=True, color="17365D")
_TITLE_FILL = PatternFill("solid", fgColor="17365D")
_SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
_INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
_FORMULA_FILL = PatternFill("solid", fgColor="E2F0D9")
_HEADER_FILL = PatternFill("solid", fgColor="365F91")
_MONEY_FORMAT = '#,##0.0,,;[Red](#,##0.0,,);-'
_PRICE_FORMAT = '$#,##0.00;[Red]($#,##0.00);-'
_PERCENT_FORMAT = '0.0%;[Red](0.0%);-'

_HISTORICAL_LINES = (
    ("net_premiums_written", 5, "Net premiums written", _MONEY_FORMAT),
    ("net_premiums_earned", 6, "Net premiums earned", _MONEY_FORMAT),
    ("losses_and_lae", 7, "Losses and loss-adjustment expenses", _MONEY_FORMAT),
    ("acquisition_expenses", 8, "Acquisition expenses", _MONEY_FORMAT),
    ("general_operating_expenses", 9, "General operating expenses", _MONEY_FORMAT),
    ("underwriting_expenses", 10, "Total underwriting expenses", _MONEY_FORMAT),
    ("loss_ratio", 11, "Reported loss ratio", _PERCENT_FORMAT),
    ("expense_ratio", 12, "Reported expense ratio", _PERCENT_FORMAT),
    ("combined_ratio", 13, "Reported combined ratio", _PERCENT_FORMAT),
    ("prior_year_reserve_development", 14, "Prior-year reserve development", _PERCENT_FORMAT),
    ("underwriting_income", 15, "Underwriting income", _MONEY_FORMAT),
    ("net_investment_income", 16, "General Insurance net investment income", _MONEY_FORMAT),
    ("invested_assets", 17, "General Insurance invested assets", _MONEY_FORMAT),
    ("unpaid_loss_reserves_beginning", 18, "Opening net loss reserves", _MONEY_FORMAT),
    ("losses_incurred_for_reserve_rollforward", 19, "Losses incurred for reserve roll-forward", _MONEY_FORMAT),
    ("losses_paid_for_reserve_rollforward", 20, "Losses paid for reserve roll-forward", _MONEY_FORMAT),
    ("reserve_other_changes", 21, "Other reserve changes", _MONEY_FORMAT),
    ("unpaid_loss_reserves", 22, "Ending net loss reserves", _MONEY_FORMAT),
    ("reinsurance_recoverable", 23, "Reinsurance recoverables", _MONEY_FORMAT),
    ("gross_loss_reserves", 24, "Ending gross loss reserves", _MONEY_FORMAT),
    ("reserve_rollforward_check", 25, "Filed reserve roll-forward check", _MONEY_FORMAT),
    ("other_operations_pretax_income", 26, "Other Operations pretax income", _MONEY_FORMAT),
    ("statutory_capital_surplus", 27, "Statutory capital and surplus", _MONEY_FORMAT),
    ("minimum_statutory_capital", 28, "Minimum required statutory capital", _MONEY_FORMAT),
    ("common_equity", 29, "Common equity attributable to parent", _MONEY_FORMAT),
    ("common_equity_distributions", 30, "Common dividends and repurchases", _MONEY_FORMAT),
    ("diluted_shares", 31, "Weighted-average diluted shares", _MONEY_FORMAT),
    ("net_income", 32, "Net income available to common", _MONEY_FORMAT),
    ("tax_rate", 33, "Effective tax rate", _PERCENT_FORMAT),
    ("reported_pretax_income", 34, "Reported consolidated pretax income", _MONEY_FORMAT),
    ("other_pretax_adjustments", 35, "Other pretax reconciliation adjustments", _MONEY_FORMAT),
)

_ASSUMPTIONS = (
    ("premiumGrowth", 41, "Net written and earned premium growth", _PERCENT_FORMAT),
    ("lossRatio", 42, "Reported loss ratio starting point", _PERCENT_FORMAT),
    ("expenseRatio", 43, "Reported expense ratio", _PERCENT_FORMAT),
    ("priorYearReserveDevelopmentRate", 44, "Forward prior-year reserve development", _PERCENT_FORMAT),
    ("netInvestmentIncomeYield", 45, "Net investment income yield", _PERCENT_FORMAT),
    ("investedAssetGrowth", 46, "Invested asset growth", _PERCENT_FORMAT),
    ("paidLossRatio", 47, "Paid losses / incurred losses", _PERCENT_FORMAT),
    ("reserveOtherChangesRate", 48, "Other reserve changes / earned premium", _PERCENT_FORMAT),
    ("otherOperationsPretaxGrowth", 49, "Other Operations pretax growth", _PERCENT_FORMAT),
    ("otherPretaxAdjustmentGrowth", 50, "Other pretax adjustments growth", _PERCENT_FORMAT),
    ("taxRate", 51, "Tax rate", _PERCENT_FORMAT),
    ("payoutRatio", 52, "Target common distribution / net income", _PERCENT_FORMAT),
    ("minimumStatutoryCapitalToPremiumRatio", 53, "Minimum statutory capital / written premium", _PERCENT_FORMAT),
    ("riskFreeRate", 54, "Risk-free rate", _PERCENT_FORMAT),
    ("equityRiskPremium", 55, "Equity risk premium", _PERCENT_FORMAT),
    ("beta", 56, "Beta", "0.00"),
    ("terminalGrowthRate", 58, "Terminal residual-income growth", _PERCENT_FORMAT),
)


def _record(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _finite_number(value: Any, label: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(float(value)):
        raise ValueError(f"Insurance model export requires a finite sourced value for {label}.")
    return float(value)


def _latest_history(insurance_model: dict[str, Any]) -> tuple[int, dict[str, Any]]:
    history = _record(insurance_model.get("history"))
    annual = history.get("annual")
    if not isinstance(annual, list) or len(annual) < 3:
        raise ValueError("Insurance model export requires three filed annual periods.")
    latest = _record(annual[-1])
    year = latest.get("year")
    if isinstance(year, bool) or not isinstance(year, int):
        raise ValueError("Insurance model export requires an integer latest fiscal year.")
    insurance = _record(latest.get("insurance"))
    return year, insurance


def _source_line(insurance: dict[str, Any], field: str, year: int) -> tuple[float, dict[str, Any]]:
    line = _record(insurance.get(field))
    value = _finite_number(line.get("value"), field)
    if line.get("source") not in {"sec_native", "derived"}:
        raise ValueError(f"FY{year} insurance model input {field} is missing or ambiguous.")
    sources = line.get("sources")
    if not isinstance(sources, list) or not sources:
        raise ValueError(f"FY{year} insurance model input {field} has no filing provenance.")
    for source in sources:
        source = _record(source)
        if not source.get("accession") or not source.get("filed") or source.get("fiscal_period") != f"FY {year}":
            raise ValueError(f"FY{year} insurance model input {field} has incomplete filing provenance.")
    return value, line


def _source_metadata(line: dict[str, Any]) -> tuple[str, str, str, str, str, str]:
    sources = [_record(source) for source in line.get("sources", [])]
    concepts = "; ".join(dict.fromkeys(str(source.get("concept") or line.get("concept") or "derived") for source in sources))
    accessions = "; ".join(dict.fromkeys(str(source.get("accession") or "") for source in sources if source.get("accession")))
    filed_dates = "; ".join(dict.fromkeys(str(source.get("filed") or "") for source in sources if source.get("filed")))
    periods = "; ".join(dict.fromkeys(str(source.get("fiscal_period") or "") for source in sources if source.get("fiscal_period")))
    units = "; ".join(dict.fromkeys(
        f"{source.get('unit') or 'unit unknown'} {source.get('unit_scale') or ''}".strip()
        for source in sources
    ))
    return concepts, accessions, filed_dates, periods, units, str(line.get("method") or "method not provided")


def _write_value(
    sheet: Worksheet,
    cell_ref: str,
    value: Any,
    *,
    formula: bool = False,
    input_cell: bool = False,
    number_format: str | None = None,
) -> None:
    cell = sheet[cell_ref]
    cell.value = value
    cell.font = _INPUT_FONT if input_cell else _FORMULA_FONT if formula else _BODY_FONT
    cell.fill = _INPUT_FILL if input_cell else _FORMULA_FILL if formula else PatternFill(fill_type=None)
    cell.alignment = Alignment(vertical="center", wrap_text=False)
    if number_format:
        cell.number_format = number_format


def _section(sheet: Worksheet, row: int, label: str, end_column: int = 7) -> None:
    for column in range(1, end_column + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = _SECTION_FILL
        cell.font = _SECTION_FONT
    sheet.cell(row=row, column=1, value=label)


def _validate_assumptions(assumptions: dict[str, Any]) -> None:
    if _finite_number(assumptions.get("forecastYears"), "forecastYears") != 5:
        raise ValueError("Insurance model workbook requires a five-year forecast.")
    required_sources = (
        "premiumGrowth", "lossRatio", "expenseRatio", "priorYearReserveDevelopmentRate",
        "netInvestmentIncomeYield", "investedAssetGrowth", "paidLossRatio", "reserveOtherChangesRate",
        "otherOperationsPretaxGrowth", "otherPretaxAdjustmentGrowth", "taxRate", "payoutRatio",
        "minimumStatutoryCapitalToPremiumRatio", "terminalGrowthRate",
    )
    sources = _record(assumptions.get("assumptionSources"))
    for field in required_sources:
        if not str(sources.get(field) or "").strip():
            raise ValueError(f"Insurance model assumption {field} must include its source or analyst-input disclosure.")
    risk_free = _finite_number(assumptions.get("riskFreeRate"), "riskFreeRate")
    erp = _finite_number(assumptions.get("equityRiskPremium"), "equityRiskPremium")
    beta = _finite_number(assumptions.get("beta"), "beta")
    growth = _finite_number(assumptions.get("terminalGrowthRate"), "terminalGrowthRate")
    if beta <= 0 or risk_free + beta * erp <= growth:
        raise ValueError("Insurance model workbook requires cost of equity to exceed terminal growth.")


def _map_model_sheet(workbook: Workbook, payload: dict[str, Any], insurance_model: dict[str, Any]) -> None:
    for existing in list(workbook.worksheets):
        workbook.remove(existing)
    sheet = workbook.create_sheet("Insurance Model")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C64"
    sheet.column_dimensions["A"].width = 44
    sheet.column_dimensions["B"].width = 20
    for column in "CDEFG":
        sheet.column_dimensions[column].width = 18
    for column, width in (("H", 34), ("I", 24), ("J", 14), ("K", 14), ("L", 32)):
        sheet.column_dimensions[column].width = width
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"

    company = _record(payload.get("company"))
    ticker = str(company.get("ticker") or "").strip().upper()
    name = str(company.get("name") or ticker or "Company").strip()
    currency = str(company.get("currency") or "USD")
    year, insurance = _latest_history(insurance_model)
    assumptions = _record(insurance_model.get("assumptions"))
    _validate_assumptions(assumptions)

    sheet.merge_cells("A1:L1")
    sheet["A1"] = f"{name} ({ticker}) — P&C Insurance Residual Income Model"
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = _TITLE_FILL
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28
    sheet.merge_cells("A2:L2")
    sheet["A2"] = f"Common-equity valuation. Amounts are {currency} actual values displayed in millions, except per share, ratios, and shares. Blue cells are editable assumptions; forecast and valuation cells are formulas."
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="404040")
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[2].height = 30

    _section(sheet, 3, f"Filed FY{year} reported operating, reserve, capital, and common-equity inputs", 8)
    for column, header in enumerate(("Reported input", f"FY{year} value", "SEC concept(s)", "Accession(s)", "Filed date(s)", "Fiscal period(s)", "Unit(s)", "Canonical method"), start=1):
        cell = sheet.cell(row=4, column=column, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for field, row, label, number_format in _HISTORICAL_LINES:
        raw_line = _record(insurance.get(field))
        raw_sources = raw_line.get("sources")
        source_records = raw_sources if isinstance(raw_sources, list) else []
        reserve_is_filed = (
            raw_line.get("source") in {"sec_native", "derived"}
            and isinstance(raw_line.get("value"), (int, float))
            and not isinstance(raw_line.get("value"), bool)
            and math.isfinite(float(raw_line.get("value")))
            and bool(source_records)
            and all(
                isinstance(source, dict)
                and source.get("accession")
                and source.get("filed")
                and source.get("fiscal_period") == f"FY {year}"
                for source in source_records
            )
        )
        missing_latest_reserve = (
            payload.get("buildStatus") == "input_required"
            and field == "unpaid_loss_reserves"
            and not reserve_is_filed
        )
        if missing_latest_reserve:
            value, line = None, raw_line
        else:
            value, line = _source_line(insurance, field, year)
        sheet.cell(row=row, column=1, value=label).font = _BODY_FONT
        _write_value(sheet, f"B{row}", value, number_format=number_format)
        if missing_latest_reserve:
            sheet[f"B{row}"].comment = Comment(
                "Missing filed ending net loss reserves. Enter the reported FY value and source on Input Required.",
                "DCF Builder Pro",
            )
            sheet.row_dimensions[row].height = 30
            continue
        concepts, accessions, filed_dates, periods, units, method = _source_metadata(line)
        for column, item in enumerate((concepts, accessions, filed_dates, periods, units, method), start=3):
            cell = sheet.cell(row=row, column=column, value=item)
            cell.font = Font(name="Arial", size=8, color="404040")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet[f"B{row}"].comment = Comment(f"Canonical method: {method}. Filing accession(s): {accessions}.", "DCF Builder Pro")
        sheet.row_dimensions[row].height = 30

    _section(sheet, 39, "Editable forecast and valuation assumptions", 2)
    for field, row, label, number_format in _ASSUMPTIONS:
        sheet.cell(row=row, column=1, value=label).font = _BODY_FONT
        _write_value(sheet, f"B{row}", _finite_number(assumptions.get(field), field), input_cell=True, number_format=number_format)
        source = _record(assumptions.get("assumptionSources")).get(field)
        sheet[f"B{row}"].comment = Comment(str(source or "Source not provided."), "DCF Builder Pro")
    for row, label, field, number_format in (
        (54, "Risk-free rate", "riskFreeRate", _PERCENT_FORMAT),
        (55, "Equity risk premium", "equityRiskPremium", _PERCENT_FORMAT),
        (56, "Beta", "beta", "0.00"),
    ):
        sheet[f"A{row}"] = label
        sheet[f"A{row}"].font = _BODY_FONT
        _write_value(sheet, f"B{row}", _finite_number(assumptions.get(field), field), input_cell=True, number_format=number_format)
    for row, field in ((54, "riskFreeRateSource"), (55, "equityRiskPremiumSource"), (56, "betaSource"), (53, "minimumStatutoryCapitalSource")):
        source = str(assumptions.get(field) or "").strip()
        if not source:
            raise ValueError(f"Insurance model assumption {field} must identify its source.")
        if row == 53:
            ratio_source = _record(assumptions.get("assumptionSources")).get("minimumStatutoryCapitalToPremiumRatio")
            source = f"Ratio: {ratio_source}. Filed minimum required capital: {source}."
        sheet[f"B{row}"].comment = Comment(source, "DCF Builder Pro")
    _write_value(sheet, "B57", "=B54+B55*B56", formula=True, number_format=_PERCENT_FORMAT)
    sheet["A57"] = "Cost of equity (CAPM)"
    sheet["A57"].font = _BODY_FONT
    _write_value(sheet, "B59", _finite_number(assumptions.get("currentPrice"), "currentPrice"), input_cell=True, number_format=_PRICE_FORMAT)
    sheet["A59"] = "Current share price"
    sheet["A59"].font = _BODY_FONT
    sheet["B59"].comment = Comment(f"Current market quote, {assumptions.get('betaSource')}; valuation context as of {assumptions.get('marketDataAsOfDate')}.", "DCF Builder Pro")
    _write_value(sheet, "B60", _finite_number(assumptions.get("dilutedSharesOutstanding"), "dilutedSharesOutstanding"), input_cell=True, number_format=_MONEY_FORMAT)
    sheet["A60"] = "Diluted shares for valuation"
    sheet["A60"].font = _BODY_FONT
    sheet["B60"].comment = Comment(f"Filed weighted-average diluted shares from FY{year}; see the sourced input table.", "DCF Builder Pro")

    _section(sheet, 62, "Five-year underwriting, reserve, capital, and residual-income forecast", 7)
    for index, column in enumerate("CDEFG", start=1):
        sheet[f"{column}64"] = f"FY{year + index}E"
        sheet[f"{column}64"].font = _HEADER_FONT
        sheet[f"{column}64"].fill = _HEADER_FILL
        sheet[f"{column}64"].alignment = Alignment(horizontal="center")
    forecast_labels = {
        65: "Net premiums written",
        66: "Net premiums earned",
        67: "Forward reserve-development assumption",
        68: "Loss ratio",
        69: "Expense ratio",
        70: "Combined ratio",
        71: "Losses and LAE",
        72: "Underwriting expenses",
        73: "Underwriting income",
        74: "Ending invested assets",
        75: "Average invested assets",
        76: "Net investment income",
        77: "Other Operations pretax income",
        78: "Other pretax adjustments",
        79: "Pretax income",
        80: "Tax expense",
        81: "Net income",
        82: "Opening net loss reserves",
        83: "Losses incurred for reserves",
        84: "Losses paid",
        85: "Other reserve changes",
        86: "Ending net loss reserves",
        87: "Ending reinsurance recoverable",
        88: "Ending gross loss reserves",
        89: "Reserve roll-forward check",
        90: "Opening statutory capital and surplus",
        91: "Required statutory capital proxy",
        92: "Statutory capital before distributions",
        93: "Distribution capacity after minimum capital",
        94: "Target common distributions",
        95: "Allowed common distributions",
        96: "Ending statutory capital and surplus",
        97: "Minimum-capital check",
        98: "Opening common equity",
        99: "Ending common equity",
        100: "Equity charge",
        101: "Residual income",
        102: "Present value of residual income",
    }
    for row, label in forecast_labels.items():
        sheet.cell(row=row, column=1, value=label).font = _BODY_FONT
    for index, column in enumerate("CDEFG", start=1):
        previous = chr(ord(column) - 1) if index > 1 else None
        prior = lambda row, historical_row: f"{previous}{row}" if previous else f"$B${historical_row}"
        formulas = {
            65: f"={prior(65, 5)}*(1+$B$41)",
            66: f"={prior(66, 6)}*(1+$B$41)",
            67: "=$B$44",
            68: "=MAX(0,$B$42+$B$14-$B$44)",
            69: "=$B$43",
            70: f"={column}68+{column}69",
            71: f"={column}66*{column}68",
            72: f"={column}66*{column}69",
            73: f"={column}66-{column}71-{column}72",
            74: f"={prior(74, 17)}*(1+$B$46)",
            75: f"=AVERAGE({prior(74, 17)},{column}74)",
            76: f"={column}75*$B$45",
            77: f"={prior(77, 26)}*(1+$B$49)",
            78: f"={prior(78, 35)}*(1+$B$50)",
            79: f"={column}73+{column}76+{column}77+{column}78",
            80: f"=MAX(0,{column}79*$B$51)",
            81: f"={column}79-{column}80",
            82: f"={prior(86, 22)}",
            83: f"={column}71",
            84: f"={column}83*$B$47",
            85: f"={column}66*$B$48",
            86: f"={column}82+{column}83-{column}84+{column}85",
            87: f"={column}86*$B$23/$B$22",
            88: f"={column}86+{column}87",
            89: f"={column}82+{column}83-{column}84+{column}85-{column}86",
            90: f"={prior(96, 27)}",
            91: f"={column}65*$B$53",
            92: f"={column}90+{column}81",
            93: f"={column}92-{column}91",
            94: f"=MAX(0,{column}81*$B$52)",
            95: f"=IF({column}93<0,0,MIN({column}94,{column}93))",
            96: f"={column}92-{column}95",
            97: f'=IF({column}93>=0,"PASS","BREACH")',
            98: f"={prior(99, 29)}",
            99: f"={column}98+{column}81-{column}95",
            100: f"={column}98*$B$57",
            101: f"={column}81-{column}100",
            102: f"={column}101/(1+$B$57)^{index}",
        }
        for row, formula in formulas.items():
            number_format = _PERCENT_FORMAT if row in (67, 68, 69, 70) else _MONEY_FORMAT
            _write_value(sheet, f"{column}{row}", formula, formula=True, number_format=number_format)

    _section(sheet, 104, "Common-equity residual-income valuation", 2)
    valuation_rows = (
        (105, "PV of forecast residual income", "=SUM(C102:G102)", _MONEY_FORMAT),
        (106, "Terminal residual income", '=IF($B$57<=$B$58,"",G101*(1+$B$58)/($B$57-$B$58))', _MONEY_FORMAT),
        (107, "PV of terminal residual income", '=IF(B106="","",B106/(1+$B$57)^5)', _MONEY_FORMAT),
        (108, "Common equity value", '=IF(B107="","",MAX(0,$B$29+B105+B107))', _MONEY_FORMAT),
        (109, "Implied share price", '=IFERROR(B108/$B$60,0)', _PRICE_FORMAT),
        (110, "Current share price", "=B59", _PRICE_FORMAT),
        (111, "Implied return", "=IFERROR(B109/B110-1,0)", _PERCENT_FORMAT),
    )
    for row, label, formula, number_format in valuation_rows:
        sheet.cell(row=row, column=1, value=label).font = Font(name="Arial", size=10, bold=row in (108, 109, 111))
        _write_value(sheet, f"B{row}", formula, formula=True, number_format=number_format)
    sheet["A112"] = "Enterprise value"
    sheet["B112"] = "Not applicable — common-equity insurance model"
    sheet["A112"].font = _BODY_FONT
    sheet["B112"].font = _BODY_FONT

    _section(sheet, 114, "Common-equity sensitivity — cost of equity and terminal growth", 7)
    sheet["B115"] = "Cost of equity / terminal growth"
    sheet["B115"].font = _HEADER_FONT
    sheet["B115"].fill = _HEADER_FILL
    for column, formula in zip("CDEFG", ("=$B$58-2%", "=$B$58-1%", "=$B$58", "=$B$58+1%", "=$B$58+2%")):
        _write_value(sheet, f"{column}115", formula, formula=True, number_format=_PERCENT_FORMAT)
        sheet[f"{column}115"].fill = _HEADER_FILL
        sheet[f"{column}115"].font = _HEADER_FONT
    for row, offset in zip(range(116, 121), ("-2%", "-1%", "", "+1%", "+2%")):
        _write_value(sheet, f"B{row}", f"=$B$57{offset}" if offset else "=$B$57", formula=True, number_format=_PERCENT_FORMAT)
        forecast_pvs = [
            f"({column}81-$B{row}*{column}98)/(1+$B{row})^{index}"
            for index, column in enumerate("CDEFG", start=1)
        ]
        _write_value(sheet, f"H{row}", "=" + "+".join(forecast_pvs), formula=True, number_format=_MONEY_FORMAT)
        _write_value(sheet, f"I{row}", f"=G81-$B{row}*G98", formula=True, number_format=_MONEY_FORMAT)
        for column in "CDEFG":
            formula = (
                f'=IF(OR($B{row}<={column}$115,$B{row}<=0),"",'
                f'MAX(0,$B$29+$H{row}+($I{row}*(1+{column}$115)/($B{row}-{column}$115))/(1+$B{row})^5))'
            )
            _write_value(sheet, f"{column}{row}", formula, formula=True, number_format=_MONEY_FORMAT)
    sheet["A121"] = "Sensitivity basis"
    sheet["B121"] = "Opening common equity plus discounted forecast and terminal residual income; cells are blank when terminal growth is not below cost of equity."
    sheet.merge_cells("B121:I121")
    sheet["A121"].font = _BODY_FONT
    sheet["B121"].font = Font(name="Arial", size=9, italic=True, color="404040")
    sheet["B121"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[121].height = 28
    sheet.auto_filter.ref = "A5:G35"


def _map_review_sheet(workbook: Workbook, payload: dict[str, Any], insurance_model: dict[str, Any]) -> None:
    review = workbook.create_sheet("Data Review")
    review.sheet_view.showGridLines = False
    company = _record(payload.get("company"))
    ticker = str(company.get("ticker") or "").upper()
    name = str(company.get("name") or ticker)
    review.merge_cells("A1:C1")
    review["A1"] = f"Data Review — {name} ({ticker})"
    review["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    review["A1"].fill = _TITLE_FILL
    review.merge_cells("A2:C2")
    review["A2"] = "Review company definitions and source notes before using this P&C common-equity model."
    review["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    review["A2"].font = Font(name="Arial", size=10, italic=True)
    for column, label in enumerate(("Status", "Area", "Review note"), start=1):
        cell = review.cell(row=5, column=column, value=label)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    ui_meta = _record(payload.get("uiMeta"))
    warnings = ui_meta.get("warnings") if isinstance(ui_meta.get("warnings"), list) else []
    notes = ui_meta.get("sourceNotes") if isinstance(ui_meta.get("sourceNotes"), list) else []
    items = [("REVIEW", "Model limitation", note) for note in warnings if isinstance(note, str)]
    items.extend(("INFO", "SEC source", note) for note in notes if isinstance(note, str))
    items.append(("INFO", "Sensitivity", "The two-way sensitivity revalues residual income at each cost-of-equity and terminal-growth combination."))
    if not items:
        raise ValueError("Insurance model export requires source notes and assumption disclosures on Data Review.")
    for row, (status, area, note) in enumerate(items, start=6):
        status_cell = review.cell(row=row, column=1, value=status)
        status_cell.font = Font(name="Arial", size=10, bold=True)
        status_cell.fill = PatternFill("solid", fgColor="FCE4D6" if status == "REVIEW" else "DDEBF7")
        review.cell(row=row, column=2, value=area)
        review.cell(row=row, column=3, value=note)
        for column in range(1, 4):
            review.cell(row=row, column=column).alignment = Alignment(wrap_text=True, vertical="top")
        review.row_dimensions[row].height = 34
    review.column_dimensions["A"].width = 14
    review.column_dimensions["B"].width = 22
    review.column_dimensions["C"].width = 112
    review.freeze_panes = "A6"
    review.auto_filter.ref = f"A5:C{5 + len(items)}"


def apply_insurance_model(workbook: Workbook, payload: dict[str, Any]) -> None:
    if payload.get("valuationModel") != "insurance_pnc_residual_income":
        raise ValueError("Insurance workbook requires valuationModel=insurance_pnc_residual_income.")
    insurance_model = _record(payload.get("insuranceModel"))
    if not insurance_model:
        raise ValueError("Insurance workbook requires dedicated insurance history and assumptions.")
    _map_model_sheet(workbook, payload, insurance_model)
    _map_review_sheet(workbook, payload, insurance_model)


def apply_incomplete_insurance_model(
    workbook: Workbook,
    payload: dict[str, Any],
    input_cells: dict[str, dict[str, str]],
) -> None:
    """Link the missing filed reserve balance and withhold dependent insurance outputs."""
    sheet = workbook["Insurance Model"]
    requirements = payload.get("requiredInputs")
    requirements = requirements if isinstance(requirements, list) else []
    requirement = next((item for item in requirements if isinstance(item, dict) and item.get("key") == "unpaid_loss_reserves"), None)
    if requirement is None:
        raise ValueError("Incomplete P&C workbook requires an unpaid_loss_reserves input.")
    identity = f"unpaid_loss_reserves:{requirement.get('fiscalYear') or requirement.get('asOfDate') or ''}"
    destination = input_cells.get(identity)
    if not destination or destination.get("sheet") != "Input Required":
        raise ValueError("Incomplete P&C workbook has no editable ending net loss reserve cell.")
    value_ref = f"'Input Required'!{destination['cell']}"
    status_ref = "'Input Required'!$B$3"
    sheet["B22"] = f'=IF(AND(ISNUMBER({value_ref}),{value_ref}>0),{value_ref},"")'
    sheet["B22"].font = Font(name="Arial", size=10, color="008000")
    sheet["B22"].fill = _FORMULA_FILL
    sheet["B22"].number_format = _MONEY_FORMAT
    sheet["B22"].comment = Comment(
        f"Linked to Input Required!{destination['cell']}; enter the filed ending net loss reserve and source reference there.",
        "DCF Builder Pro",
    )
    sheet["B25"] = f'=IF({status_ref}<>"READY","",B18+B19-B20+B21-B22)'
    sheet["B25"].font = Font(name="Arial", size=10, color="008000")
    sheet["B25"].fill = _FORMULA_FILL
    sheet["B25"].number_format = _MONEY_FORMAT

    formula_cells = [
        f"{column}{row}"
        for column in "CDEFG"
        for row in range(65, 103)
    ]
    formula_cells.extend(f"B{row}" for row in (*range(105, 110), 111))
    formula_cells.extend(f"{column}{row}" for row in range(116, 121) for column in "CDEFGHI")
    for cell_ref in formula_cells:
        cell = sheet[cell_ref]
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.value = f'=IF({status_ref}<>"READY","",{cell.value[1:]})'
