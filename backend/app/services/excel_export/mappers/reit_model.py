from __future__ import annotations

import math
from typing import Any

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from .incomplete import gate_formulas_on_ready, require_input_destination

_INPUT_FONT = Font(name="Arial", size=10, color="0000FF")
_FORMULA_FONT = Font(name="Arial", size=10, color="000000")
_BODY_FONT = Font(name="Arial", size=10, color="000000")
_HEADER_FONT = Font(name="Arial", size=9, bold=True, color="FFFFFF")
_SECTION_FONT = Font(name="Arial", size=10, bold=True, color="17365D")
_TITLE_FILL = PatternFill("solid", fgColor="17365D")
_SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
_INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
_FORMULA_FILL = PatternFill("solid", fgColor="E2F0D9")
_HEADER_FILL = PatternFill("solid", fgColor="365F91")
_MONEY_FORMAT = '#,##0.0,,;[Red](#,##0.0,,);-'
_PRICE_FORMAT = '$#,##0.00;[Red]($#,##0.00);-'
_PERCENT_FORMAT = '0.0%;[Red](0.0%);-'
_RATIO_FORMAT = '0.00x'

_HISTORICAL_LINES = (
    ("nareit_bridge_net_income", 5, "Net earnings attributable to common in FFO bridge", _MONEY_FORMAT),
    ("real_estate_depreciation", 6, "Real estate depreciation and amortization", _MONEY_FORMAT),
    ("disposition_gains_nareit_adjustment", 7, "NAREIT property-disposition adjustment, net of tax", _MONEY_FORMAT),
    ("nci_nareit_adjustment", 8, "NAREIT noncontrolling-interest adjustment", _MONEY_FORMAT),
    ("unconsolidated_nareit_adjustment", 9, "NAREIT proportionate unconsolidated adjustment", _MONEY_FORMAT),
    ("nareit_ffo", 11, "Filed NAREIT FFO attributable to common", _MONEY_FORMAT),
    ("modified_ffo_fx_adjustment", 13, "Modified FFO FX, derivative, and other adjustment", _MONEY_FORMAT),
    ("modified_ffo_deferred_tax_adjustment", 14, "Modified FFO deferred-tax adjustment", _MONEY_FORMAT),
    ("modified_ffo_current_tax_adjustment", 15, "Modified FFO current-tax adjustment", _MONEY_FORMAT),
    ("modified_ffo_nci_adjustment", 16, "Modified FFO noncontrolling-interest adjustment", _MONEY_FORMAT),
    ("modified_ffo_unconsolidated_adjustment", 17, "Modified FFO unconsolidated-entity adjustment", _MONEY_FORMAT),
    ("modified_ffo", 19, "Filed FFO as modified by Prologis", _MONEY_FORMAT),
    ("core_ffo_disposition_adjustment", 21, "Core FFO development and land disposition adjustment", _MONEY_FORMAT),
    ("core_ffo_tax_adjustment", 22, "Core FFO current tax on dispositions", _MONEY_FORMAT),
    ("core_ffo_debt_extinguishment_adjustment", 23, "Core FFO debt-extinguishment adjustment", _MONEY_FORMAT),
    ("core_ffo_nci_adjustment", 24, "Core FFO noncontrolling-interest adjustment", _MONEY_FORMAT),
    ("core_ffo_unconsolidated_adjustment", 25, "Core FFO unconsolidated-entity adjustment", _MONEY_FORMAT),
    ("core_ffo", 27, "Filed Core FFO attributable to common", _MONEY_FORMAT),
    ("tenant_improvements_and_lease_commissions", 29, "Tenant improvements and leasing commissions", _MONEY_FORMAT),
    ("property_improvements", 30, "Property improvements", _MONEY_FORMAT),
    ("analyst_affo", 32, "Analyst-defined AFFO", _MONEY_FORMAT),
    ("same_store_noi_net_effective", 33, "Prologis-share same-store NOI, net effective", _MONEY_FORMAT),
    ("same_store_noi_cash", 34, "Prologis-share same-store NOI, cash basis", _MONEY_FORMAT),
    ("occupancy", 35, "Operating-portfolio occupancy", _PERCENT_FORMAT),
    ("real_estate_segment_noi", 36, "Real Estate segment NOI", _MONEY_FORMAT),
    ("strategic_capital_segment_noi", 37, "Strategic Capital segment NOI", _MONEY_FORMAT),
    ("common_distributions", 38, "Cash dividends on common and preferred stock", _MONEY_FORMAT),
)

_BALANCE_LINES = (
    ("netIncome", 39, "Net income available to common"),
    ("commonEquity", 40, "Common equity attributable to parent"),
    ("cash", 41, "Cash and cash equivalents"),
    ("longTermDebt", 42, "Reported long-term debt"),
    ("preferredEquity", 43, "Preferred equity"),
    ("nonControllingInterest", 44, "Noncontrolling-interest equity"),
    ("dilutedShares", 45, "Weighted-average diluted shares"),
)


def _record(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _finite_number(value: Any, label: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(float(value)):
        raise ValueError(f"REIT model export requires a finite sourced value for {label}.")
    return float(value)


def _source_line(line: Any, field: str, year: int) -> tuple[float, dict[str, Any]]:
    record = _record(line)
    value = _finite_number(record.get("value"), field)
    if record.get("source") not in {"sec_native", "derived"}:
        raise ValueError(f"FY{year} REIT model input {field} is missing or ambiguous.")
    sources = record.get("sources")
    if not isinstance(sources, list) or not sources:
        raise ValueError(f"FY{year} REIT model input {field} has no filing provenance.")
    for source in sources:
        source = _record(source)
        if not source.get("accession") or not source.get("filed"):
            raise ValueError(f"FY{year} REIT model input {field} has incomplete filing provenance.")
        if record.get("source") == "sec_native" and source.get("fiscal_period") != f"FY {year}":
            raise ValueError(f"FY{year} REIT model input {field} has mismatched fiscal-period provenance.")
    return value, record


def _source_metadata(line: dict[str, Any]) -> tuple[str, str, str, str, str, str]:
    sources = [_record(source) for source in line.get("sources", [])]
    concepts = "; ".join(dict.fromkeys(str(source.get("concept") or line.get("concept") or "derived") for source in sources))
    accessions = "; ".join(dict.fromkeys(str(source.get("accession") or "") for source in sources if source.get("accession")))
    filed_dates = "; ".join(dict.fromkeys(str(source.get("filed") or "") for source in sources if source.get("filed")))
    periods = "; ".join(dict.fromkeys(str(source.get("fiscal_period") or "") for source in sources if source.get("fiscal_period")))
    units = "; ".join(dict.fromkeys(
        f"{source.get('unit') or 'unit unknown'} {source.get('unit_scale') or ''}".strip()
        for source in sources
    ))
    return concepts, accessions, filed_dates, periods, units, str(line.get("method") or "method not provided")


def _write_value(
    sheet: Worksheet,
    cell_ref: str,
    value: Any,
    *,
    formula: bool = False,
    input_cell: bool = False,
    number_format: str | None = None,
) -> None:
    cell = sheet[cell_ref]
    cell.value = value
    cell.font = _INPUT_FONT if input_cell else _FORMULA_FONT if formula else _BODY_FONT
    cell.fill = _INPUT_FILL if input_cell else _FORMULA_FILL if formula else PatternFill(fill_type=None)
    cell.alignment = Alignment(vertical="center", wrap_text=False)
    if number_format:
        cell.number_format = number_format


def _section(sheet: Worksheet, row: int, label: str, end_column: int = 7) -> None:
    for column in range(1, end_column + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = _SECTION_FILL
        cell.font = _SECTION_FONT
    sheet.cell(row=row, column=1, value=label)


def _latest_history(reit_model: dict[str, Any]) -> tuple[list[dict[str, Any]], int]:
    history = _record(reit_model.get("history"))
    annual = history.get("annual")
    if not isinstance(annual, list) or len(annual) != 3:
        raise ValueError("REIT model export requires exactly three filed annual periods.")
    rows = [_record(item) for item in annual]
    years = history.get("years")
    if not isinstance(years, list) or len(years) != 3:
        raise ValueError("REIT model export requires three fiscal year labels.")
    if [row.get("year") for row in rows] != years:
        raise ValueError("REIT history fiscal years do not align.")
    latest_year = rows[-1].get("year")
    if isinstance(latest_year, bool) or not isinstance(latest_year, int):
        raise ValueError("REIT model export requires an integer latest fiscal year.")
    return rows, latest_year


def _validate_assumptions(assumptions: dict[str, Any]) -> None:
    if _finite_number(assumptions.get("forecastYears"), "forecastYears") != 5:
        raise ValueError("REIT model workbook requires a five-year forecast.")
    sources = _record(assumptions.get("assumptionSources"))
    for field in ("sameStoreNoiGrowth", "targetOccupancy", "recurringCapexRatio", "payoutRatio", "terminalGrowthRate", "navCapRate"):
        if not str(sources.get(field) or "").strip():
            raise ValueError(f"REIT model assumption {field} requires a source or explicit analyst-input disclosure.")
    risk_free = _finite_number(assumptions.get("riskFreeRate"), "riskFreeRate")
    erp = _finite_number(assumptions.get("equityRiskPremium"), "equityRiskPremium")
    beta = _finite_number(assumptions.get("beta"), "beta")
    growth = _finite_number(assumptions.get("terminalGrowthRate"), "terminalGrowthRate")
    if beta <= 0 or risk_free + beta * erp <= growth:
        raise ValueError("REIT model workbook requires cost of equity above terminal growth.")


def _map_model_sheet(workbook: Workbook, payload: dict[str, Any], reit_model: dict[str, Any]) -> None:
    for existing in list(workbook.worksheets):
        workbook.remove(existing)
    sheet = workbook.create_sheet("REIT Model")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "B68"
    sheet.column_dimensions["A"].width = 54
    sheet.column_dimensions["B"].width = 19
    for column in "CDEFG":
        sheet.column_dimensions[column].width = 17
    for column, width in (("H", 34), ("I", 24), ("J", 16), ("K", 16), ("L", 32)):
        sheet.column_dimensions[column].width = width
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"

    company = _record(payload.get("company"))
    ticker = str(company.get("ticker") or "").strip().upper()
    name = str(company.get("name") or ticker or "Company").strip()
    currency = str(company.get("currency") or "USD")
    annual, year = _latest_history(reit_model)
    assumptions = _record(reit_model.get("assumptions"))
    _validate_assumptions(assumptions)

    sheet.merge_cells("A1:L1")
    sheet["A1"] = f"{name} ({ticker}) — Equity REIT AFFO and NAV Model"
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = _TITLE_FILL
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28
    sheet.merge_cells("A2:L2")
    sheet["A2"] = f"Common-equity AFFO DCF with a property NAV cross-check. Amounts are {currency} actual values displayed in millions, except per share, ratios, and shares. Blue cells are editable assumptions; forecast and valuation cells are formulas."
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="404040")
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[2].height = 30

    _section(sheet, 3, "Filed annual operations and FFO/AFFO reconciliation", 4)
    for column, header in enumerate(("Reported measure / reconciliation item", "FY2023", "FY2024", f"FY{year}"), start=1):
        cell = sheet.cell(row=4, column=column, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for field, row, label, number_format in _HISTORICAL_LINES:
        sheet.cell(row=row, column=1, value=label).font = _BODY_FONT
        for column, item in enumerate(annual, start=2):
            reit = _record(item.get("reit"))
            line = _record(reit.get(field))
            amount = line.get("value")
            if isinstance(amount, (int, float)) and not isinstance(amount, bool) and math.isfinite(float(amount)):
                cell_ref = f"{chr(64 + column)}{row}"
                _write_value(sheet, cell_ref, float(amount), number_format=number_format)
                _source_line(line, field, int(item["year"]))
                concepts, accessions, filed_dates, periods, units, method = _source_metadata(line)
                sheet[cell_ref].comment = Comment(
                    f"Canonical method: {method}. SEC concept(s): {concepts}. Filing accession(s): {accessions}. Filed date(s): {filed_dates}. Fiscal period(s): {periods}. Unit(s): {units}.",
                    "DCF Builder Pro",
                )
            elif field == "occupancy":
                sheet[f"{chr(64 + column)}{row}"] = None
            else:
                sheet[f"{chr(64 + column)}{row}"] = None

    for field, row, label in _BALANCE_LINES:
        sheet.cell(row=row, column=1, value=label).font = _BODY_FONT
        for column, item in enumerate(annual, start=2):
            line = _record(item.get(field))
            amount = line.get("value")
            if isinstance(amount, (int, float)) and not isinstance(amount, bool) and math.isfinite(float(amount)):
                cell_ref = f"{chr(64 + column)}{row}"
                _write_value(sheet, cell_ref, float(amount), number_format=_PERCENT_FORMAT if field == "taxRate" else _MONEY_FORMAT)
                _source_line(line, field, int(item["year"]))
                concepts, accessions, filed_dates, periods, units, method = _source_metadata(line)
                sheet[cell_ref].comment = Comment(
                    f"Canonical method: {method}. SEC concept(s): {concepts}. Filing accession(s): {accessions}. Filed date(s): {filed_dates}. Fiscal period(s): {periods}. Unit(s): {units}.",
                    "DCF Builder Pro",
                )

    for column in "BCD":
        _write_value(sheet, f"{column}10", f"=SUM({column}5:{column}9)", formula=True, number_format=_MONEY_FORMAT)
        _write_value(sheet, f"{column}12", f"={column}10-{column}11", formula=True, number_format=_MONEY_FORMAT)
        _write_value(sheet, f"{column}18", f"=SUM({column}11,{column}13:{column}17)", formula=True, number_format=_MONEY_FORMAT)
        _write_value(sheet, f"{column}20", f"={column}18-{column}19", formula=True, number_format=_MONEY_FORMAT)
        _write_value(sheet, f"{column}26", f"=SUM({column}19,{column}21:{column}25)", formula=True, number_format=_MONEY_FORMAT)
        _write_value(sheet, f"{column}28", f"={column}26-{column}27", formula=True, number_format=_MONEY_FORMAT)
        _write_value(sheet, f"{column}31", f"=SUM({column}29:{column}30)", formula=True, number_format=_MONEY_FORMAT)
        _write_value(sheet, f"{column}32", f"={column}27-{column}31", formula=True, number_format=_MONEY_FORMAT)
    for row, label in (
        (10, "Calculated NAREIT FFO"), (12, "NAREIT FFO reconciliation check"),
        (18, "Calculated FFO as modified by Prologis"), (20, "Modified FFO reconciliation check"),
        (26, "Calculated Core FFO"), (28, "Core FFO reconciliation check"),
        (31, "Recurring property-capex proxy"),
    ):
        sheet.cell(row=row, column=1, value=label).font = _BODY_FONT

    if annual[-1].get("reit", {}).get("same_store_noi_growth", {}).get("value") is not None:
        _write_value(sheet, "D46", "=D33/C33-1", formula=True, number_format=_PERCENT_FORMAT)
    sheet["A46"] = "Same-store net-effective NOI growth (FY2025 vs FY2024)"
    sheet["A46"].font = _BODY_FONT
    sheet["D46"].comment = Comment("Calculated from the FY2025 10-K comparative same-store NOI rows, which use the same reported ownership share and population.", "DCF Builder Pro")

    _section(sheet, 48, "Editable forecast, valuation, and NAV assumptions", 2)
    assumption_inputs = (
        ("sameStoreNoiGrowth", 51, "Same-store NOI growth proxy", _PERCENT_FORMAT),
        ("targetOccupancy", 52, "Target operating-portfolio occupancy", _PERCENT_FORMAT),
        ("recurringCapexRatio", 53, "Recurring-capex proxy / Core FFO", _PERCENT_FORMAT),
        ("payoutRatio", 54, "Common-and-preferred distributions / AFFO", _PERCENT_FORMAT),
        ("terminalGrowthRate", 55, "Perpetual AFFO growth", _PERCENT_FORMAT),
        ("navCapRate", 56, "Selected NAV capitalization rate", _PERCENT_FORMAT),
        ("riskFreeRate", 57, "Risk-free rate", _PERCENT_FORMAT),
        ("equityRiskPremium", 58, "Equity risk premium", _PERCENT_FORMAT),
        ("beta", 59, "Beta", "0.00"),
    )
    assumption_sources = _record(assumptions.get("assumptionSources"))
    for field, row, label, number_format in assumption_inputs:
        sheet.cell(row=row, column=1, value=label).font = _BODY_FONT
        missing_same_store_growth = (
            payload.get("buildStatus") == "input_required"
            and field == "sameStoreNoiGrowth"
            and assumptions.get(field) is None
        )
        _write_value(
            sheet,
            f"B{row}",
            None if missing_same_store_growth else _finite_number(assumptions.get(field), field),
            input_cell=not missing_same_store_growth,
            number_format=number_format,
        )
        comment = "Enter the latest filed ratio and source on Input Required." if missing_same_store_growth else str(assumption_sources.get(field) or "Source not provided.")
        sheet[f"B{row}"].comment = Comment(comment, "DCF Builder Pro")
    _write_value(sheet, "B60", "=B57+B58*B59", formula=True, number_format=_PERCENT_FORMAT)
    sheet["A60"] = "Cost of equity (CAPM)"
    sheet["A60"].font = _BODY_FONT
    for row, field in ((57, "riskFreeRateSource"), (58, "equityRiskPremiumSource"), (59, "betaSource")):
        source = str(assumptions.get(field) or "").strip()
        if not source:
            raise ValueError(f"REIT assumption {field} must identify a current source.")
        sheet[f"B{row}"].comment = Comment(source, "DCF Builder Pro")

    for row, field, label, number_format in (
        (61, "currentPrice", "Current share price", _PRICE_FORMAT),
        (62, "dilutedSharesOutstanding", "Diluted shares for valuation", _MONEY_FORMAT),
        (63, "marketCapitalization", "Current market capitalization", _MONEY_FORMAT),
    ):
        sheet.cell(row=row, column=1, value=label).font = _BODY_FONT
        _write_value(sheet, f"B{row}", _finite_number(assumptions.get(field), field), input_cell=True, number_format=number_format)
    sheet["B61"].comment = Comment(f"Current market quote as of {assumptions.get('marketDataAsOfDate')} from {assumptions.get('betaSource')}.", "DCF Builder Pro")
    sheet["B62"].comment = Comment(f"FY{year} filed weighted-average diluted shares; see the historical source table.", "DCF Builder Pro")
    sheet["B63"].comment = Comment(f"Current live market capitalization as of {assumptions.get('marketDataAsOfDate')} from {assumptions.get('betaSource')}.", "DCF Builder Pro")
    sheet["A64"] = "Current market-implied cap rate"
    sheet["A64"].font = _BODY_FONT
    _write_value(sheet, "B64", "=D36/(B63+D42+D43+D44-D41)", formula=True, number_format=_PERCENT_FORMAT)
    sheet["B56"].comment = Comment(
        f"Editable initial assumption. The starting value is market-implied from filed FY{year} Real Estate segment NOI and current market enterprise value; it is a circular cross-check, not an independent market cap-rate estimate. Source: {assumption_sources.get('navCapRate')}.",
        "DCF Builder Pro",
    )

    _section(sheet, 66, "Five-year same-store NOI, Core FFO, AFFO, and distribution forecast", 7)
    for index, column in enumerate("CDEFG", start=1):
        sheet[f"{column}67"] = f"FY{year + index}E"
        sheet[f"{column}67"].font = _HEADER_FONT
        sheet[f"{column}67"].fill = _HEADER_FILL
        sheet[f"{column}67"].alignment = Alignment(horizontal="center")
    forecast_labels = {
        68: "Operating-portfolio occupancy",
        69: "Real Estate segment NOI forecast (same-store growth proxy)",
        70: "Real Estate segment NOI growth",
        71: "Core FFO forecast",
        72: "Recurring-capex proxy",
        73: "Analyst-defined AFFO",
        74: "Common-and-preferred cash distributions",
        75: "AFFO distribution coverage",
        76: "Present value of AFFO",
    }
    for row, label in forecast_labels.items():
        sheet.cell(row=row, column=1, value=label).font = _BODY_FONT
    for index, column in enumerate("CDEFG", start=1):
        previous = chr(ord(column) - 1) if index > 1 else None
        prior_noi = f"{previous}69" if previous else "$D$36"
        prior_occ = f"{previous}68" if previous else "$D$35"
        prior_ffo = f"{previous}71" if previous else "$D$27"
        formulas = {
            68: "=$B$52" if index == 1 else f"={previous}68",
            69: f"={prior_noi}*(1+$B$51)*({column}68/{prior_occ})",
            70: f"=IFERROR({column}69/{prior_noi}-1,0)",
            71: f"={prior_ffo}*(1+$B$51)*({column}68/{prior_occ})",
            72: f"={column}71*$B$53",
            73: f"={column}71-{column}72",
            74: f"={column}73*$B$54",
            75: f"=IFERROR({column}73/{column}74,0)",
            76: f"={column}73/(1+$B$60)^{index}",
        }
        for row, formula in formulas.items():
            number_format = _PERCENT_FORMAT if row in (68, 70) else _RATIO_FORMAT if row == 75 else _MONEY_FORMAT
            _write_value(sheet, f"{column}{row}", formula, formula=True, number_format=number_format)

    _section(sheet, 79, "Common-equity AFFO DCF valuation", 2)
    valuation_rows = (
        (81, "PV of forecast AFFO", "=SUM(C76:G76)", _MONEY_FORMAT),
        (82, "Terminal AFFO value", '=IF($B$60<=$B$55,"",G73*(1+$B$55)/($B$60-$B$55))', _MONEY_FORMAT),
        (83, "PV of terminal AFFO", '=IF(B82="","",B82/(1+$B$60)^5)', _MONEY_FORMAT),
        (84, "Common equity value", '=IF(B83="","",MAX(0,B81+B83))', _MONEY_FORMAT),
        (85, "Implied share price", "=IFERROR(B84/B62,0)", _PRICE_FORMAT),
        (86, "Current share price", "=B61", _PRICE_FORMAT),
        (87, "Implied return", "=IFERROR(B85/B86-1,0)", _PERCENT_FORMAT),
    )
    for row, label, formula, number_format in valuation_rows:
        sheet.cell(row=row, column=1, value=label).font = Font(name="Arial", size=10, bold=row in (84, 85, 87))
        _write_value(sheet, f"B{row}", formula, formula=True, number_format=number_format)
    sheet["A88"] = "Enterprise value"
    sheet["B88"] = "Not applicable — common-equity AFFO model"
    sheet["A88"].font = _BODY_FONT
    sheet["B88"].font = _BODY_FONT

    _section(sheet, 90, "Real Estate segment NAV cross-check", 2)
    nav_rows = (
        (91, "FY reported Real Estate segment NOI", "=D36", _MONEY_FORMAT),
        (92, "Selected cap rate", "=B56", _PERCENT_FORMAT),
        (93, "Implied property value", "=IFERROR(B91/B92,0)", _MONEY_FORMAT),
        (94, "Add cash", "=D41", _MONEY_FORMAT),
        (95, "Less long-term debt", "=-D42", _MONEY_FORMAT),
        (96, "Less preferred equity", "=-D43", _MONEY_FORMAT),
        (97, "Less noncontrolling interests", "=-D44", _MONEY_FORMAT),
        (98, "Real Estate segment NAV to common", "=SUM(B93:B97)", _MONEY_FORMAT),
        (99, "NAV per diluted share", "=IFERROR(B98/B62,0)", _PRICE_FORMAT),
        (100, "Current share price", "=B61", _PRICE_FORMAT),
        (101, "NAV premium / (discount)", "=IFERROR(B99/B100-1,0)", _PERCENT_FORMAT),
    )
    for row, label, formula, number_format in nav_rows:
        sheet.cell(row=row, column=1, value=label).font = Font(name="Arial", size=10, bold=row in (98, 99, 101))
        _write_value(sheet, f"B{row}", formula, formula=True, number_format=number_format)

    _section(sheet, 103, "AFFO equity value sensitivity — cost of equity and terminal growth", 9)
    sheet["B104"] = "Cost of equity / terminal growth"
    sheet["B104"].font = _HEADER_FONT
    sheet["B104"].fill = _HEADER_FILL
    for column, formula in zip("CDEFG", ("=$B$55-2%", "=$B$55-1%", "=$B$55", "=$B$55+1%", "=$B$55+2%")):
        _write_value(sheet, f"{column}104", formula, formula=True, number_format=_PERCENT_FORMAT)
        sheet[f"{column}104"].fill = _HEADER_FILL
        sheet[f"{column}104"].font = _HEADER_FONT
    for row, offset in zip(range(105, 110), ("-2%", "-1%", "", "+1%", "+2%")):
        _write_value(sheet, f"B{row}", f"=$B$60{offset}" if offset else "=$B$60", formula=True, number_format=_PERCENT_FORMAT)
        periods = [f"{column}73/(1+$B{row})^{index}" for index, column in enumerate("CDEFG", start=1)]
        _write_value(sheet, f"H{row}", "=" + "+".join(periods), formula=True, number_format=_MONEY_FORMAT)
        _write_value(sheet, f"I{row}", "=G73", formula=True, number_format=_MONEY_FORMAT)
        for column in "CDEFG":
            formula = (
                f'=IF(OR($B{row}<={column}$104,$B{row}<=0),"",'
                f'MAX(0,$H{row}+($I{row}*(1+{column}$104)/($B{row}-{column}$104))/(1+$B{row})^5))'
            )
            _write_value(sheet, f"{column}{row}", formula, formula=True, number_format=_MONEY_FORMAT)
    sheet["A110"] = "Sensitivity basis"
    sheet["B110"] = "Discounted analyst AFFO plus terminal AFFO at each cost-of-equity / perpetual-growth pair."
    sheet.merge_cells("B110:I110")
    sheet["A110"].font = _BODY_FONT
    sheet["B110"].font = Font(name="Arial", size=9, italic=True, color="404040")
    sheet["B110"].alignment = Alignment(wrap_text=True, vertical="top")

    _section(sheet, 112, "Real Estate segment NAV sensitivity — NOI and capitalization rate", 7)
    sheet["B113"] = "Selected cap rate / NOI factor"
    sheet["B113"].font = _HEADER_FONT
    sheet["B113"].fill = _HEADER_FILL
    for column, formula in zip("CDEFG", ("=90%", "=95%", "=100%", "=105%", "=110%")):
        _write_value(sheet, f"{column}113", formula, formula=True, number_format=_PERCENT_FORMAT)
        sheet[f"{column}113"].fill = _HEADER_FILL
        sheet[f"{column}113"].font = _HEADER_FONT
    for row, offset in zip(range(114, 119), ("-1%", "-0.5%", "", "+0.5%", "+1%")):
        _write_value(sheet, f"B{row}", f"=$B$56{offset}" if offset else "=$B$56", formula=True, number_format=_PERCENT_FORMAT)
        for column in "CDEFG":
            _write_value(sheet, f"{column}{row}", f"={column}$113*$D$36/$B{row}+$D$41-$D$42-$D$43-$D$44", formula=True, number_format=_MONEY_FORMAT)
    sheet["A119"] = "NAV sensitivity scope"
    sheet["B119"] = "Property-only NAV; excludes Strategic Capital and other corporate/non-property assets and liabilities."
    sheet.merge_cells("B119:G119")
    sheet["A119"].font = _BODY_FONT
    sheet["B119"].font = Font(name="Arial", size=9, italic=True, color="404040")
    sheet["B119"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[119].height = 28
    sheet.auto_filter.ref = "A5:D45"


def _map_review_sheet(workbook: Workbook, payload: dict[str, Any], reit_model: dict[str, Any]) -> None:
    review = workbook.create_sheet("Data Review")
    review.sheet_view.showGridLines = False
    company = _record(payload.get("company"))
    ticker = str(company.get("ticker") or "").upper()
    name = str(company.get("name") or ticker)
    review.merge_cells("A1:C1")
    review["A1"] = f"Data Review — {name} ({ticker})"
    review["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    review["A1"].fill = _TITLE_FILL
    review.merge_cells("A2:C2")
    review["A2"] = "NAREIT and issuer Core FFO are filed non-GAAP measures. Analyst AFFO and the NAV cross-check use explicit assumptions."
    review["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    review["A2"].font = Font(name="Arial", size=10, italic=True)
    for column, label in enumerate(("Status", "Area", "Review note"), start=1):
        cell = review.cell(row=5, column=column, value=label)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    ui_meta = _record(payload.get("uiMeta"))
    warnings = ui_meta.get("warnings") if isinstance(ui_meta.get("warnings"), list) else []
    notes = ui_meta.get("sourceNotes") if isinstance(ui_meta.get("sourceNotes"), list) else []
    items = [("REVIEW", "Model limitation", note) for note in warnings if isinstance(note, str)]
    items.extend(("INFO", "SEC source", note) for note in notes if isinstance(note, str))
    items.extend((
        ("INFO", "FFO/AFFO", "Analyst AFFO equals reported Core FFO less tenant improvements, lease commissions, and all property improvements treated as recurring."),
        ("INFO", "NAV", "The selected cap rate begins at the current market-implied cap rate. This NAV cross-check excludes the Strategic Capital segment and other corporate/non-property assets and liabilities."),
        ("INFO", "Distributions", "The reported cash-flow line combines common and preferred dividends; the payout ratio includes both because separate preferred cash dividends are not disclosed in that line."),
        ("INFO", "Shares", "Per-share valuation uses the latest filed weighted-average diluted shares, not period-end shares."),
    ))
    if not items:
        raise ValueError("REIT model export requires source notes and analyst-assumption disclosures on Data Review.")
    for row, (status, area, note) in enumerate(items, start=6):
        cell = review.cell(row=row, column=1, value=status)
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="FCE4D6" if status == "REVIEW" else "DDEBF7")
        review.cell(row=row, column=2, value=area)
        review.cell(row=row, column=3, value=note)
        for column in range(1, 4):
            review.cell(row=row, column=column).alignment = Alignment(wrap_text=True, vertical="top")
        review.row_dimensions[row].height = 34
    review.column_dimensions["A"].width = 14
    review.column_dimensions["B"].width = 22
    review.column_dimensions["C"].width = 115
    review.freeze_panes = "A6"
    review.auto_filter.ref = f"A5:C{5 + len(items)}"


def apply_reit_model(workbook: Workbook, payload: dict[str, Any]) -> None:
    if payload.get("valuationModel") != "reit_affo":
        raise ValueError("REIT workbook requires valuationModel=reit_affo.")
    reit_model = _record(payload.get("reitModel"))
    if not reit_model:
        raise ValueError("REIT workbook requires dedicated REIT history and assumptions.")
    _map_model_sheet(workbook, payload, reit_model)
    _map_review_sheet(workbook, payload, reit_model)


def apply_incomplete_reit_model(
    workbook: Workbook,
    payload: dict[str, Any],
    input_cells: dict[str, dict[str, str]],
) -> None:
    """Link missing filed same-store growth and withhold downstream REIT values."""
    sheet = workbook["REIT Model"]
    requirement, destination = require_input_destination(
        payload, input_cells, "same_store_noi_growth",
        missing_message="Incomplete REIT workbook requires a same_store_noi_growth input.",
        destination_message="Incomplete REIT workbook has no editable same-store NOI growth cell.",
    )
    value_ref = f"'Input Required'!{destination['cell']}"
    status_ref = "'Input Required'!$B$3"
    sheet["B51"] = f'=IF(AND(ISNUMBER({value_ref}),{value_ref}>-1),{value_ref},"")'
    sheet["B51"].font = Font(name="Arial", size=10, color="008000")
    sheet["B51"].fill = _FORMULA_FILL
    sheet["B51"].number_format = _PERCENT_FORMAT
    sheet["B51"].comment = Comment(
        f"Linked to Input Required!{destination['cell']}; enter the latest filed same-store NOI growth and source reference there.",
        "DCF Builder Pro",
    )

    formula_cells = [
        f"{column}{row}"
        for column in "CDEFG"
        for row in range(68, 77)
    ]
    formula_cells.extend(f"B{row}" for row in (*range(81, 86), 87, 64, *range(91, 102)))
    formula_cells.extend(f"{column}{row}" for row in range(105, 110) for column in "CDEFGHI")
    formula_cells.extend(f"{column}{row}" for row in range(114, 119) for column in "CDEFG")
    gate_formulas_on_ready(sheet, status_ref, formula_cells)
