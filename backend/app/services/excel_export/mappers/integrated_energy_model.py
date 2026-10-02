from __future__ import annotations

import math
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

_INPUT_FONT = Font(name="Arial", size=10, color="0000FF")
_FORMULA_FONT = Font(name="Arial", size=10, color="000000")
_BODY_FONT = Font(name="Arial", size=10, color="000000")
_SMALL_FONT = Font(name="Arial", size=9, color="404040")
_HEADER_FONT = Font(name="Arial", size=9, bold=True, color="FFFFFF")
_SECTION_FONT = Font(name="Arial", size=10, bold=True, color="17365D")
_TITLE_FILL = PatternFill("solid", fgColor="17365D")
_SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
_INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
_FORMULA_FILL = PatternFill("solid", fgColor="E2F0D9")
_HEADER_FILL = PatternFill("solid", fgColor="365F91")
_MONEY_FORMAT = '#,##0.0;[Red](#,##0.0);-'
_PERCENT_FORMAT = '0.0%;[Red](0.0%);0.0%'
_PRICE_FORMAT = '$0.00;[Red]($0.00);-'
_PRICE_PER_SHARE_FORMAT = '$#,##0.00;[Red]($#,##0.00);-'
_VOLUME_FORMAT = '#,##0.0;[Red](#,##0.0);-'
_RESERVE_FORMAT = '#,##0.00;[Red](#,##0.00);-'
_MULTIPLE_FORMAT = '0.0x;[Red](0.0x);-'

_ENERGY_FIELDS = (
    'weighted_average_diluted_shares', 'current_debt', 'long_term_debt', 'interest_bearing_debt',
    'crude_oil_production', 'ngl_production', 'bitumen_production', 'synthetic_oil_production',
    'liquids_production', 'natural_gas_production_available_for_sale', 'oil_equivalent_production',
    'average_crude_price', 'average_ngl_price', 'average_bitumen_price', 'average_synthetic_oil_price',
    'average_natural_gas_price', 'average_production_cost_per_oil_equivalent_barrel', 'proved_oil_equivalent_reserves',
    'proved_developed_oil_equivalent_reserves', 'proved_undeveloped_oil_equivalent_reserves', 'upstream_earnings_gaap',
    'energy_products_earnings_gaap', 'chemical_products_earnings_gaap', 'specialty_products_earnings_gaap',
    'corporate_financing_earnings_gaap', 'upstream_depreciation_and_depletion',
    'energy_products_depreciation_and_depletion', 'chemical_products_depreciation_and_depletion',
    'specialty_products_depreciation_and_depletion', 'upstream_ppe_additions_including_noncash',
    'energy_products_ppe_additions_including_noncash', 'chemical_products_ppe_additions_including_noncash',
    'specialty_products_ppe_additions_including_noncash', 'cash_capex', 'operating_working_capital_investment',
    'corporate_interest_revenue', 'brent_2026_earnings_sensitivity', 'henry_hub_2026_earnings_sensitivity',
    'ttf_2026_earnings_sensitivity',
)

_INPUTS = (
    (6, 'Crude price change / barrel', 'crudePriceChange', 'price'),
    (7, 'NGL price change / barrel', 'nglPriceChange', 'price'),
    (8, 'Bitumen price change / barrel', 'bitumenPriceChange', 'price'),
    (9, 'Synthetic oil price change / barrel', 'syntheticOilPriceChange', 'price'),
    (10, 'Natural gas price change / Mcf', 'naturalGasPriceChange', 'price'),
    (11, 'Annual liquids production growth', 'liquidsProductionGrowth', 'percent'),
    (12, 'Annual natural gas production growth', 'naturalGasProductionGrowth', 'percent'),
    (13, 'Annual production cost change / BOE', 'productionCostChangePerBoe', 'price'),
    (14, 'Energy Products earnings growth', 'energyProductsEarningsGrowth', 'percent'),
    (15, 'Chemical Products earnings growth', 'chemicalProductsEarningsGrowth', 'percent'),
    (16, 'Specialty Products earnings growth', 'specialtyProductsEarningsGrowth', 'percent'),
    (17, 'Corporate operating earnings growth', 'corporateOperatingEarningsGrowth', 'percent'),
    (18, 'Consolidated revenue growth', 'revenueGrowth', 'percent'),
    (19, 'Cash CapEx / revenue', 'cashCapexPctRevenue', 'percent'),
    (20, 'Depreciation and depletion / revenue', 'depreciationPctRevenue', 'percent'),
    (21, 'Working-capital investment / revenue', 'workingCapitalInvestmentPctRevenue', 'percent'),
    (22, 'Annual reserve replacement ratio', 'reserveReplacementRatio', 'percent'),
    (24, 'Current share price', 'currentPrice', 'price_per_share'),
    (25, 'Current market capitalization (USD mm)', 'marketCapitalization', 'money_usdmm'),
    (26, 'Current common shares (millions)', 'commonSharesOutstanding', 'shares_millions'),
    (27, 'Cash and equivalents (USD mm)', 'cash', 'money_usdmm'),
    (28, 'Total interest-bearing debt (USD mm)', 'debt', 'money_usdmm'),
    (29, 'Non-operating marketable securities (USD mm)', 'marketableSecurities', 'money_usdmm'),
    (30, 'Noncontrolling interest (USD mm)', 'nonControllingInterest', 'money_usdmm'),
    (31, 'Preferred equity (USD mm)', 'preferredEquity', 'money_usdmm'),
    (33, 'Risk-free rate', 'riskFreeRate', 'percent'),
    (34, 'Equity risk premium', 'equityRiskPremium', 'percent'),
    (35, 'Levered beta', 'beta', 'multiple'),
    (36, 'Current cost of debt', 'costOfDebt', 'percent'),
    (37, 'Effective tax rate', 'taxRate', 'percent'),
    (38, 'Debt weight', 'debtWeight', 'percent'),
    (39, 'Equity weight', 'equityWeight', 'percent'),
    (42, 'Terminal FCFF growth', 'terminalGrowthRate', 'percent'),
)


def _record(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _finite(value: Any, name: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(float(value)):
        raise ValueError(f"Integrated energy workbook requires a finite value for {name}.")
    return float(value)


def _line_value(line: Any, name: str, year: int) -> float:
    record = _record(line)
    if record.get('source') not in {'sec_native', 'derived'}:
        raise ValueError(f"FY{year} XOM energy workbook input {name} is missing or ambiguous.")
    value = _finite(record.get('value'), name)
    sources = record.get('sources')
    if not isinstance(sources, list) or not sources or any(not _record(source).get('accession') or not _record(source).get('filed') for source in sources):
        raise ValueError(f"FY{year} XOM energy workbook input {name} has incomplete SEC provenance.")
    return value


def _write(
    sheet: Worksheet,
    coordinate: str,
    value: Any,
    *,
    formula: bool = False,
    input_cell: bool = False,
    number_format: str | None = None,
    wrap: bool = False,
) -> None:
    cell = sheet[coordinate]
    cell.value = value
    cell.font = _INPUT_FONT if input_cell else _FORMULA_FONT if formula else _BODY_FONT
    cell.fill = _INPUT_FILL if input_cell else _FORMULA_FILL if formula else PatternFill(fill_type=None)
    cell.alignment = Alignment(vertical='center', wrap_text=wrap)
    if number_format:
        cell.number_format = number_format


def _section(sheet: Worksheet, row: int, title: str, end_column: int = 13) -> None:
    for column in range(1, end_column + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = _SECTION_FILL
        cell.font = _SECTION_FONT
    sheet.cell(row=row, column=1, value=title)


def _amount_mm(value: float) -> float:
    return value / 1_000_000


def _validate_payload_model(model: dict[str, Any]) -> tuple[list[dict[str, Any]], list[int], dict[str, Any]]:
    history = _record(model.get('history'))
    annual, years, assumptions = history.get('annual'), history.get('years'), _record(model.get('assumptions'))
    if not isinstance(annual, list) or len(annual) != 3 or not isinstance(years, list) or len(years) != 3:
        raise ValueError('Integrated energy workbook requires three annual source periods.')
    rows = [_record(row) for row in annual]
    if [row.get('year') for row in rows] != years or any(years[index] != years[index - 1] + 1 for index in range(1, len(years))):
        raise ValueError('XOM annual energy history must be aligned and consecutive.')
    if _finite(assumptions.get('forecastYears'), 'forecastYears') != 5 or _finite(assumptions.get('baseYear'), 'baseYear') != years[-1]:
        raise ValueError('Integrated energy workbook requires five forecast years and a matching base year.')
    sources = _record(assumptions.get('assumptionSources'))
    if not sources or any(not isinstance(value, str) or not value.strip() for value in sources.values()):
        raise ValueError('Integrated energy workbook requires source or analyst-input disclosure for every assumption.')
    wacc = _finite(assumptions.get('wacc'), 'wacc')
    growth = _finite(assumptions.get('terminalGrowthRate'), 'terminalGrowthRate')
    if wacc <= 0 or growth < 0 or growth >= wacc:
        raise ValueError('Integrated energy terminal growth must remain below WACC.')
    return rows, years, assumptions


def _source_register_rows(model: dict[str, Any]) -> list[list[Any]]:
    output: list[list[Any]] = []
    history = _record(model.get('history'))
    rows = history.get('annual') if isinstance(history.get('annual'), list) else []
    for item in rows:
        row = _record(item)
        year = row.get('year')
        energy = _record(row.get('energy'))
        candidates = {
            **{f'energy.{key}': value for key, value in energy.items()},
            'revenue': row.get('revenue'), 'net income': row.get('netIncome'),
            'interest expense': row.get('interestExpense'), 'tax rate': row.get('taxRate'),
            'depreciation and depletion': row.get('depreciation'), 'CFO': row.get('cashFlowFromOperations'),
            'cash': row.get('cash'), 'current debt': row.get('currentDebt'), 'long-term debt': row.get('longTermDebt'),
            'total debt': row.get('debt'), 'marketable securities': row.get('marketableSecurities'),
            'preferred equity': row.get('preferredEquity'), 'noncontrolling interest': row.get('nonControllingInterest'),
            'weighted-average diluted shares': row.get('dilutedShares'),
        }
        for field, raw_line in candidates.items():
            line = _record(raw_line)
            sources = line.get('sources') if isinstance(line.get('sources'), list) else []
            if sources:
                for source in sources:
                    source = _record(source)
                    output.append([
                        year, field, line.get('value'), source.get('reported_value'), source.get('unit'), source.get('unit_scale'),
                        line.get('source'), line.get('method'), source.get('concept') or line.get('concept'), source.get('label'),
                        source.get('statement'), source.get('fiscal_period'), source.get('accession'), source.get('filed'), source.get('report_date'),
                    ])
            else:
                output.append([year, field, line.get('value'), None, None, None, line.get('source'), line.get('method'), line.get('concept'), None, None, None, None, None, None])
    for field, source in _record(_record(model.get('assumptions')).get('assumptionSources')).items():
        output.append(['Assumption', field, None, None, None, None, 'analyst_input_or_market', source, None, None, None, None, None, None, None])
    return output


def _map_data_review(workbook: Workbook, model: dict[str, Any]) -> None:
    sheet = workbook.create_sheet('Data Review')
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'A2'
    headers = ['Fiscal year', 'Source/model field', 'Canonical value', 'Reported value', 'Unit', 'Scale', 'Source type', 'Normalization / source note', 'SEC concept', 'Filed label', 'SEC statement/table', 'Fiscal period', 'Accession', 'Filed date', 'Report date']
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for row in _source_register_rows(model):
        sheet.append(row)
    widths = [14, 48, 20, 18, 20, 20, 22, 50, 42, 52, 54, 16, 24, 15, 15]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(ord('A') + index - 1)].width = width
    for row in range(2, sheet.max_row + 1):
        for cell in sheet[row]:
            cell.font = _BODY_FONT
            cell.alignment = Alignment(vertical='top', wrap_text=cell.column >= 8)
        wrap_lines = [math.ceil(len(str(sheet.cell(row=row, column=column).value or '')) / widths[column - 1]) for column in range(8, 12)]
        sheet.row_dimensions[row].height = max(15, min(60, 12 * max(wrap_lines or [1])))
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def _map_assumptions(sheet: Worksheet, assumptions: dict[str, Any]) -> None:
    _write(sheet, 'K4', 'Editable model inputs and current market data')
    sheet['K4'].fill = _SECTION_FILL
    sheet['K4'].font = _SECTION_FONT
    _write(sheet, 'K5', 'Blue font = editable analyst input or current market data', wrap=True)
    sheet['K5'].font = _SMALL_FONT
    sources = _record(assumptions.get('assumptionSources'))
    for row, label, key, kind in _INPUTS:
        value = _finite(assumptions.get(key), key)
        if kind == 'money_usdmm' or kind == 'shares_millions':
            value /= 1_000_000
        number_format = {
            'price': _PRICE_FORMAT,
            'price_per_share': _PRICE_PER_SHARE_FORMAT,
            'percent': _PERCENT_FORMAT,
            'multiple': _MULTIPLE_FORMAT,
            'money_usdmm': _MONEY_FORMAT,
            'shares_millions': _VOLUME_FORMAT,
        }[kind]
        _write(sheet, f'K{row}', label)
        _write(sheet, f'L{row}', value, input_cell=True, number_format=number_format)
        note = sources.get(key)
        if note:
            _write(sheet, f'M{row}', note, wrap=True)
            sheet.row_dimensions[row].height = max(15, 13 * math.ceil(len(note) / 82))
    _write(sheet, 'K23', 'Upstream earnings conversion factor')
    _write(sheet, 'L23', '=D25/D24', formula=True, number_format=_PERCENT_FORMAT)
    _write(sheet, 'M23', sources.get('upstreamEarningsConversionFactor'), wrap=True)
    sheet.row_dimensions[23].height = max(28, 13 * math.ceil(len(str(sources.get('upstreamEarningsConversionFactor') or '')) / 82))
    _write(sheet, 'K40', 'Cost of equity — CAPM')
    _write(sheet, 'L40', '=L33+L35*L34', formula=True, number_format=_PERCENT_FORMAT)
    _write(sheet, 'M40', 'Risk-free rate + beta × equity-risk premium.', wrap=True)
    _write(sheet, 'K41', 'Weighted average cost of capital')
    _write(sheet, 'L41', '=L39*L40+L38*L36*(1-L37)', formula=True, number_format=_PERCENT_FORMAT)
    _write(sheet, 'M41', sources.get('wacc'), wrap=True)
    _write(sheet, 'K43', 'XOM 2026 Brent earnings sensitivity (USD mm / $1/bbl)')
    _write(sheet, 'L43', _finite(assumptions.get('filedBrentSensitivity'), 'filedBrentSensitivity') / 1_000_000, number_format=_MONEY_FORMAT)
    _write(sheet, 'K44', 'XOM 2026 Henry Hub earnings sensitivity (USD mm / $0.10/MMBtu)')
    _write(sheet, 'L44', _finite(assumptions.get('filedHenryHubSensitivity'), 'filedHenryHubSensitivity') / 1_000_000, number_format=_MONEY_FORMAT)
    _write(sheet, 'K45', 'XOM 2026 TTF earnings sensitivity (USD mm / $0.10/MMBtu)')
    _write(sheet, 'L45', _finite(assumptions.get('filedTTFSensitivity'), 'filedTTFSensitivity') / 1_000_000, number_format=_MONEY_FORMAT)


def _map_forecasts(sheet: Worksheet, rows: list[dict[str, Any]], years: list[int]) -> None:
    final_actual_col = chr(ord('B') + len(rows) - 1)
    for period, col in enumerate('EFGHI', start=1):
        previous = chr(ord(col) - 1)
        year = years[-1] + period
        _write(sheet, f'{col}4', f'FY{year}E')
        sheet[f'{col}4'].font = _HEADER_FONT
        sheet[f'{col}4'].fill = _HEADER_FILL
        _write(sheet, f'{col}51', f'={final_actual_col}5*(1+$L$11)^{period}', formula=True, number_format=_VOLUME_FORMAT)
        _write(sheet, f'{col}52', f'={final_actual_col}6*(1+$L$11)^{period}', formula=True, number_format=_VOLUME_FORMAT)
        _write(sheet, f'{col}53', f'={final_actual_col}7*(1+$L$11)^{period}', formula=True, number_format=_VOLUME_FORMAT)
        _write(sheet, f'{col}54', f'={final_actual_col}8*(1+$L$11)^{period}', formula=True, number_format=_VOLUME_FORMAT)
        _write(sheet, f'{col}55', f'=({final_actual_col}9-SUM({final_actual_col}5:{final_actual_col}8))*(1+$L$11)^{period}', formula=True, number_format=_VOLUME_FORMAT)
        _write(sheet, f'{col}56', f'=SUM({col}51:{col}55)', formula=True, number_format=_VOLUME_FORMAT)
        _write(sheet, f'{col}57', f'={final_actual_col}11*(1+$L$12)^{period}', formula=True, number_format=_VOLUME_FORMAT)
        _write(sheet, f'{col}58', f'={col}56+{col}57/6', formula=True, number_format=_VOLUME_FORMAT)
        _write(sheet, f'{col}59', f'={final_actual_col}15+$L$6*{period}', formula=True, number_format=_PRICE_FORMAT)
        _write(sheet, f'{col}60', f'={final_actual_col}16+$L$7*{period}', formula=True, number_format=_PRICE_FORMAT)
        _write(sheet, f'{col}61', f'={final_actual_col}17+$L$8*{period}', formula=True, number_format=_PRICE_FORMAT)
        _write(sheet, f'{col}62', f'={final_actual_col}18+$L$9*{period}', formula=True, number_format=_PRICE_FORMAT)
        _write(sheet, f'{col}63', f'={final_actual_col}19+$L$10*{period}', formula=True, number_format=_PRICE_FORMAT)
        _write(sheet, f'{col}64', f'={final_actual_col}20+$L$13*{period}', formula=True, number_format=_PRICE_FORMAT)
        _write(sheet, f'{col}66', f'={col}51*{col}59*365/1000+{col}52*{col}60*365/1000+{col}53*{col}61*365/1000+{col}54*{col}62*365/1000+{col}55*{col}59*365/1000+{col}57*{col}63*365/1000', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}67', f'={col}58*{col}64*365/1000', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}68', f'={col}66-{col}67', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}69', f'=$D$25+({col}68-$D$24)*$L$23', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}70', f'=$D$28*(1+$L$14)^{period}', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}71', f'=$D$29*(1+$L$15)^{period}', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}72', f'=$D$30*(1+$L$16)^{period}', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}73', f'=$D$35*(1+$L$17)^{period}', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}74', f'=SUM({col}69:{col}73)', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}75', f'=$D$36*(1+$L$18)^{period}', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}76', f'={col}75*$L$20', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}77', f'={col}75*$L$19', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}78', f'={col}75*$L$21', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}79', f'={col}74+{col}76-{col}77-{col}78', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}80', f'=1/(1+$L$41)^{period}', formula=True, number_format=_PERCENT_FORMAT)
        _write(sheet, f'{col}81', f'={col}79*{col}80', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}82', f'={"$D$45" if period == 1 else previous + "85"}', formula=True, number_format=_RESERVE_FORMAT)
        _write(sheet, f'{col}83', f'={col}58*365/1000000', formula=True, number_format=_RESERVE_FORMAT)
        _write(sheet, f'{col}84', f'={col}83*$L$22', formula=True, number_format=_RESERVE_FORMAT)
        _write(sheet, f'{col}85', f'={col}82+{col}84-{col}83', formula=True, number_format=_RESERVE_FORMAT)
        _write(sheet, f'{col}86', f'=IFERROR({col}85/{col}83,0)', formula=True, number_format=_MULTIPLE_FORMAT)
        _write(sheet, f'{col}87', f'={col}85-({col}82+{col}84-{col}83)', formula=True, number_format=_RESERVE_FORMAT)
        _write(sheet, f'{col}88', f'={col}51*365/1000*$L$23', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}89', '=$L$43', formula=True, number_format=_MONEY_FORMAT)


def _map_valuation(sheet: Worksheet) -> None:
    _section(sheet, 90, 'Enterprise-to-common-equity DCF valuation', 10)
    for row, label, formula, number_format in (
        (91, 'PV of five-year FCFF', '=SUM(E81:I81)', _MONEY_FORMAT),
        (92, 'Terminal value (Gordon growth)', '=I79*(1+$L$42)/($L$41-$L$42)', _MONEY_FORMAT),
        (93, 'PV of terminal value', '=B92*I80', _MONEY_FORMAT),
        (94, 'Enterprise value', '=B91+B93', _MONEY_FORMAT),
        (95, 'Common equity value', '=B94+L27+L29-L28-L30-L31', _MONEY_FORMAT),
        (96, 'Current share price', '=L24', _PRICE_PER_SHARE_FORMAT),
        (97, 'Implied value per diluted share', '=B95/L26', _PRICE_PER_SHARE_FORMAT),
        (98, 'Upside / (downside)', '=B97/B96-1', _PERCENT_FORMAT),
    ):
        _write(sheet, f'A{row}', label)
        _write(sheet, f'B{row}', formula, formula=True, number_format=number_format)
    _section(sheet, 101, 'Implied value per share sensitivity — WACC vs. terminal growth', 7)
    for index, column in enumerate('CDEFG', start=-2):
        _write(sheet, f'{column}102', f'=$L$42+{index}*0.005', formula=True, number_format=_PERCENT_FORMAT)
    for row, delta in zip(range(103, 108), (-0.01, -0.005, 0, 0.005, 0.01), strict=True):
        _write(sheet, f'B{row}', f'=$L$41+({delta})', formula=True, number_format=_PERCENT_FORMAT)
        for column in 'CDEFG':
            pv = '+'.join(f'${forecastCol}$79/(1+$B{row})^{period}' for period, forecastCol in enumerate('EFGHI', start=1))
            terminal = f'$I$79*(1+{column}$102)/($B{row}-{column}$102)/(1+$B{row})^5'
            bridge = '+$L$27+$L$29-$L$28-$L$30-$L$31'
            _write(sheet, f'{column}{row}', f'=({pv}+{terminal}{bridge})/$L$26', formula=True, number_format=_PRICE_PER_SHARE_FORMAT)


def apply_integrated_energy_model(workbook: Workbook, payload: dict[str, Any]) -> None:
    model = _record(payload.get('integratedEnergyModel'))
    rows, years, assumptions = _validate_payload_model(model)
    company = _record(payload.get('company'))
    ticker = str(company.get('ticker') or '').upper()
    if ticker != 'XOM':
        raise ValueError('The initial integrated-energy workbook is issuer-specific to XOM.')
    for existing in list(workbook.worksheets):
        workbook.remove(existing)
    sheet = workbook.create_sheet('Integrated Energy Model')
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'E5'
    sheet.sheet_view.zoomScale = 85
    sheet.column_dimensions['A'].width = 58
    for column in 'BCDEFGHI':
        sheet.column_dimensions[column].width = 15
    sheet.column_dimensions['J'].width = 3
    sheet.column_dimensions['K'].width = 47
    sheet.column_dimensions['L'].width = 17
    sheet.column_dimensions['M'].width = 76
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    name = str(company.get('name') or ticker)
    sheet.merge_cells('A1:M1')
    sheet['A1'] = f'{name} ({ticker}) — Integrated Energy FCFF DCF'
    sheet['A1'].font = Font(name='Arial', size=16, bold=True, color='FFFFFF')
    sheet['A1'].fill = _TITLE_FILL
    sheet.row_dimensions[1].height = 28
    sheet.merge_cells('A2:M2')
    sheet['A2'] = 'Filed actuals display in USD millions, thousand barrels per day, million cubic feet per day, and billions of oil-equivalent reserves. Blue cells are editable, formula cells are green, and Data Review retains source periods and SEC accession lineage.'
    sheet['A2'].font = _SMALL_FONT
    sheet['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    sheet.row_dimensions[2].height = 30
    _section(sheet, 3, 'Physical production, realized prices, reserves, and filed segment earnings')
    sheet['A4'] = 'Fiscal year / operating schedule'
    sheet['A4'].font = _HEADER_FONT
    sheet['A4'].fill = _HEADER_FILL
    for row, label in (
        (5, 'Crude oil production (thousand bbl/day)'), (6, 'NGL production (thousand bbl/day)'),
        (7, 'Bitumen production (thousand bbl/day)'), (8, 'Synthetic oil production (thousand bbl/day)'),
        (9, 'Total liquids production (thousand bbl/day)'), (10, 'Liquids component reconciliation check'),
        (11, 'Natural gas production available for sale (MMcf/day)'),
        (12, 'Oil-equivalent production (thousand boe/day)'), (13, 'Oil-equivalent conversion check'),
        (15, 'Crude price (USD/bbl)'), (16, 'NGL price (USD/bbl)'), (17, 'Bitumen price (USD/bbl)'),
        (18, 'Synthetic oil price (USD/bbl)'), (19, 'Natural gas price (USD/Mcf)'), (20, 'Production cost (USD/boe)'),
        (22, 'Modeled production revenue proxy (USD mm)'), (23, 'Production costs (USD mm)'),
        (24, 'Production gross margin proxy (USD mm)'), (25, 'Upstream GAAP segment earnings (USD mm)'),
        (28, 'Energy Products GAAP segment earnings (USD mm)'), (29, 'Chemical Products GAAP segment earnings (USD mm)'),
        (30, 'Specialty Products GAAP segment earnings (USD mm)'), (31, 'Corporate and Financing GAAP earnings (USD mm)'),
        (32, 'Corporate interest revenue (USD mm)'), (33, 'Consolidated interest expense (USD mm)'),
        (34, 'Effective income tax rate'), (35, 'Corporate operating earnings after tax (USD mm)'),
        (36, 'Consolidated sales and other operating revenue (USD mm)'),
        (37, 'Segment-to-net-income reconciliation check (USD mm)'),
        (38, 'Consolidated net income (USD mm)'),
        (39, 'Operating NOPAT after financing adjustment (USD mm)'),
        (40, 'Consolidated depreciation and depletion (USD mm)'), (41, 'Total cash CapEx (Non-GAAP, USD mm)'),
        (42, 'Operating working-capital investment (USD mm)'), (43, 'Historical FCFF (USD mm)'),
        (45, 'Total proved reserves (BBOE)'), (46, 'Proved developed reserves (BBOE)'),
        (47, 'Proved undeveloped reserves (BBOE)'), (48, 'Developed + undeveloped reserve check (BBOE)'),
        (50, 'Forecast year / upstream production and commodity revenue'),
        (51, 'Crude production (thousand bbl/day)'), (52, 'NGL production (thousand bbl/day)'),
        (53, 'Bitumen production (thousand bbl/day)'), (54, 'Synthetic oil production (thousand bbl/day)'),
        (55, 'Other liquid volume residual (thousand bbl/day)'), (56, 'Total liquids production (thousand bbl/day)'),
        (57, 'Natural gas production (MMcf/day)'), (58, 'Oil-equivalent production (thousand boe/day)'),
        (59, 'Crude realized price (USD/bbl)'), (60, 'NGL realized price (USD/bbl)'),
        (61, 'Bitumen realized price (USD/bbl)'), (62, 'Synthetic oil price (USD/bbl)'),
        (63, 'Natural gas realized price (USD/Mcf)'), (64, 'Production cost (USD/boe)'),
        (65, 'Production economics and segment forecasts'), (66, 'Upstream production revenue proxy (USD mm)'),
        (67, 'Upstream production costs (USD mm)'), (68, 'Upstream production gross margin proxy (USD mm)'),
        (69, 'Upstream GAAP earnings forecast (USD mm)'), (70, 'Energy Products earnings (USD mm)'),
        (71, 'Chemical Products earnings (USD mm)'), (72, 'Specialty Products earnings (USD mm)'),
        (73, 'Corporate operating earnings after tax (USD mm)'), (74, 'Unlevered NOPAT (USD mm)'),
        (75, 'Consolidated revenue scale (USD mm)'), (76, 'Depreciation and depletion (USD mm)'),
        (77, 'Cash CapEx (USD mm)'), (78, 'Working-capital investment (USD mm)'),
        (79, 'Unlevered free cash flow (USD mm)'), (80, 'Discount factor'), (81, 'Present value of FCFF (USD mm)'),
        (82, 'Beginning proved reserves (BBOE)'), (83, 'Annual oil-equivalent production (BBOE)'),
        (84, 'Reserve additions under editable replacement ratio (BBOE)'), (85, 'Ending proved reserves (BBOE)'),
        (86, 'Reserve life (years)'), (87, 'Reserve roll-forward check (BBOE)'),
        (88, 'Modeled Brent earnings sensitivity check (USD mm per $1/bbl)'),
        (89, 'Filed XOM FY2026 Brent sensitivity (USD mm per $1/bbl)'),
    ):
        sheet.cell(row=row, column=1, value=label).font = Font(name='Arial', size=10, bold=row in (9,12,24,25,35,37,38,43,48,56,58,68,69,74,79,81,85,88,89))
    for row, title in ((4, 'Filed historical schedule'), (50, 'Forecast schedule — production and volume'), (65, 'Segment earnings, reinvestment, and free cash flow'), (90, 'Enterprise DCF valuation')):
        _section(sheet, row, title)

    for index, item in enumerate(rows):
        col = chr(ord('B') + index)
        year = int(item['year'])
        energy = _record(item.get('energy'))
        for row, field, label, kind in (
            (5, 'crude_oil_production', 'crude oil production', 'volume'),
            (6, 'ngl_production', 'NGL production', 'volume'),
            (7, 'bitumen_production', 'bitumen production', 'volume'),
            (8, 'synthetic_oil_production', 'synthetic oil production', 'volume'),
            (9, 'liquids_production', 'total liquids production', 'volume'),
            (11, 'natural_gas_production_available_for_sale', 'natural gas available for sale', 'gas'),
            (12, 'oil_equivalent_production', 'oil-equivalent production', 'volume'),
            (15, 'average_crude_price', 'crude price', 'price'),
            (16, 'average_ngl_price', 'NGL price', 'price'),
            (17, 'average_bitumen_price', 'bitumen price', 'price'),
            (18, 'average_synthetic_oil_price', 'synthetic oil price', 'price'),
            (19, 'average_natural_gas_price', 'natural gas price', 'price'),
            (20, 'average_production_cost_per_oil_equivalent_barrel', 'production cost per BOE', 'price'),
            (25, 'upstream_earnings_gaap', 'Upstream GAAP earnings', 'money'),
            (28, 'energy_products_earnings_gaap', 'Energy Products earnings', 'money'),
            (29, 'chemical_products_earnings_gaap', 'Chemical Products earnings', 'money'),
            (30, 'specialty_products_earnings_gaap', 'Specialty Products earnings', 'money'),
            (31, 'corporate_financing_earnings_gaap', 'Corporate and Financing earnings', 'money'),
            (32, 'corporate_interest_revenue', 'Corporate interest revenue', 'money'),
            (45, 'proved_oil_equivalent_reserves', 'proved reserves', 'reserve'),
            (46, 'proved_developed_oil_equivalent_reserves', 'proved developed reserves', 'reserve'),
            (47, 'proved_undeveloped_oil_equivalent_reserves', 'proved undeveloped reserves', 'reserve'),
        ):
            crude_line = _record(energy.get('crude_oil_production'))
            crude_sources = crude_line.get('sources')
            crude_is_filed = (
                crude_line.get('source') in {'sec_native', 'derived'}
                and isinstance(crude_line.get('value'), (int, float))
                and not isinstance(crude_line.get('value'), bool)
                and math.isfinite(float(crude_line.get('value')))
                and isinstance(crude_sources, list) and bool(crude_sources)
                and all(_record(source).get('accession') and _record(source).get('filed') for source in crude_sources)
            )
            missing_crude_production = (
                payload.get('buildStatus') == 'input_required'
                and field == 'crude_oil_production'
                and not crude_is_filed
            )
            if missing_crude_production:
                continue
            value = _line_value(energy.get(field), label, year)
            shown = value / 1_000 if kind == 'volume' else value / 1_000_000 if kind == 'gas' else _amount_mm(value) if kind == 'money' else value / 1_000_000_000 if kind == 'reserve' else value
            number_format = _VOLUME_FORMAT if kind in {'volume','gas'} else _MONEY_FORMAT if kind == 'money' else _RESERVE_FORMAT if kind == 'reserve' else _PRICE_FORMAT
            _write(sheet, f'{col}{row}', shown, number_format=number_format)
        _write(sheet, f'{col}10', f'={col}9-SUM({col}5:{col}8)', formula=True, number_format=_VOLUME_FORMAT)
        _write(sheet, f'{col}13', f'={col}12-({col}9+{col}11/6)', formula=True, number_format=_VOLUME_FORMAT)
        _write(sheet, f'{col}22', f'={col}5*{col}15*365/1000+{col}6*{col}16*365/1000+{col}7*{col}17*365/1000+{col}8*{col}18*365/1000+({col}9-SUM({col}5:{col}8))*{col}15*365/1000+{col}11*{col}19*365/1000', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}23', f'={col}12*{col}20*365/1000', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}24', f'={col}22-{col}23', formula=True, number_format=_MONEY_FORMAT)
        corporateInterestExpense = _line_value(item.get('interestExpense'), 'interest expense', year)
        tax_rate = _line_value(item.get('taxRate'), 'effective tax rate', year)
        _write(sheet, f'{col}33', _amount_mm(corporateInterestExpense), number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}34', tax_rate, number_format=_PERCENT_FORMAT)
        _write(sheet, f'{col}35', f'={col}31-({col}32-{col}33)*(1-{col}34)', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}36', _amount_mm(_line_value(item.get('revenue'), 'consolidated revenue', year)), number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}38', _amount_mm(_line_value(item.get('netIncome'), 'consolidated net income', year)), number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}37', f'=SUM({col}25,{col}28:{col}31)-{col}38', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}39', f'=SUM({col}25,{col}28:{col}30,{col}35)', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}40', _amount_mm(_line_value(item.get('depreciation'), 'depreciation and depletion', year)), number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}41', _amount_mm(_line_value(energy.get('cash_capex'), 'cash CapEx', year)), number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}42', _amount_mm(_line_value(energy.get('operating_working_capital_investment'), 'working-capital investment', year)), number_format=_MONEY_FORMAT)
        _write(sheet, f'{col}43', f'={col}39+{col}40-{col}41-{col}42', formula=True, number_format=_MONEY_FORMAT)
        sheet.cell(row=4, column=2 + index, value=f'FY{year}A').font = _HEADER_FONT
        sheet.cell(row=4, column=2 + index).fill = _HEADER_FILL
    _write(sheet, 'L23', '=D25/D24', formula=True, number_format=_PERCENT_FORMAT)
    _map_forecasts(sheet, rows, years)
    _map_assumptions(sheet, assumptions)
    _map_valuation(sheet)
    _map_data_review(workbook, model)
    sheet.freeze_panes = 'E5'
    sheet.print_title_rows = '1:4'
    sheet.print_area = 'A1:M107'
    sheet.auto_filter.ref = 'A4:I58'


def apply_incomplete_integrated_energy_model(
    workbook: Workbook,
    payload: dict[str, Any],
    input_cells: dict[str, dict[str, str]],
) -> None:
    """Link a missing filed crude-production value and withhold dependent energy outputs."""
    sheet = workbook['Integrated Energy Model']
    requirements = payload.get('requiredInputs')
    requirements = requirements if isinstance(requirements, list) else []
    model = _record(payload.get('integratedEnergyModel'))
    history = _record(model.get('history'))
    years = history.get('years') if isinstance(history.get('years'), list) else []
    if not years:
        raise ValueError('Incomplete integrated-energy workbook requires aligned annual history.')
    status_ref = "'Input Required'!$B$3"
    found_requirement = False
    for requirement in requirements:
        if not isinstance(requirement, dict) or requirement.get('key') != 'crude_oil_production':
            continue
        fiscal_year = requirement.get('fiscalYear')
        if not isinstance(fiscal_year, int) or fiscal_year not in years:
            raise ValueError('XOM crude production input must identify one fiscal year.')
        identity = f"crude_oil_production:{fiscal_year}"
        destination = input_cells.get(identity)
        if not destination or destination.get('sheet') != 'Input Required':
            raise ValueError(f'Incomplete integrated-energy workbook has no editable crude production input for FY{fiscal_year}.')
        column = chr(ord('B') + years.index(fiscal_year))
        value_ref = f"'Input Required'!{destination['cell']}"
        sheet[f'{column}5'] = f'=IF(AND({status_ref}="READY",ISNUMBER({value_ref}),{value_ref}>0),{value_ref}/1000,"")'
        sheet[f'{column}5'].font = _FORMULA_FONT
        sheet[f'{column}5'].fill = _FORMULA_FILL
        sheet[f'{column}5'].number_format = _VOLUME_FORMAT
        for row, formula, number_format in (
            (10, f'={column}9-SUM({column}5:{column}8)', _VOLUME_FORMAT),
            (22, f'={column}5*{column}15*365/1000+{column}6*{column}16*365/1000+{column}7*{column}17*365/1000+{column}8*{column}18*365/1000+({column}9-SUM({column}5:{column}8))*{column}15*365/1000+{column}11*{column}19*365/1000', _MONEY_FORMAT),
            (24, f'={column}22-{column}23', _MONEY_FORMAT),
        ):
            sheet[f'{column}{row}'] = f'=IF({status_ref}<>"READY","",{formula[1:]})'
            sheet[f'{column}{row}'].font = _FORMULA_FONT
            sheet[f'{column}{row}'].fill = _FORMULA_FILL
            sheet[f'{column}{row}'].number_format = number_format
        found_requirement = True
    if not found_requirement:
        raise ValueError('Incomplete integrated-energy workbook requires at least one crude_oil_production input.')

    sheet['L23'] = f'=IF({status_ref}<>"READY","",IFERROR(D25/D24,""))'
    sheet['L23'].font = _FORMULA_FONT
    sheet['L23'].fill = _FORMULA_FILL
    sheet['L23'].number_format = _PERCENT_FORMAT
    formula_cells = [
        f'{column}{row}'
        for column in 'EFGHI'
        for row in range(51, 90)
    ]
    formula_cells.extend(f'B{row}' for row in (*range(91, 96), 97, 98))
    formula_cells.extend(f'{column}{row}' for row in range(103, 108) for column in 'CDEFG')
    for cell_ref in formula_cells:
        cell = sheet[cell_ref]
        if isinstance(cell.value, str) and cell.value.startswith('='):
            cell.value = f'=IF({status_ref}<>"READY","",{cell.value[1:]})'
