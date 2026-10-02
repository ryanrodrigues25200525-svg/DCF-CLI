from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

from .review import _map_data_review_sheet
from .utils import _resolve_amount_scale_divisor


_NAVY = "17365D"
_BLUE = "365F91"
_PALE_BLUE = "DDEBF7"
_WHITE = "FFFFFF"
_GREEN = "008000"
_BLACK = "000000"
_THIN = Side(style="thin", color="B7C9D6")
_AMOUNT = "#,##0.0;(#,##0.0);-"
_MULTIPLE = "0.0x"
_PRICE = "$0.00;($0.00);-"
_PERCENT = "0.0%;(0.0%);-"


def _record(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _input_destination(input_cells: dict[str, dict[str, str]], key: str, period: int | str | None = None) -> str:
    identity = f"{key}:{period if period is not None else ''}"
    destination = input_cells.get(identity)
    if destination is None and isinstance(period, str):
        destination = input_cells.get(f"{key}:")
    if not destination or destination.get("sheet") != "Input Required":
        raise ValueError(f"Utility workbook has no editable input cell for {identity}.")
    return destination["cell"]


def _link_input(model, cell: str, key: str, period: int | str | None, divisor: float, input_cells: dict[str, dict[str, str]], number_format: str) -> None:
    destination = _input_destination(input_cells, key, period)
    model[cell] = f'=IF(\'Input Required\'!$B$3<>"READY","",\'Input Required\'!{destination}/{divisor})'
    model[cell].font = Font(name="Arial", size=10, color=_GREEN)
    model[cell].number_format = number_format


def _build_utility_sheet(
    workbook: Workbook,
    payload: dict[str, Any],
    input_cells: dict[str, dict[str, str]] | None,
) -> None:
    valuation_model = payload.get("valuationModel")
    if valuation_model != "utility_dcf":
        raise ValueError("Utility workbook requires valuationModel=utility_dcf.")
    utility = _record(payload.get("utilityModel"))
    incomplete = payload.get("buildStatus") == "input_required"
    market = _record(payload.get("market"))
    if incomplete:
        base_year = utility.get("baseYear")
        forecast_years = utility.get("forecastYears")
        risk_free = utility.get("riskFreeRate")
        erp = utility.get("equityRiskPremium")
        beta = utility.get("beta")
        current_price = utility.get("currentPrice") if utility.get("currentPrice") is not None else market.get("currentPrice")
        shares = utility.get("dilutedShares") if utility.get("dilutedShares") is not None else market.get("sharesDiluted")
        utility_as_of_date = utility.get("asOfDate")
    else:
        assumptions = _record(utility.get("assumptions"))
        base_year = assumptions.get("baseYear")
        forecast_years = 5
        risk_free = assumptions.get("riskFreeRate")
        erp = assumptions.get("equityRiskPremium")
        beta = assumptions.get("beta")
        current_price = assumptions.get("currentPrice")
        shares = assumptions.get("dilutedShares")
        utility_as_of_date = assumptions.get("asOfDate")
    if not isinstance(base_year, int) or forecast_years != 5:
        raise ValueError("Utility workbook requires a fiscal base year and five forecast years.")
    utility_requirements = payload.get("requiredInputs") if isinstance(payload.get("requiredInputs"), list) else []
    utility_requirement_keys = {str(item.get("key")) for item in utility_requirements if isinstance(item, dict)}
    market_fields = (
        ("risk_free_rate", risk_free), ("equity_risk_premium", erp), ("beta", beta),
        ("current_share_price", current_price), ("diluted_shares", shares),
    )
    missing_market_fields = {key for key, value in market_fields if value is None}
    if not incomplete and any(not isinstance(value, (int, float)) for _, value in market_fields):
        raise ValueError("Utility workbook requires live risk-free rate, ERP, beta, price, and shares.")
    if incomplete and not missing_market_fields.issubset(utility_requirement_keys):
        raise ValueError("Utility workbook is missing a required market-input row.")
    if any(isinstance(value, (int, float)) and float(value) <= 0 for _, value in market_fields):
        raise ValueError("Utility market and CAPM inputs must be positive.")

    for existing in list(workbook.worksheets):
        if incomplete and existing.title in {"Input Required", "Data Review"}:
            continue
        workbook.remove(existing)
    model = workbook.create_sheet("Utility Model", 0 if incomplete else None)
    model.sheet_view.showGridLines = False
    ticker = str(_record(payload.get("company")).get("ticker") or "").upper()
    name = str(_record(payload.get("company")).get("name") or ticker)
    divisor = _resolve_amount_scale_divisor(payload)
    status_ref = "'Input Required'!$B$3" if incomplete else None

    model["A1"] = f"Regulated Utility Rate-Base DDM — {ticker}"
    model["A1"].font = Font(name="Arial", size=16, bold=True, color=_WHITE)
    model["A1"].fill = PatternFill("solid", fgColor=_NAVY)
    model.merge_cells("A1:G1")
    model["A2"] = name
    model["A2"].font = Font(name="Arial", size=11, bold=True, color=_NAVY)
    model["A3"] = (
        "Direct common-equity dividend DCF. Rate-base additions, depreciation, authorized returns, and payout inputs are visible below; "
        "no enterprise value or corporate FCFF is calculated."
    )
    model.merge_cells("A3:G3")
    model["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    model.row_dimensions[3].height = 34
    if incomplete:
        model["A4"] = "Model status"
        model["B4"] = f"={status_ref}"
        model["B4"].font = Font(name="Arial", size=10, bold=True, color=_GREEN)
    model["A6"] = f"FY{base_year} regulated rate base"
    if incomplete:
        assert input_cells is not None
        _link_input(model, "B6", "jurisdictional_rate_base", base_year, divisor, input_cells, _AMOUNT)
    else:
        model["B6"] = assumptions["baseRateBase"] / divisor
        model["B6"].number_format = _AMOUNT
    model["B6"].fill = PatternFill("solid", fgColor=_PALE_BLUE)

    model["A8"] = "Rate-base roll-forward"
    model["A8"].font = Font(name="Arial", size=11, bold=True, color=_NAVY)
    years = [base_year + offset for offset in range(1, 6)]
    for index, year in enumerate(years, start=3):
        cell = model.cell(row=7, column=index, value=year)
        cell.number_format = '"FY"0'
        cell.font = Font(name="Arial", size=10, bold=True, color=_WHITE)
        cell.fill = PatternFill("solid", fgColor=_BLUE)
        cell.alignment = Alignment(horizontal="center")

    for row, label in (
        (9, "Opening regulated rate base"),
        (10, "Approved additions entering rate base"),
        (11, "Rate-base depreciation and retirements"),
        (12, "Closing regulated rate base"),
        (13, "Average regulated rate base"),
        (15, "Authorized common-equity ratio"),
        (16, "Allowed return on common equity"),
        (17, "Allowed equity earnings proxy"),
        (18, "Common dividend payout ratio"),
        (19, "Forecast common dividends"),
        (20, "Dividend discount factor"),
        (21, "Present value of common dividends"),
    ):
        model.cell(row=row, column=1, value=label)

    for index, year in enumerate(years, start=3):
        column = get_column_letter(index)
        prior_column = get_column_letter(index - 1)
        period = index - 2
        if index == 3:
            model[f"{column}9"] = "=$B$6"
        else:
            model[f"{column}9"] = f"={prior_column}12"
        if incomplete:
            assert input_cells is not None
            _link_input(model, f"{column}10", "rate_base_additions", year, divisor, input_cells, _AMOUNT)
            _link_input(model, f"{column}11", "rate_base_depreciation", year, divisor, input_cells, _AMOUNT)
        else:
            model[f"{column}10"] = assumptions["rateBaseAdditions"][period - 1] / divisor
            model[f"{column}11"] = assumptions["rateBaseDepreciation"][period - 1] / divisor
            model[f"{column}10"].number_format = _AMOUNT
            model[f"{column}11"].number_format = _AMOUNT
        model[f"{column}12"] = f'=IF({status_ref}<>"READY","",{column}9+{column}10-{column}11)' if incomplete else f"={column}9+{column}10-{column}11"
        model[f"{column}13"] = f'=IF({status_ref}<>"READY","",AVERAGE({column}9,{column}12))' if incomplete else f"=AVERAGE({column}9,{column}12)"
        if incomplete:
            _link_input(model, f"{column}15", "authorized_equity_ratio", base_year, 1, input_cells, _PERCENT)
            _link_input(model, f"{column}16", "allowed_roe", base_year, 1, input_cells, _PERCENT)
            _link_input(model, f"{column}18", "dividend_payout_ratio", base_year, 1, input_cells, _PERCENT)
        else:
            model[f"{column}15"] = assumptions["authorizedEquityRatio"]
            model[f"{column}16"] = assumptions["allowedRoe"]
            model[f"{column}18"] = assumptions["dividendPayoutRatio"]
            for row in (15, 16, 18):
                model.cell(row=row, column=index).number_format = _PERCENT
        model[f"{column}17"] = f'=IF({status_ref}<>"READY","",{column}13*{column}15*{column}16)' if incomplete else f"={column}13*{column}15*{column}16"
        model[f"{column}19"] = f'=IF({status_ref}<>"READY","",{column}17*{column}18)' if incomplete else f"={column}17*{column}18"
        model[f"{column}20"] = f'=IF({status_ref}<>"READY","",1/(1+$B$24)^{period})' if incomplete else f"=1/(1+$B$24)^{period}"
        model[f"{column}21"] = f'=IF({status_ref}<>"READY","",{column}19*{column}20)' if incomplete else f"={column}19*{column}20"
        for row in (9, 10, 11, 12, 13, 17, 19, 21):
            model.cell(row=row, column=index).number_format = _AMOUNT
        for row in (15, 16, 18):
            model.cell(row=row, column=index).number_format = _PERCENT
        model[f"{column}20"].number_format = "0.000x"
        for row in (9, 12, 13, 17, 19, 21):
            model.cell(row=row, column=index).font = Font(name="Arial", size=10, color=_BLACK if not incomplete else _GREEN)

    model["A23"] = "Common-equity dividend DCF"
    model["A23"].font = Font(name="Arial", size=11, bold=True, color=_NAVY)
    for row, label in (
        (24, "Cost of equity (CAPM)"),
        (25, "Terminal dividend growth"),
        (26, "Terminal value"),
        (27, "Present value of terminal value"),
        (28, "Present value of forecast dividends"),
        (29, "Common-equity value"),
        (30, "Diluted shares (millions)"),
        (31, "Implied value per share"),
        (32, "Current share price"),
        (33, "Implied upside / (downside)"),
    ):
        model.cell(row=row, column=1, value=label)
    model["I4"] = "Current Market Inputs"
    model["I4"].font = Font(name="Arial", size=10, bold=True, color=_WHITE)
    model["I4"].fill = PatternFill("solid", fgColor=_BLUE)
    model["J4"] = "Value"
    model["J4"].font = Font(name="Arial", size=10, bold=True, color=_WHITE)
    model["J4"].fill = PatternFill("solid", fgColor=_BLUE)
    model["K4"] = "Source / as-of date"
    model["K4"].font = Font(name="Arial", size=10, bold=True, color=_WHITE)
    model["K4"].fill = PatternFill("solid", fgColor=_BLUE)
    market_sources = _record(utility.get("assumptionSources")) if incomplete else _record(_record(utility.get("assumptions")).get("assumptionSources"))
    for row, label, value, number_format, source_key, source_name in (
        (5, "Current risk-free rate", risk_free, _PERCENT, "risk_free_rate", "riskFreeRate"),
        (6, "Current equity-risk premium", erp, _PERCENT, "equity_risk_premium", "equityRiskPremium"),
        (7, "Current utility beta", beta, _MULTIPLE, "beta", "beta"),
    ):
        model.cell(row=row, column=9, value=label)
        if incomplete and source_key in missing_market_fields:
            assert input_cells is not None
            _link_input(model, f"J{row}", source_key, str(utility_as_of_date), 1, input_cells, number_format)
        else:
            model.cell(row=row, column=10, value=float(value))
            model.cell(row=row, column=10).number_format = number_format
            model.cell(row=row, column=10).font = Font(name="Arial", size=10, color=_GREEN)
        model.cell(row=row, column=11, value=market_sources.get(source_name, ""))
        model.cell(row=row, column=11).alignment = Alignment(wrap_text=True, vertical="top")
    coe_formula = "=J5+J7*J6"
    model["B24"] = f'=IF({status_ref}<>"READY","",{coe_formula[1:]})' if incomplete else coe_formula
    if incomplete:
        assert input_cells is not None
        _link_input(model, "B25", "terminal_growth_rate", None, 1, input_cells, _PERCENT)
    else:
        model["B25"] = assumptions["terminalGrowthRate"]
    model["B24"].number_format = _PERCENT
    model["B25"].number_format = _PERCENT
    status = f'{status_ref}<>"READY"' if incomplete else "FALSE"
    model["B26"] = f'=IF({status},"",IF(B24<=B25,"",G19*(1+B25)/(B24-B25)))'
    model["B27"] = f'=IF({status},"",IF(B24<=B25,"",B26/(1+B24)^5))'
    model["B28"] = f'=IF({status},"",SUM(C21:G21))'
    model["B29"] = f'=IF({status},"",IF(B24<=B25,"",B27+B28))'
    if incomplete and "diluted_shares" in missing_market_fields:
        assert input_cells is not None
        _link_input(model, "B30", "diluted_shares", base_year, divisor, input_cells, _AMOUNT)
    else:
        model["B30"] = float(shares) / divisor
    model["B31"] = f'=IF({status},"",IF(B24<=B25,"",IFERROR(B29/B30,0)))'
    if incomplete and "current_share_price" in missing_market_fields:
        assert input_cells is not None
        _link_input(model, "B32", "current_share_price", str(utility_as_of_date), 1, input_cells, _PRICE)
    else:
        model["B32"] = float(current_price)
    model["K8"] = market_sources.get("currentPrice", "")
    model["K9"] = market_sources.get("dilutedShares", "")
    model["B33"] = f'=IF({status},"",IF(B24<=B25,"",IFERROR(B31/B32-1,0)))'
    model["B26"].number_format = _AMOUNT
    model["B27"].number_format = _AMOUNT
    model["B28"].number_format = _AMOUNT
    model["B29"].number_format = _AMOUNT
    model["B30"].number_format = _AMOUNT
    model["B31"].number_format = _PRICE
    model["B32"].number_format = _PRICE
    model["B33"].number_format = _PERCENT

    model["A35"] = "Per-share sensitivity — cost of equity vs. terminal growth"
    model["A35"].font = Font(name="Arial", size=11, bold=True, color=_NAVY)
    growth_offsets = (-0.02, -0.01, 0, 0.01, 0.02)
    coe_offsets = (-0.02, -0.01, 0, 0.01, 0.02)
    model["B36"] = "Cost of equity / terminal growth"
    for index, offset in enumerate(growth_offsets, start=3):
        cell = model.cell(row=36, column=index, value=offset)
        cell.number_format = _PERCENT
        cell.font = Font(name="Arial", size=9, bold=True, color=_WHITE)
        cell.fill = PatternFill("solid", fgColor=_BLUE)
    for index, offset in enumerate(coe_offsets, start=37):
        model.cell(row=index, column=2, value=offset).number_format = _PERCENT
        model.cell(row=index, column=2).font = Font(name="Arial", size=9, bold=True, color=_WHITE)
        model.cell(row=index, column=2).fill = PatternFill("solid", fgColor=_BLUE)
        pv = "+".join(
            f"{get_column_letter(forecast_column)}19/(1+$B$24+$B{index})^{period}"
            for period, forecast_column in enumerate(range(3, 8), start=1)
        )
        for column in range(3, 8):
            header = get_column_letter(column)
            terminal_growth = f"($B$25+{header}$36)"
            sensitivity = (
                f'=IF({status},"",IF($B$24+$B{index}<={terminal_growth},"",'
                f'IFERROR(({pv}+($G$19*(1+{terminal_growth})/($B$24+$B{index}-{terminal_growth}))'
                f'/(1+$B$24+$B{index})^5)/$B$30,"")))'
            )
            model.cell(row=index, column=column, value=sensitivity).number_format = _PRICE

    for row in (12, 17, 19, 21, 29, 31, 33):
        model.cell(row=row, column=1).font = Font(name="Arial", size=10, bold=True)
        model.cell(row=row, column=2).font = Font(name="Arial", size=10, bold=True)
        model.cell(row=row, column=1).fill = PatternFill("solid", fgColor=_PALE_BLUE)
        model.cell(row=row, column=2).fill = PatternFill("solid", fgColor=_PALE_BLUE)
    for row in range(6, 34):
        for column in (1, 2):
            model.cell(row=row, column=column).border = Border(bottom=_THIN)
            model.cell(row=row, column=column).alignment = Alignment(vertical="center", wrap_text=True)
    for column, width in {"A": 44, "B": 19, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "I": 32, "J": 15, "K": 52}.items():
        model.column_dimensions[column].width = width
    model.freeze_panes = "C8"

    if incomplete:
        if input_cells is None:
            raise ValueError("Incomplete utility workbook requires an input register.")
    else:
        _map_data_review_sheet(workbook, payload)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"


def apply_utility_model(workbook: Workbook, payload: dict[str, Any]) -> None:
    _build_utility_sheet(workbook, payload, None)


def apply_incomplete_utility_model(
    workbook: Workbook,
    payload: dict[str, Any],
    input_cells: dict[str, dict[str, str]],
) -> None:
    _build_utility_sheet(workbook, payload, input_cells)
