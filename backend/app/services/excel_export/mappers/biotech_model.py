from __future__ import annotations

import re
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

from .review import _map_data_review_sheet
from .utils import _resolve_amount_scale_divisor


_NAVY, _BLUE, _PALE_BLUE, _WHITE, _GREEN, _BLACK = "17365D", "365F91", "DDEBF7", "FFFFFF", "008000", "000000"
_THIN = Side(style="thin", color="B7C9D6")
_AMOUNT = "#,##0.0;(#,##0.0);-"
_PERCENT = "0.0%;(0.0%);-"
_PRICE = "$0.00;($0.00);-"


def _record(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _rnpv_assets(raw_assets: Any) -> list[dict[str, Any]]:
    if not isinstance(raw_assets, list):
        return []
    stage_pattern = re.compile(r"(?:phase\s*(?:[2-3]|1\s*/\s*[2-3])|registrational|regulatory)", re.IGNORECASE)
    return [asset for asset in raw_assets if isinstance(asset, dict)
        and asset.get("development_status") != "explicitly_paused"
        and isinstance(asset.get("stage"), str) and stage_pattern.search(asset["stage"])]


def _other_pipeline_assets(raw_assets: Any) -> list[dict[str, Any]]:
    if not isinstance(raw_assets, list):
        return []
    modeled_ids = {str(asset.get("asset_id")) for asset in _rnpv_assets(raw_assets) if isinstance(asset, dict)}
    return [asset for asset in raw_assets if isinstance(asset, dict) and str(asset.get("asset_id")) not in modeled_ids]


def _destination(input_cells: dict[str, dict[str, str]], key: str, year: int | None = None) -> str:
    identity = f"{key}:{year if year is not None else ''}"
    destination = input_cells.get(identity)
    if not destination or destination.get("sheet") != "Input Required":
        raise ValueError(f"Biotech workbook has no editable input cell for {identity}.")
    return destination["cell"]


def _link_input(sheet, cell: str, input_cells: dict[str, dict[str, str]], key: str, year: int | None, divisor: float, number_format: str) -> None:
    destination = _destination(input_cells, key, year)
    sheet[cell] = f'=IF(\'Input Required\'!$B$3<>"READY","",\'Input Required\'!{destination}/{divisor})'
    sheet[cell].font = Font(name="Arial", size=9, color=_GREEN)
    sheet[cell].number_format = number_format


def apply_biotech_model(workbook: Workbook, payload: dict[str, Any]) -> None:
    _build_biotech_model(workbook, payload, None)


def apply_incomplete_biotech_model(workbook: Workbook, payload: dict[str, Any], input_cells: dict[str, dict[str, str]]) -> None:
    _build_biotech_model(workbook, payload, input_cells)


def _build_biotech_model(
    workbook: Workbook,
    payload: dict[str, Any],
    input_cells: dict[str, dict[str, str]] | None,
) -> None:
    if payload.get("valuationModel") != "biotech_pipeline_rnpv":
        raise ValueError("Biotech workbook requires valuationModel=biotech_pipeline_rnpv.")
    incomplete = payload.get("buildStatus") == "input_required"
    biotech = _record(payload.get("biotechModel"))
    if incomplete:
        base_year = biotech.get("baseYear")
        forecast_years = biotech.get("forecastYears")
        revenue_base = biotech.get("commercialRevenueBase")
        pipeline_assets = biotech.get("pipelineAssets")
        risk_free = biotech.get("riskFreeRate")
        erp = biotech.get("equityRiskPremium")
        beta = biotech.get("beta")
        assumption_sources = _record(biotech.get("assumptionSources"))
        assumptions = {}
    else:
        assumptions = _record(biotech.get("assumptions"))
        base_year = assumptions.get("baseYear")
        forecast_years = 10
        revenue_base = assumptions.get("commercialRevenueBase")
        pipeline_assets = biotech.get("pipelineAssets")
        risk_free = assumptions.get("riskFreeRate")
        erp = assumptions.get("equityRiskPremium")
        beta = assumptions.get("beta")
        assumption_sources = _record(assumptions.get("assumptionSources"))
    if not isinstance(base_year, int) or forecast_years != 10:
        raise ValueError("Biotech workbook requires a filed fiscal base year and ten forecast years.")
    if not isinstance(revenue_base, (int, float)) or revenue_base < 0:
        raise ValueError("Biotech workbook requires non-negative filed consolidated revenue (zero values a pipeline-only franchise).")
    if any(not isinstance(value, (int, float)) or value <= 0 for value in (risk_free, erp, beta)):
        raise ValueError("Biotech workbook requires a current risk-free rate, ERP, and beta.")
    assets = _rnpv_assets(pipeline_assets)
    other_assets = _other_pipeline_assets(pipeline_assets)
    if not assets:
        raise ValueError("Biotech workbook requires disclosed source-backed late-stage pipeline assets.")
    if incomplete and input_cells is None:
        raise ValueError("Incomplete biotech workbook requires an input register.")

    market = _record(payload.get("market"))
    divisor = _resolve_amount_scale_divisor(payload)
    for existing in list(workbook.worksheets):
        if incomplete and existing.title in {"Input Required", "Data Review"}:
            continue
        workbook.remove(existing)
    summary = workbook.create_sheet("Biotech Model", 0 if incomplete else None)
    pipeline = workbook.create_sheet("Pipeline Valuation", 1 if incomplete else None)
    summary.sheet_view.showGridLines = False
    pipeline.sheet_view.showGridLines = False
    ticker = str(_record(payload.get("company")).get("ticker") or "").upper()
    name = str(_record(payload.get("company")).get("name") or ticker)
    status_ref = "'Biotech Model'!$B$4" if incomplete else '"READY"'
    input_status_ref = "'Input Required'!$B$3" if incomplete else '"READY"'

    summary["A1"] = f"Commercial DCF and Pipeline rNPV — {ticker}"
    summary["A1"].font = Font(name="Arial", size=16, bold=True, color=_WHITE)
    summary["A1"].fill = PatternFill("solid", fgColor=_NAVY)
    summary.merge_cells("A1:O1")
    summary["A2"] = name
    summary["A2"].font = Font(name="Arial", size=11, bold=True, color=_NAVY)
    summary["A3"] = "Commercial revenue starts from reported consolidated revenue (which can include collaboration and milestone revenue). Pipeline assumptions are blank and analyst-owned."
    summary.merge_cells("A3:O3")
    summary["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    summary.row_dimensions[3].height = 34
    amount_unit = {1.0: "USD actual", 1_000.0: "USD thousands", 1_000_000.0: "USD millions", 1_000_000_000.0: "USD billions"}.get(divisor, "USD scaled")
    share_unit = {1.0: "actual", 1_000.0: "thousands", 1_000_000.0: "millions", 1_000_000_000.0: "billions"}.get(divisor, "scaled")
    summary["A5"] = f"Units: {amount_unit}; diluted shares in {share_unit}; per-share value in USD."
    summary.merge_cells("A5:F5")
    summary["A5"].font = Font(name="Arial", size=9, italic=True, color=_BLACK)
    summary["A4"] = "Model status"
    summary["B4"] = (
        f'=IF({input_status_ref}<>"READY",{input_status_ref},IF(OR($M$15<2%,$M$15>50%),'
        '"INPUT ERROR — WACC outside supported range",IF(OR($M$16>10%,$M$16>$M$15-0.5%),'
        '"INPUT ERROR — terminal growth must be at least 0.5% below WACC","READY")))'
        if incomplete else "READY"
    )
    summary["B4"].font = Font(name="Arial", size=10, bold=True, color=_GREEN)
    summary["A6"] = f"FY{base_year} reported consolidated revenue base"
    summary["B6"] = float(revenue_base) / divisor
    summary["B6"].number_format = _AMOUNT

    summary["A8"] = "Commercial franchise DCF"
    summary["A8"].font = Font(name="Arial", size=11, bold=True, color=_NAVY)
    forecast_columns = [get_column_letter(index + 2) for index in range(10)]
    for index, column in enumerate(forecast_columns):
        year = base_year + index + 1
        summary.cell(row=9, column=index + 2, value=year)
        summary.cell(row=9, column=index + 2).number_format = '"FY"0'
        summary.cell(row=9, column=index + 2).font = Font(name="Arial", size=9, bold=True, color=_WHITE)
        summary.cell(row=9, column=index + 2).fill = PatternFill("solid", fgColor=_BLUE)
        if incomplete:
            assert input_cells is not None
            _link_input(summary, f"{column}10", input_cells, "commercial_revenue_growth", year, 1, _PERCENT)
            _link_input(summary, f"{column}12", input_cells, "commercial_fcf_margin", None, 1, _PERCENT)
        else:
            summary[f"{column}10"] = assumptions["commercialRevenueGrowth"][index]
            summary[f"{column}10"].number_format = _PERCENT
            summary[f"{column}12"] = assumptions["commercialFcfMargin"]
            summary[f"{column}12"].number_format = _PERCENT
        prior_revenue = "$B$6" if index == 0 else f"{forecast_columns[index - 1]}11"
        summary[f"{column}11"] = f'=IF({status_ref}<>"READY","",{prior_revenue}*(1+{column}10))'
        summary[f"{column}13"] = f'=IF({status_ref}<>"READY","",{column}11*{column}12)'
        summary[f"{column}14"] = f'=IF({status_ref}<>"READY","",1/(1+$M$15)^{index + 1})'
        summary[f"{column}15"] = f'=IF({status_ref}<>"READY","",{column}13*{column}14)'
        for row in (11, 13, 15):
            summary[f"{column}{row}"].number_format = _AMOUNT
            summary[f"{column}{row}"].font = Font(name="Arial", size=9, color=_BLACK)
        summary[f"{column}14"].number_format = "0.000x"
    for row, label in (
        (10, "Commercial revenue growth"), (11, "Commercial revenue"), (12, "Post-tax commercial FCFF margin"),
        (13, "Commercial FCFF"), (14, "Discount factor"), (15, "PV of commercial FCFF"),
    ):
        summary.cell(row=row, column=1, value=label)

    summary["L4"] = "WACC and market inputs"
    summary["L4"].font = Font(name="Arial", size=10, bold=True, color=_WHITE)
    summary["L4"].fill = PatternFill("solid", fgColor=_BLUE)
    summary["M4"] = "Value"
    summary["M4"].font = Font(name="Arial", size=9, bold=True, color=_WHITE)
    summary["M4"].fill = PatternFill("solid", fgColor=_BLUE)
    summary["N4"] = "Source / as-of date"
    summary["N4"].font = Font(name="Arial", size=9, bold=True, color=_WHITE)
    summary["N4"].fill = PatternFill("solid", fgColor=_BLUE)
    for row, label, value, number_format, field in (
        (5, "Risk-free rate", risk_free, _PERCENT, "riskFreeRate"),
        (6, "Equity-risk premium", erp, _PERCENT, "equityRiskPremium"),
        (7, "Beta", beta, "0.000x", "beta"),
    ):
        summary.cell(row=row, column=12, value=label)
        summary.cell(row=row, column=13, value=float(value))
        summary.cell(row=row, column=13).number_format = number_format
        summary.cell(row=row, column=13).font = Font(name="Arial", size=9, color=_GREEN)
        summary.cell(row=row, column=14, value=assumption_sources.get(field, ""))
        summary.cell(row=row, column=14).alignment = Alignment(wrap_text=True, vertical="top")
    summary["L8"] = "Cost of equity"
    summary["M8"] = "=M5+M7*M6"
    summary["M8"].number_format = _PERCENT
    summary["L9"] = "Cost of debt"
    summary["L10"] = "Normalized tax rate"
    if incomplete:
        assert input_cells is not None
        _link_input(summary, "M9", input_cells, "biotech_cost_of_debt", None, 1, _PERCENT)
        _link_input(summary, "M10", input_cells, "biotech_normalized_tax_rate", None, 1, _PERCENT)
    else:
        summary["M9"] = assumptions["costOfDebt"]
        summary["M10"] = assumptions["normalizedTaxRate"]
        summary["M9"].number_format = _PERCENT
        summary["M10"].number_format = _PERCENT
    summary["L11"] = "Market capitalization"
    summary["M11"] = float(market.get("marketCap") or 0) / divisor
    summary["M11"].number_format = _AMOUNT
    summary["L12"] = "Interest-bearing debt"
    summary["M12"] = float(market.get("debt") or 0) / divisor
    summary["M12"].number_format = _AMOUNT
    summary["L13"] = "Equity weight"
    summary["M13"] = "=IFERROR(M11/(M11+M12),1)"
    summary["M13"].number_format = _PERCENT
    summary["L14"] = "Debt weight"
    summary["M14"] = "=IFERROR(M12/(M11+M12),0)"
    summary["M14"].number_format = _PERCENT
    summary["L15"] = "WACC"
    summary["M15"] = f'=IF({input_status_ref}<>"READY","",M8*M13+M9*(1-M10)*M14)' if incomplete else "=M8*M13+M9*(1-M10)*M14"
    summary["M15"].number_format = _PERCENT
    summary["L16"] = "Terminal growth"
    if incomplete:
        assert input_cells is not None
        _link_input(summary, "M16", input_cells, "terminal_growth_rate", None, 1, _PERCENT)
    else:
        summary["M16"] = assumptions["terminalGrowthRate"]
        summary["M16"].number_format = _PERCENT

    summary["L18"] = "Commercial franchise terminal value"
    summary["M18"] = f'=IF({status_ref}<>"READY","",IF($M$15<=$M$16,"",$K$13*(1+$M$16)/($M$15-$M$16)))'
    summary["L19"] = "PV of commercial terminal value"
    summary["M19"] = f'=IF({status_ref}<>"READY","",IF($M$15<=$M$16,"",$M$18/(1+$M$15)^10))'
    summary["L20"] = "PV of commercial forecast"
    summary["M20"] = f'=IF({status_ref}<>"READY","",SUM($B$15:$K$15))'
    summary["L21"] = "Commercial franchise enterprise value"
    summary["M21"] = f'=IF({status_ref}<>"READY","",IF($M$15<=$M$16,"",$M$19+$M$20))'
    summary["L22"] = "Individually forecast late-stage asset rNPV"
    summary["M22"] = "='Pipeline Valuation'!$B$3"
    summary["L23"] = "Other SEC-listed pipeline asset rNPV"
    summary["M23"] = "='Pipeline Valuation'!$B$4"
    summary["L24"] = "Pipeline rNPV outside the SEC-mapped asset inventory"
    if incomplete:
        assert input_cells is not None
        _link_input(summary, "M24", input_cells, "unmapped_pipeline_rnpv", None, divisor, _AMOUNT)
    else:
        summary["M24"] = assumptions["otherPipelineRnpv"] / divisor
        summary["M24"].number_format = _AMOUNT
    summary["L25"] = "Enterprise value"
    summary["M25"] = f'=IF({status_ref}<>"READY","",IF($M$15<=$M$16,"",SUM(M21:M24)))'
    for row, label, market_key in (
        (26, "Add cash and cash equivalents", "cash"),
        (27, "Add marketable securities / nonoperating assets", "nonOperatingAssets"),
        (28, "Less interest-bearing debt", "debt"),
        (29, "Less noncontrolling interest", "minorityInterest"),
        (30, "Less preferred equity", "preferredEquity"),
    ):
        summary.cell(row=row, column=12, value=label)
        summary.cell(row=row, column=13, value=float(market.get(market_key) or 0) / divisor)
        summary.cell(row=row, column=13).number_format = _AMOUNT
    summary["L31"] = "Common-equity value"
    summary["M31"] = f'=IF({status_ref}<>"READY","",IF($M$15<=$M$16,"",M25+M26+M27-M28-M29-M30))'
    summary["M31"].number_format = _AMOUNT
    summary["L32"] = f"Diluted shares ({share_unit})"
    summary["M32"] = float(market.get("sharesDiluted") or 0) / divisor
    summary["M32"].number_format = _AMOUNT
    summary["L33"] = "Implied value per share"
    summary["M33"] = f'=IF({status_ref}<>"READY","",IF($M$15<=$M$16,"",IFERROR(M31/M32,0)))'
    summary["M33"].number_format = _PRICE
    summary["L34"] = "Current share price"
    summary["M34"] = float(market.get("currentPrice") or 0)
    summary["M34"].number_format = _PRICE
    summary["L35"] = "Implied upside / (downside)"
    summary["M35"] = f'=IF({status_ref}<>"READY","",IF($M$15<=$M$16,"",IFERROR(M33/M34-1,0)))'
    summary["M35"].number_format = _PERCENT

    pipeline.sheet_view.showGridLines = False
    pipeline_forecast_years = 35
    pipeline_forecast_end = get_column_letter(16 + pipeline_forecast_years - 1)
    pipeline["A1"] = f"Risk-adjusted Pipeline Valuation — {ticker}"
    pipeline["A1"].font = Font(name="Arial", size=16, bold=True, color=_WHITE)
    pipeline["A1"].fill = PatternFill("solid", fgColor=_NAVY)
    pipeline.merge_cells(f"A1:{pipeline_forecast_end}1")
    pipeline["A2"] = "SEC identity, stage, status, partner, and filing narrative are reported facts. The 35-year schedule captures launch and post-LOE cash flows; sales, probability, economics, and residual rNPV remain analyst inputs."
    pipeline.merge_cells(f"A2:{pipeline_forecast_end}2")
    pipeline["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    pipeline["A3"] = "Total individually forecast late-stage asset rNPV"
    pipeline["B3"] = f'=IF({status_ref}<>"READY","",SUM(O6:O{5 + len(assets)}))'
    pipeline["B3"].number_format = _AMOUNT
    pipeline["B3"].font = Font(name="Arial", size=10, bold=True, color=_GREEN)
    pipeline["O3"] = "Discount factor"
    pipeline["O4"] = "Risk-adjusted asset FCFF"
    for cell in (pipeline["O3"], pipeline["O4"]):
        cell.font = Font(name="Arial", size=9, bold=True, color=_NAVY)
    pipeline["A4"] = "Other SEC-listed pipeline asset rNPV"
    other_section_row = 7 + len(assets) + 1
    other_header_row = other_section_row + 1
    other_first_row = other_header_row + 1
    other_last_row = other_first_row + len(other_assets) - 1
    if other_assets:
        pipeline["B4"] = f'=IF({status_ref}<>"READY","",SUM(G{other_first_row}:G{other_last_row}))'
    else:
        pipeline["B4"] = f'=IF({status_ref}<>"READY","",0)'
    pipeline["B4"].number_format = _AMOUNT
    pipeline["B4"].font = Font(name="Arial", size=10, bold=True, color=_GREEN)
    header_labels = (
        "Asset ID", "SEC asset label", "Stage / status / partner", "SEC source and filing basis", "Include",
        "Launch year", "Peak sales", "Years to peak", "Exclusivity year", "Post-LOE erosion", "PoS",
        "Retained share", "Post-tax FCFF margin", "Development-cost PV", "Asset rNPV",
    )
    for column, label in enumerate(header_labels, start=1):
        cell = pipeline.cell(row=5, column=column, value=label)
        cell.font = Font(name="Arial", size=9, bold=True, color=_WHITE)
        cell.fill = PatternFill("solid", fgColor=_BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for index in range(pipeline_forecast_years):
        column = 16 + index
        cell = pipeline.cell(row=5, column=column, value=base_year + index + 1)
        cell.number_format = '"FY"0'
        cell.font = Font(name="Arial", size=9, bold=True, color=_WHITE)
        cell.fill = PatternFill("solid", fgColor=_BLUE)
        discount = pipeline.cell(row=3, column=column, value=f'=IF(\'Biotech Model\'!$B$4<>"READY","",1/(1+\'Biotech Model\'!$M$15)^{index + 1})')
        discount.number_format = "0.000x"
        total_cashflow = pipeline.cell(row=4, column=column, value=f'=IF(\'Biotech Model\'!$B$4<>"READY","",SUM({get_column_letter(column)}6:{get_column_letter(column)}{5 + len(assets)}))')
        total_cashflow.number_format = _AMOUNT

    for row, asset in enumerate(assets, start=6):
        asset_id = str(asset.get("asset_id") or "")
        pipeline.cell(row=row, column=1, value=asset_id)
        pipeline.cell(row=row, column=2, value=str(asset.get("asset_name") or asset_id))
        stage_status_partner = [str(item) for item in (asset.get("stage"), asset.get("development_status")) if item]
        if asset.get("partner"):
            stage_status_partner.append(f"Partner: {asset['partner']}")
        pipeline.cell(row=row, column=3, value="; ".join(stage_status_partner))
        pipeline.cell(row=row, column=4, value=f"SEC accession {asset.get('accession_number')}; filed {asset.get('filing_date')}; report date {asset.get('report_date')}; {asset.get('form')}; {asset.get('source_statement')}")
        pipeline.cell(row=row, column=4).alignment = Alignment(wrap_text=True, vertical="top")
        asset_inputs = (
            (5, "include", 1, "0"),
            (6, "launch_year", 1, "0"),
            (7, "peak_sales", divisor, _AMOUNT),
            (8, "years_to_peak", 1, "0"),
            (9, "exclusivity_year", 1, "0"),
            (10, "post_loe_erosion", 1, _PERCENT),
            (11, "probability_of_success", 1, _PERCENT),
            (12, "retained_share", 1, _PERCENT),
            (13, "contribution_margin", 1, _PERCENT),
            (14, "development_cost_pv", divisor, _AMOUNT),
        )
        for column, suffix, scale, number_format in asset_inputs:
            key = f"asset_{suffix}:{asset_id}"
            if incomplete:
                assert input_cells is not None
                _link_input(pipeline, f"{get_column_letter(column)}{row}", input_cells, key, None, scale, number_format)
            else:
                value_key = {
                    "include": "include", "launch_year": "launchYear", "peak_sales": "peakSales",
                    "years_to_peak": "yearsToPeak", "exclusivity_year": "exclusivityYear",
                    "post_loe_erosion": "postLoeErosion", "probability_of_success": "probabilityOfSuccess",
                    "retained_share": "retainedShare", "contribution_margin": "contributionMargin",
                    "development_cost_pv": "developmentCostPv",
                }[suffix]
                raw_asset = next((item for item in assumptions["assets"] if item["assetId"] == asset_id), None)
                if raw_asset is None:
                    raise ValueError(f"Complete biotech export is missing model assumptions for {asset_id}.")
                pipeline.cell(row=row, column=column, value=raw_asset[value_key] / scale)
                pipeline.cell(row=row, column=column).number_format = number_format
        for index in range(pipeline_forecast_years):
            column = get_column_letter(16 + index)
            year = base_year + index + 1
            include_ref = f"E{row}"
            launch_ref = f"F{row}"
            peak_ref = f"G{row}"
            ramp_ref = f"H{row}"
            exclusivity_ref = f"I{row}"
            erosion_ref = f"J{row}"
            probability_ref = f"K{row}"
            retained_ref = f"L{row}"
            margin_ref = f"M{row}"
            sales_formula = (
                f'IF({column}$5<{launch_ref},0,{peak_ref}*MIN(1,({column}$5-{launch_ref}+1)/{ramp_ref})'
                f'*IF({column}$5<={exclusivity_ref},1,(1-{erosion_ref})^({column}$5-{exclusivity_ref})))'
            )
            pipeline[f"{column}{row}"] = (
                f'=IF(\'Biotech Model\'!$B$4<>"READY","",IF({include_ref}=0,0,'
                f'{sales_formula}*{probability_ref}*{retained_ref}*{margin_ref}))'
            )
            pipeline[f"{column}{row}"].number_format = _AMOUNT
        fcff_range = f"P{row}:{pipeline_forecast_end}{row}"
        pipeline[f"O{row}"] = f'=IF(\'Biotech Model\'!$B$4<>"READY","",SUMPRODUCT({fcff_range},$P$3:${pipeline_forecast_end}$3)-E{row}*N{row})'
        pipeline[f"O{row}"].number_format = _AMOUNT
        pipeline[f"O{row}"].font = Font(name="Arial", size=9, bold=True, color=_GREEN)

    summary["M22"] = "='Pipeline Valuation'!$B$3"
    summary["M22"].number_format = _AMOUNT
    summary["M23"] = "='Pipeline Valuation'!$B$4"
    summary["M23"].number_format = _AMOUNT

    pipeline["A" + str(other_section_row)] = "Other SEC-listed assets — source-mapped rNPV inputs"
    pipeline.merge_cells(start_row=other_section_row, start_column=1, end_row=other_section_row, end_column=7)
    pipeline.cell(row=other_section_row, column=1).font = Font(name="Arial", size=11, bold=True, color=_WHITE)
    pipeline.cell(row=other_section_row, column=1).fill = PatternFill("solid", fgColor=_NAVY)
    other_headers = ("Asset ID", "SEC asset label", "Stage / status / partner", "SEC source and filing basis", "Include", "Analyst-entered rNPV (PV)", "Included rNPV")
    for column, label in enumerate(other_headers, start=1):
        cell = pipeline.cell(row=other_header_row, column=column, value=label)
        cell.font = Font(name="Arial", size=9, bold=True, color=_WHITE)
        cell.fill = PatternFill("solid", fgColor=_BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row, asset in enumerate(other_assets, start=other_first_row):
        asset_id = str(asset.get("asset_id") or "")
        pipeline.cell(row=row, column=1, value=asset_id)
        pipeline.cell(row=row, column=2, value=str(asset.get("asset_name") or asset_id))
        stage_status_partner = [str(item) for item in (asset.get("stage"), asset.get("development_status")) if item]
        if asset.get("partner"):
            stage_status_partner.append(f"Partner: {asset['partner']}")
        pipeline.cell(row=row, column=3, value="; ".join(stage_status_partner))
        pipeline.cell(row=row, column=4, value=f"SEC accession {asset.get('accession_number')}; filed {asset.get('filing_date')}; report date {asset.get('report_date')}; {asset.get('form')}; {asset.get('source_statement')}")
        pipeline.cell(row=row, column=4).alignment = Alignment(wrap_text=True, vertical="top")
        if incomplete:
            assert input_cells is not None
            _link_input(pipeline, f"E{row}", input_cells, f"asset_scope_include:{asset_id}", None, 1, "0")
            _link_input(pipeline, f"F{row}", input_cells, f"asset_scope_rnpv:{asset_id}", None, divisor, _AMOUNT)
        else:
            raw_asset = next((item for item in assumptions.get("otherAssets", []) if item["assetId"] == asset_id), None)
            if raw_asset is None:
                raise ValueError(f"Complete biotech export is missing other-pipeline rNPV for {asset_id}.")
            pipeline.cell(row=row, column=5, value=raw_asset["include"])
            pipeline.cell(row=row, column=6, value=raw_asset["rnpv"] / divisor)
            pipeline.cell(row=row, column=6).number_format = _AMOUNT
        pipeline.cell(row=row, column=7, value=f'=IF(\'Biotech Model\'!$B$4<>"READY","",E{row}*F{row})')
        pipeline.cell(row=row, column=7).number_format = _AMOUNT

    if incomplete:
        assert input_cells is not None
        unmapped_destination = _destination(input_cells, "unmapped_pipeline_rnpv")
        summary["M24"] = f'=IF(\'Input Required\'!$B$3<>"READY","",\'Input Required\'!{unmapped_destination}/{divisor})'
        summary["M24"].font = Font(name="Arial", size=9, color=_GREEN)
        summary["M24"].number_format = _AMOUNT
    summary["A37"] = "Formula guide: Commercial revenue = prior year × (1 + growth). Late-stage asset FCFF = sales × PoS × retained share × post-tax margin. Asset rNPV = discounted FCFF − included development-cost PV. Early, paused, and other SEC-listed programs use explicit include controls and analyst-entered PVs; those PVs stay fixed in sensitivity. Unmapped pipeline is separate."
    summary.merge_cells("A37:O37")
    summary["A37"].alignment = Alignment(wrap_text=True, vertical="top")
    summary["A37"].font = Font(name="Arial", size=9, color=_BLACK)
    summary.row_dimensions[37].height = 38

    summary["A39"] = "Sensitivity: implied value per share"
    summary["A39"].font = Font(name="Arial", size=11, bold=True, color=_NAVY)
    summary["A40"] = "WACC / terminal growth"
    summary["A40"].font = Font(name="Arial", size=9, bold=True, color=_WHITE)
    summary["A40"].fill = PatternFill("solid", fgColor=_BLUE)
    terminal_offsets = (-0.01, -0.005, 0.0, 0.005, 0.01)
    wacc_offsets = (-0.02, -0.01, 0.0, 0.01, 0.02)
    for index, offset in enumerate(terminal_offsets, start=2):
        cell = summary.cell(row=40, column=index, value=f'=IF({status_ref}<>"READY","",MAX(0,$M$16+({offset})))')
        cell.number_format = _PERCENT
        cell.font = Font(name="Arial", size=9, bold=True, color=_WHITE)
        cell.fill = PatternFill("solid", fgColor=_BLUE)
    source_costs = f"SUMPRODUCT('Pipeline Valuation'!$E$6:$E${5 + len(assets)},'Pipeline Valuation'!$N$6:$N${5 + len(assets)})"
    for row_offset, offset in enumerate(wacc_offsets):
        row = 41 + row_offset
        commercial_terms = [f"${get_column_letter(index + 2)}$13/(1+$A${row})^{index + 1}" for index in range(10)]
        pipeline_terms = [f"'Pipeline Valuation'!${get_column_letter(index + 16)}$4/(1+$A${row})^{index + 1}" for index in range(pipeline_forecast_years)]
        wacc_cell = summary.cell(row=row, column=1, value=f'=IF({status_ref}<>"READY","",$M$15+({offset}))')
        wacc_cell.number_format = _PERCENT
        wacc_cell.font = Font(name="Arial", size=9, bold=True, color=_WHITE)
        wacc_cell.fill = PatternFill("solid", fgColor=_BLUE)
        for column in range(2, 7):
            growth_ref = f"{get_column_letter(column)}$40"
            terminal_value = f"$K$13*(1+{growth_ref})/($A{row}-{growth_ref})/(1+$A{row})^10"
            equity_value = (
                f"({'+'.join(commercial_terms)}+{terminal_value}+{'+'.join(pipeline_terms)}-"
                f"{source_costs}+'Pipeline Valuation'!$B$4+$M$24+$M$26+$M$27-$M$28-$M$29-$M$30)/$M$32"
            )
            cell = summary.cell(row=row, column=column, value=f'=IF({status_ref}<>"READY","",IF(OR($A{row}<2%,$A{row}>50%,{growth_ref}>10%,$A{row}<={growth_ref}+0.5%),"",{equity_value}))')
            cell.number_format = _PRICE
            cell.font = Font(name="Arial", size=9, color=_BLACK)
            cell.border = Border(bottom=Side(style="thin", color="D0D7DE"))

    for row, width in ((1, 25), (2, 40), (3, 15), (4, 15), (5, 15)):
        summary.row_dimensions[row].height = 24 if row in (1, 3) else summary.row_dimensions[row].height
        pipeline.row_dimensions[row].height = 30 if row == 2 else pipeline.row_dimensions[row].height
    for column, width in {
        "A": 28, "B": 28, "C": 36, "D": 75, "E": 10, "F": 17, "G": 18, "H": 12,
        "I": 13, "J": 16, "K": 12, "L": 14, "M": 20, "N": 20, "O": 18,
        **{get_column_letter(16 + index): 15 for index in range(pipeline_forecast_years)},
    }.items():
        pipeline.column_dimensions[column].width = width
    pipeline.freeze_panes = "P6"
    pipeline.auto_filter.ref = f"A5:{pipeline_forecast_end}{5 + len(assets)}"
    pipeline.row_dimensions[5].height = 44
    for row in range(6, 6 + len(assets)):
        pipeline.row_dimensions[row].height = 70
    if other_assets:
        pipeline.row_dimensions[other_header_row].height = 38
        for row in range(other_first_row, other_last_row + 1):
            pipeline.row_dimensions[row].height = 70
    for column, width in {"A": 36, "B": 16, "C": 13, "D": 15, "E": 15, "F": 15, "G": 15, "H": 15, "I": 22, "J": 22, "K": 22, "L": 22, "M": 16}.items():
        summary.column_dimensions[column].width = width
    summary.freeze_panes = "B9"
    for row in range(5, 36):
        for column in range(1, 14):
            cell = summary.cell(row=row, column=column)
            if cell.value is not None:
                cell.border = Border(bottom=Side(style="thin", color="D0D7DE"))
                if column == 12:
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in range(6, 6 + len(assets)):
        for column in range(1, 16 + pipeline_forecast_years):
            cell = pipeline.cell(row=row, column=column)
            cell.alignment = Alignment(vertical="center", wrap_text=column in (2, 3, 4))
            if column in (1, 2, 3, 4):
                cell.border = Border(bottom=_THIN)
    for row in range(other_first_row, other_last_row + 1):
        for column in range(1, 8):
            cell = pipeline.cell(row=row, column=column)
            cell.alignment = Alignment(vertical="center", wrap_text=column in (2, 3, 4))
            cell.border = Border(bottom=_THIN)

    if not incomplete:
        _map_data_review_sheet(workbook, payload)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
