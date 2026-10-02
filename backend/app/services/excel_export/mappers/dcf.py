from __future__ import annotations

import math
from copy import copy
from datetime import date, datetime
from typing import Any

from openpyxl.cell.cell import MergedCell
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .utils import (
    DCF_TIMELINE_COLUMNS,
    RECALC_COLUMNS,
    SHEET_ASSUMPTION_BREAKDOWN,
    SHEET_COVER,
    SHEET_DATA_RECALCULATED,
    SHEET_DATA_ORIGINAL,
    SHEET_DCF_BASE,
    SHEET_DCF_BEAR,
    SHEET_DCF_BULL,
    SHEET_OUTPUTS,
    _apply_scenario_snapshot_to_sheet,
    _axis_from_bounds,
    _clear_sensitivity_blocks,
    _enforce_core_public_dcf_formulas,
    _enforce_outputs_bridge_formulas,
    _fallback_revenue_ebit_matrix,
    _fallback_wacc_terminal_matrix,
    _fiscal_year_end_date,
    _force_set,
    _historical_value_for_year,
    _infer_revenue_growth_rate,
    _link_dcf_income_statement_to_recalculated_data,
    _matrix_values,
    _normalize_public_dcf_assumption_block,
    _numeric_list,
    _safe_date,
    _safe_set,
    _safe_set_or_clear,
    _sanitize_terminal_growth_rate,
    _sanitize_wacc_rate,
    _scale,
    _scenario_snapshot,
    _sheet,
    _set_comment,
    _set_percent_axis_column,
    _set_percent_axis_row,
    _to_float,
)

SCENARIO_BASE = "base"
SCENARIO_BULL = "bull"
SCENARIO_BEAR = "bear"

def _required_positive_integer(value: Any, label: str, default: int) -> int:
    parsed = _to_float(value)
    if parsed is None:
        parsed = float(default)
    if parsed < 1 or not parsed.is_integer():
        raise ValueError(f"{label} must be a positive whole number.")
    return int(parsed)


def _map_dcf_base_inputs(dcf_base: Worksheet, payload: dict[str, Any], divisor: float) -> None:
    assumptions = payload.get("assumptions", {})
    assumptions = assumptions if isinstance(assumptions, dict) else {}
    market = payload.get("market", {})
    market = market if isinstance(market, dict) else {}
    transaction = payload.get("transaction", {})
    transaction = transaction if isinstance(transaction, dict) else {}
    company = payload.get("company", {})
    company = company if isinstance(company, dict) else {}
    forecasts = payload.get("forecasts", []) or []
    key_metrics = (payload.get("uiMeta") or {}).get("keyMetrics") or {}

    tax_rate = _to_float(assumptions.get("taxRate"))
    da_pct = _to_float(assumptions.get("daPctRevenue"))
    if da_pct is None:
        da_pct = _to_float(assumptions.get("deaRatio"))
    ebit_margin = _to_float(assumptions.get("ebitMargin"))
    if ebit_margin is None:
        ebit_margin = _to_float(assumptions.get("ebitMarginTarget"))
    steady_state_ebit_margin = _to_float(assumptions.get("ebitMarginSteadyState"))
    if steady_state_ebit_margin is None:
        steady_state_ebit_margin = ebit_margin
    gross_margin = _to_float(assumptions.get("grossMargin"))
    historicals = payload.get("historicals", {})
    historicals = historicals if isinstance(historicals, dict) else {}
    historical_years = [year for year in (historicals.get("years", []) or []) if _to_float(year) is not None]
    latest_year = int(historical_years[-1]) if historical_years else None
    if latest_year is not None:
        revenue = _historical_value_for_year(payload, latest_year, statement="income", keys=["Total Revenue", "Revenue"])
        if gross_margin is None:
            gross_profit = _historical_value_for_year(payload, latest_year, statement="income", keys=["Gross Profit", "GrossProfit"])
            gross_margin = gross_profit / revenue if gross_profit is not None and revenue else None
        if ebit_margin is None:
            ebit = _historical_value_for_year(payload, latest_year, statement="income", keys=["Operating Income (EBIT)", "EBIT"])
            ebit_margin = ebit / revenue if ebit is not None and revenue else None
    terminal_assumptions = assumptions.get("terminal") or {}
    exit_multiple = _to_float(terminal_assumptions.get("exitMultiple"))
    wacc_assumptions = assumptions.get("wacc")
    wacc_assumptions = wacc_assumptions if isinstance(wacc_assumptions, dict) else {}
    base_wacc_assumption = _sanitize_wacc_rate(assumptions.get("waccRate") or wacc_assumptions.get("waccRate"))
    terminal_growth = _sanitize_terminal_growth_rate(
        terminal_assumptions.get("g"),
        reference_wacc=base_wacc_assumption,
    )
    revenue_growth = _to_float(assumptions.get("revenueGrowthStage1"))
    if revenue_growth is None:
        revenue_growth = _to_float(assumptions.get("revenueGrowth"))
    if revenue_growth is None:
        revenue_growth = _to_float(assumptions.get("revenueGrowthRate"))
    if revenue_growth is None:
        revenue_growth = _infer_revenue_growth_rate(forecasts)
    if revenue_growth is None:
        revenue_growth = 0.06

    shares = _to_float(market.get("sharesDiluted")) or 0.0
    price = _to_float(market.get("currentPrice")) or 0.0
    if shares <= 0 or price <= 0:
        raise ValueError("The DCF workbook requires positive diluted shares and a current share price for per-share outputs.")
    market_cap_raw = _to_float(market.get("marketCap"))
    if market_cap_raw is None and shares > 0 and price > 0:
        market_cap_raw = shares * price

    debt_raw = _to_float(market.get("marketValueDebt"))
    if debt_raw is None:
        debt_raw = _to_float(market.get("debt"))
    cash_raw = _to_float(market.get("cash"))
    non_operating_assets_raw = _to_float(market.get("nonOperatingAssets"))
    minority_interest_raw = _to_float(market.get("minorityInterest"))
    preferred_equity_raw = _to_float(market.get("preferredEquity"))
    for label, value in (
        ("debt", debt_raw),
        ("cash", cash_raw),
        ("non-operating assets", non_operating_assets_raw),
        ("minority interest", minority_interest_raw),
        ("preferred equity", preferred_equity_raw),
    ):
        if value is None:
            raise ValueError(f"A sourced {label} amount is required for the equity bridge.")

    equity_raw = market_cap_raw

    net_debt_raw = _to_float(market.get("netDebt"))
    if net_debt_raw is None:
        net_debt_raw = debt_raw - cash_raw

    explicit_purchase_price = None
    if isinstance(transaction, dict):
        explicit_purchase_price = _to_float(transaction.get("purchasePrice"))
    if explicit_purchase_price is None:
        explicit_purchase_price = _to_float(payload.get("purchasePrice"))

    enterprise_raw = explicit_purchase_price if explicit_purchase_price is not None else _to_float(key_metrics.get("enterpriseValue"))
    if enterprise_raw is None and equity_raw is not None and net_debt_raw is not None:
        enterprise_raw = equity_raw + net_debt_raw

    _safe_set(dcf_base, "C9", _scale(enterprise_raw, divisor))
    _safe_set(dcf_base, "C11", _scale(equity_raw, divisor))
    _safe_set(dcf_base, "F17", _scale(cash_raw, divisor))
    _safe_set(dcf_base, "F18", _scale(debt_raw, divisor))
    _safe_set(dcf_base, "F19", _scale(non_operating_assets_raw, divisor))
    _safe_set(dcf_base, "E20", "Minority Interest")
    dcf_base["F20"]._style = copy(dcf_base["F9"]._style)
    dcf_base["F20"].number_format = dcf_base["F9"].number_format
    _force_set(dcf_base, "F20", _scale(minority_interest_raw, divisor))
    _safe_set(dcf_base, "E21", "Preferred Equity")
    dcf_base["F21"]._style = copy(dcf_base["F9"]._style)
    dcf_base["F21"].number_format = dcf_base["F9"].number_format
    _force_set(dcf_base, "F21", _scale(preferred_equity_raw, divisor))

    first_forecast = forecasts[0] if forecasts and isinstance(forecasts[0], dict) else {}
    capex_ratio = _to_float(assumptions.get("capexPctRevenue"))
    if capex_ratio is None:
        capex_ratio = _to_float(assumptions.get("capexRatio"))
    if capex_ratio is None:
        first_capex = _to_float(first_forecast.get("capex"))
        first_revenue = _to_float(first_forecast.get("revenue"))
        capex_ratio = abs(first_capex / first_revenue) if first_capex is not None and first_revenue else None
    nwc_ratio = _to_float(assumptions.get("nwcPctRevenue"))
    if capex_ratio is None or not math.isfinite(capex_ratio) or capex_ratio < 0:
        raise ValueError("The operating DCF requires an editable, non-negative capex-to-revenue assumption.")
    if nwc_ratio is None:
        nwc_ratio = 0.0
    if ebit_margin is None or not math.isfinite(ebit_margin) or not 0 < ebit_margin <= 1:
        raise ValueError("The operating DCF requires an editable EBIT margin assumption between 0% and 100%.")
    if gross_margin is None or not math.isfinite(gross_margin) or not 0 < gross_margin <= 1:
        raise ValueError("The operating DCF requires an editable gross margin assumption between 0% and 100%.")

    _safe_set(dcf_base, "E9", "CapEx % of Revenue")
    _safe_set(dcf_base, "F9", capex_ratio)
    dcf_base["F9"].number_format = "0.0%"
    operating_archetype = str(company.get("operatingArchetype") or company.get("operating_archetype") or "")
    if operating_archetype == "subscription_software":
        _safe_set(dcf_base, "E10", "Aggregate operating NWC residual % of Revenue")
    else:
        _safe_set(dcf_base, "E10", "NWC intensity overlay % of Revenue")
    _safe_set(dcf_base, "F10", nwc_ratio)
    dcf_base["F10"].number_format = "0.0%"
    _safe_set(dcf_base, "F11", tax_rate)
    _safe_set(dcf_base, "F13", da_pct)
    _safe_set(dcf_base, "F14", revenue_growth)
    _safe_set(dcf_base, "E14", "Stage 1 Revenue Growth")
    _safe_set(dcf_base, "E15", "Starting EBIT Margin")
    _safe_set(dcf_base, "F15", ebit_margin)
    _safe_set(dcf_base, "E16", "Forecast Gross Margin")
    _safe_set(dcf_base, "F16", gross_margin)
    _safe_set(dcf_base, "E17", "Steady-State EBIT Margin")
    _safe_set(dcf_base, "F17", steady_state_ebit_margin)
    _safe_set(dcf_base, "E19", "Stage 2 Revenue Growth")
    stage2_growth = _to_float(assumptions.get("revenueGrowthStage2"))
    if stage2_growth is None:
        stage2_growth = revenue_growth
    _safe_set(dcf_base, "F19", stage2_growth)
    dcf_base["F15"].number_format = "0.0%"
    dcf_base["F16"].number_format = "0.0%"
    dcf_base["F14"].number_format = "0.0%"
    dcf_base["F17"]._style = copy(dcf_base["F15"]._style)
    dcf_base["F17"].number_format = "0.0%"
    dcf_base["F19"]._style = copy(dcf_base["F14"]._style)
    dcf_base["F19"].number_format = "0.0%"

    margin_convergence_years = _required_positive_integer(
        assumptions.get("ebitMarginConvergenceYears") or assumptions.get("marginRampYears"),
        "EBIT-margin convergence period",
        5,
    )
    growth_stage1_years = _required_positive_integer(
        assumptions.get("revenueGrowthStage1Years"), "Revenue-growth Stage 1 duration", 3
    )
    growth_fade_years = _required_positive_integer(
        assumptions.get("revenueGrowthFadeYears"), "Revenue-growth fade duration", 4
    )
    for label_ref, value_ref, label, value, number_format in (
        ("H13", "I13", "EBIT Margin Convergence (Years)", margin_convergence_years, "0"),
        ("H14", "I14", "Revenue Growth Stage 1 (Years)", growth_stage1_years, "0"),
        ("H15", "I15", "Revenue Growth Fade (Years)", growth_fade_years, "0"),
        ("H16", "I16", "Diluted Shares (millions)", _scale(shares, divisor), "#,##0.0;(#,##0.0);-"),
        ("H17", "I17", "Current Share Price", price, "$0.00"),
    ):
        _safe_set(dcf_base, label_ref, label)
        _safe_set(dcf_base, value_ref, value)
        dcf_base[label_ref]._style = copy(dcf_base["H9"]._style)
        dcf_base[value_ref]._style = copy(dcf_base["F14"]._style)
        dcf_base[value_ref].number_format = number_format

    _safe_set(dcf_base, "C16", exit_multiple)
    _safe_set(dcf_base, "Q103", terminal_growth)

    _set_comment(dcf_base, "F9", "Editable capex assumption: capex as a percentage of revenue; forecast capex recalculates as revenue multiplied by this input.")
    _set_comment(
        dcf_base,
        "F10",
        "Editable residual working-capital assumption. For subscription software this reconciles filed aggregate noncash current assets less current liabilities to separately disclosed operating drivers; for other operating models it is an explicit overlay on working-capital days.",
    )
    _set_comment(dcf_base, "F11", "Source: Tax assumption from payload.assumptions.taxRate.")
    _set_comment(dcf_base, "F14", "Editable revenue-growth assumption for the first forecast stage; initialized from filed revenue history, not issuer guidance.")
    _set_comment(dcf_base, "F15", "Editable starting EBIT margin, initialized from the latest filed margin or three-year average according to the operating archetype.")
    _set_comment(dcf_base, "F16", "Editable forecast gross margin initialized from the latest filed gross profit divided by filed revenue; forecast cost of revenue reconciles to this input.")
    _set_comment(dcf_base, "F17", "Editable steady-state EBIT margin, initialized from the three-year filed average.")
    _set_comment(dcf_base, "F19", "Editable Stage 2 revenue growth assumption; the forecast fades toward this rate after the Stage 1 duration.")
    _set_comment(dcf_base, "I13", "Editable analyst assumption for the forecast years used to fade EBIT margin to its steady-state assumption.")
    _set_comment(dcf_base, "I14", "Editable forecast duration for Stage 1 revenue growth.")
    _set_comment(dcf_base, "I15", "Editable forecast duration for the linear fade from Stage 1 to Stage 2 revenue growth.")
    _set_comment(dcf_base, "I16", "Source: current diluted shares from the market payload, displayed in millions to match workbook currency units.")
    _set_comment(dcf_base, "I17", "Source: current share price from the market payload.")
    _set_comment(dcf_base, "C16", "Source: Terminal exit multiple from payload.assumptions.terminal.exitMultiple.")
    _set_comment(dcf_base, "F20", "Source: Noncontrolling interest from the SEC balance sheet; not an equity plug.")
    _set_comment(dcf_base, "F21", "Source: Preferred equity from SEC canonical history; explicit zero only when the filing has no preferred stock fact.")

    as_of = _safe_date(company.get("asOfDate"))
    if as_of is not None:
        _safe_set(dcf_base, "I9", as_of)

    closing_date_raw = (
        transaction.get("closingDate")
        or transaction.get("closeDate")
        or payload.get("closingDate")
    )
    _safe_set_or_clear(dcf_base, "I10", _safe_date(closing_date_raw))

    fiscal_year = as_of.year if as_of is not None else datetime.now().year
    fiscal_end = _fiscal_year_end_date(company.get("fiscalYearEnd"), fiscal_year)
    if fiscal_end is None:
        fiscal_end = date(fiscal_year, 12, 31)
    _safe_set(dcf_base, "I11", fiscal_end)



def _sync_shared_scenario_inputs(scenario_sheet: Worksheet, dcf_base: Worksheet, *, nwc_multiplier: float) -> None:
    # Preserve scenario-specific formulas and apply only payload-driven shared assumptions.
    for cell in (
        "C9", "C11", "C12", "C17", "F9", "F11", "F13", "F14", "F15", "F16", "F17", "F19", "F20", "F21",
        "C16", "I9", "I10", "I11", "I13", "I14", "I15", "I16", "I17", "Q103", "AA17", "AA18", "AA19",
    ):
        scenario_sheet[cell].value = dcf_base[cell].value

    base_nwc_change = _to_float(dcf_base["F10"].value)
    if base_nwc_change is not None:
        scenario_sheet["F10"].value = base_nwc_change * max(0.0, nwc_multiplier)



def _harden_growth_rate_formulas(*scenario_sheets: Worksheet) -> None:
    # Guard CAGR calculations against divide-by-zero in low-data scenarios.
    formula_map = {
        "S24": '=IFERROR((I24/H24)^(1/(COLUMNS(H24:I24)-1))-1,"")',
        "T24": '=IFERROR((Q24/J24)^(1/(COLUMNS(J24:Q24)-1))-1,"")',
        "S27": '=IFERROR((I27/H27)^(1/(COLUMNS(H27:I27)-1))-1,"")',
        "T27": '=IFERROR((Q27/J27)^(1/(COLUMNS(J27:Q27)-1))-1,"")',
    }
    for sheet in scenario_sheets:
        for cell_ref, formula in formula_map.items():
            _force_set(sheet, cell_ref, formula)
        # Harden any remaining CAGR-style formulas in S/T columns that still
        # divide by historical anchors without IFERROR wrappers.
        for row in range(20, 131):
            for col in ("S", "T"):
                cell_ref = f"{col}{row}"
                value = sheet[cell_ref].value
                if not (isinstance(value, str) and value.startswith("=")):
                    continue
                upper_value = value.upper()
                if "IFERROR(" in upper_value:
                    continue
                if "COLUMNS(" not in upper_value:
                    continue
                expression = value[1:].lstrip("+")
                _force_set(sheet, cell_ref, f'=IFERROR({expression},"")')



def _normalize_dcf_waterfall_formulas(
    data_recalc: Worksheet,
    timeline_years: list[int | None],
    historical_years: set[int],
    *scenario_sheets: Worksheet,
) -> None:
    # D&A is shown as an operating memo line, not a second operating expense.
    for sheet in scenario_sheets:
        for idx, col in enumerate(DCF_TIMELINE_COLUMNS):
            _force_set(sheet, f"{col}48", f"=-({col}36+{col}39+{col}45)")
            _force_set(sheet, f"{col}49", f"=IFERROR(-{col}48/{col}20,0)")
            _force_set(sheet, f"{col}51", f"={col}32+{col}48")

            year = timeline_years[idx] if idx < len(timeline_years) else None
            if year is None:
                for row in range(18, 116):
                    _force_set(sheet, f"{col}{row}", None)
                continue
            if year in historical_years:
                recalc_col = RECALC_COLUMNS[idx]
                capex_ref = f"'Data Given (Recalculated)'!{recalc_col}35"
                da_ref = f"'Data Given (Recalculated)'!{recalc_col}36"
                nwc_ref = f"'Data Given (Recalculated)'!{recalc_col}41"

                _force_set(
                    sheet,
                    f"{col}65",
                    f'=IF({capex_ref}="","",-{capex_ref})'
                    if data_recalc[f"{recalc_col}35"].value is not None
                    else '=""',
                )
                _force_set(
                    sheet,
                    f"{col}68",
                    f'=IF({da_ref}="","",{da_ref})'
                    if data_recalc[f"{recalc_col}36"].value is not None
                    else '=""',
                )
                _force_set(sheet, f"{col}69", f"=IFERROR({col}68/{col}20,0)")
                _force_set(
                    sheet,
                    f"{col}77",
                    f'=IF({nwc_ref}="","",-{nwc_ref})'
                    if data_recalc[f"{recalc_col}41"].value is not None
                    else '=""',
                )
            else:
                _force_set(sheet, f"{col}65", f"=-{col}20*$F$9")
                _force_set(sheet, f"{col}68", f"={col}20*$F$13")
                _force_set(sheet, f"{col}69", "=$F$13")
                days_sheet = "'Data Given (Recalculated)'!"
                current_nwc = (
                    f"({col}20*{days_sheet}$C$50/365+{col}30*{days_sheet}$C$51/365-"
                    f"{col}30*{days_sheet}$C$52/365+{col}20*$F$10)"
                )
                prior_year = timeline_years[idx - 1] if idx > 0 else None
                if idx > 0 and prior_year in historical_years:
                    prior_recalc_col = RECALC_COLUMNS[idx - 1]
                    prior_nwc = f"{days_sheet}{prior_recalc_col}40"
                elif idx > 0:
                    prior_col = DCF_TIMELINE_COLUMNS[idx - 1]
                    prior_nwc = (
                        f"({prior_col}20*{days_sheet}$C$50/365+{prior_col}30*{days_sheet}$C$51/365-"
                        f"{prior_col}30*{days_sheet}$C$52/365+{prior_col}20*$F$10)"
                    )
                else:
                    prior_nwc = "0"
                _force_set(sheet, f"{col}77", f"=IFERROR(-({current_nwc}-{prior_nwc}),\"\")")
            _force_set(sheet, f"{col}66", f"=IFERROR(-{col}65/{col}20,0)")
            _force_set(sheet, f"{col}70", f"=IFERROR(-{col}68/{col}65,0)")

        for col in DCF_TIMELINE_COLUMNS:
            sheet[f"{col}78"].number_format = sheet["J78"].number_format
            sheet[f"{col}79"].number_format = sheet["J79"].number_format


def apply_incomplete_operating_dcf(
    workbook: Workbook,
    payload: dict[str, Any],
    input_cells: dict[str, dict[str, str]],
    timeline_years: list[int | None],
    historical_years: set[int],
) -> None:
    requirements = payload.get("requiredInputs")
    requirements = requirements if isinstance(requirements, list) else []
    capex_requirements = [
        item for item in requirements
        if isinstance(item, dict) and item.get("key") == "capex" and isinstance(item.get("fiscalYear"), int)
    ]
    if not capex_requirements:
        raise ValueError("Incomplete operating DCF requires an exact historical CapEx input manifest.")

    recalculated = _sheet(workbook, SHEET_DATA_RECALCULATED)
    for requirement in capex_requirements:
        year = int(requirement["fiscalYear"])
        if year not in timeline_years:
            raise ValueError(f"FY{year} CapEx input is outside the workbook timeline.")
        timeline_index = timeline_years.index(year)
        input_identity = f"capex:{year}"
        destination = input_cells.get(input_identity)
        if not destination or destination.get("sheet") != "Input Required":
            raise ValueError(f"FY{year} CapEx input has no editable workbook cell.")
        input_cell = destination["cell"]
        recalculated_column = RECALC_COLUMNS[timeline_index]
        _force_set(
            recalculated,
            f"{recalculated_column}35",
            f'=IF(\'Input Required\'!$B$3="READY",\'Input Required\'!{input_cell},"")',
        )

    recent_years = sorted(historical_years)[-3:]
    if len(recent_years) != 3:
        raise ValueError("Incomplete operating DCF requires three historical years for a CapEx ratio.")
    capex_ratio_terms: list[str] = []
    for year in recent_years:
        try:
            timeline_index = timeline_years.index(year)
        except ValueError as error:
            raise ValueError(f"FY{year} is outside the incomplete DCF timeline.") from error
        column = RECALC_COLUMNS[timeline_index]
        capex_ratio_terms.append(
            f"ABS('{SHEET_DATA_RECALCULATED}'!{column}35/'{SHEET_DATA_RECALCULATED}'!{column}12)"
        )
    capex_ratio_formula = (
        '=IF(\'Input Required\'!$B$3<>"READY","",AVERAGE('
        + ",".join(capex_ratio_terms)
        + "))"
    )
    for sheet_name in (SHEET_DCF_BASE, SHEET_DCF_BULL, SHEET_DCF_BEAR):
        sheet = _sheet(workbook, sheet_name)
        _force_set(sheet, "F9", capex_ratio_formula)
        sheet["F9"].number_format = "0.0%"
        sheet["F9"].font = Font(name="Arial", size=10, color="008000")
        _set_comment(
            sheet,
            "F9",
            "Editable formula assumption: average filed CapEx / revenue for the latest three fiscal years. The formula remains blank until every required actual and source reference is entered.",
        )

    def guard_formula(sheet: Worksheet, coordinate: str) -> None:
        formula = sheet[coordinate].value
        if not (isinstance(formula, str) and formula.startswith("=")):
            return
        if "'Input Required'!$B$3" in formula:
            return
        _force_set(sheet, coordinate, f'=IF(\'Input Required\'!$B$3<>"READY","",{formula[1:]})')

    forecast_columns = {
        DCF_TIMELINE_COLUMNS[index]
        for index, year in enumerate(timeline_years)
        if year is not None and year not in historical_years
    }
    for sheet_name in (SHEET_DCF_BASE, SHEET_DCF_BULL, SHEET_DCF_BEAR):
        sheet = _sheet(workbook, sheet_name)
        for column in forecast_columns:
            for row in range(18, 91):
                guard_formula(sheet, f"{column}{row}")
        for row in range(91, 131):
            for column_index in range(1, sheet.max_column + 1):
                coordinate = sheet.cell(row=row, column=column_index).coordinate
                guard_formula(sheet, coordinate)
        for coordinate in ("C9", "C10", "C12", "C13", "C14"):
            guard_formula(sheet, coordinate)

    for sheet_name in (SHEET_COVER, SHEET_OUTPUTS):
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell, MergedCell):
                    guard_formula(sheet, cell.coordinate)

    for index, year in enumerate(timeline_years):
        if year is None or year in historical_years:
            continue
        column = RECALC_COLUMNS[index]
        for row in range(12, 42):
            guard_formula(recalculated, f"{column}{row}")



def _normalize_public_dcf_layout(
    outputs: Worksheet,
    dcf_base: Worksheet,
    dcf_bull: Worksheet,
    dcf_bear: Worksheet,
    data_recalc: Worksheet,
    payload: dict[str, Any],
    divisor: float,
    timeline_years: list[int | None],
    historical_years: set[int],
) -> None:
    _link_dcf_income_statement_to_recalculated_data(timeline_years, historical_years, dcf_base, dcf_bull, dcf_bear)
    _normalize_public_dcf_assumption_block(dcf_base, dcf_bull, dcf_bear, payload, divisor)
    _enforce_core_public_dcf_formulas(dcf_base, dcf_bull, dcf_bear)
    _normalize_dcf_waterfall_formulas(data_recalc, timeline_years, historical_years, dcf_base, dcf_bull, dcf_bear)
    _enforce_outputs_bridge_formulas(outputs)



def _sync_scenario_formula_backbone(dcf_base: Worksheet, *scenario_sheets: Worksheet) -> None:
    # Keep formula topology identical across Base/Bull/Bear to avoid scenario drift.
    scenario_specific_formula_cells = {"F12"}
    for row in dcf_base.iter_rows(min_row=1, max_row=dcf_base.max_row, min_col=1, max_col=dcf_base.max_column):
        for base_cell in row:
            formula = base_cell.value
            if not (isinstance(formula, str) and formula.startswith("=")):
                continue
            if base_cell.coordinate in scenario_specific_formula_cells:
                continue
            for scenario in scenario_sheets:
                scenario_cell = scenario[base_cell.coordinate]
                if isinstance(scenario_cell, MergedCell):
                    continue
                scenario_cell.value = formula



def _apply_capex_schedule_to_dcf(
    dcf_base: Worksheet,
    dcf_bull: Worksheet,
    dcf_bear: Worksheet,
    payload: dict[str, Any],
    timeline_years: list[int | None],
    divisor: float,
) -> None:
    # Projection-period CapEx rows are formula-driven from assumptions.
    # Avoid stamping hardcoded year-by-year values that can drift.
    return



def _apply_scenario_snapshots_to_dcf(
    dcf_base: Worksheet,
    dcf_bull: Worksheet,
    dcf_bear: Worksheet,
    payload: dict[str, Any],
    timeline_years: list[int | None],
    divisor: float,
) -> None:
    base_snapshot = _scenario_snapshot(payload, SCENARIO_BASE)
    bull_snapshot = _scenario_snapshot(payload, SCENARIO_BULL)
    bear_snapshot = _scenario_snapshot(payload, SCENARIO_BEAR)

    if base_snapshot is None and bull_snapshot is None and bear_snapshot is None:
        return

    if base_snapshot is not None:
        _apply_scenario_snapshot_to_sheet(dcf_base, base_snapshot, timeline_years, divisor)
    if bull_snapshot is not None:
        _apply_scenario_snapshot_to_sheet(dcf_bull, bull_snapshot, timeline_years, divisor)
    if bear_snapshot is not None:
        _apply_scenario_snapshot_to_sheet(dcf_bear, bear_snapshot, timeline_years, divisor)

    for sheet in (dcf_base, dcf_bull, dcf_bear):
        first_forecast_idx = next(
            (
                idx
                for idx, col in enumerate(DCF_TIMELINE_COLUMNS)
                if isinstance(sheet[f"{col}18"].value, str)
                and str(sheet[f"{col}18"].value).endswith("E")
            ),
            None,
        )
        if first_forecast_idx is None or first_forecast_idx == 0:
            continue
        first_forecast_col = DCF_TIMELINE_COLUMNS[first_forecast_idx]
        prior_col = DCF_TIMELINE_COLUMNS[first_forecast_idx - 1]
        _force_set(
            sheet,
            f"{first_forecast_col}20",
            f"={prior_col}20*(1+{first_forecast_col}21)",
        )



def _finalize_assumption_block_cleanup(*scenario_sheets: Worksheet) -> None:
    # Ensure assumption area is clean and free of duplicate helper numbers.
    for sheet in scenario_sheets:
        _safe_set(sheet, "B18", "Income Statement")
        _safe_set_or_clear(sheet, "B19", None)
        _safe_set(sheet, "E14", "Stage 1 Revenue Growth")
        _safe_set(sheet, "E15", "Starting EBIT Margin")
        _safe_set(sheet, "E16", "Forecast Gross Margin")
        _safe_set(sheet, "E17", "Steady-State EBIT Margin")
        _safe_set_or_clear(sheet, "F18", None)
        _safe_set(sheet, "E19", "Stage 2 Revenue Growth")
        for address in ("C18", "C19", "D18", "D19", "E18", "G18", "G19"):
            _safe_set_or_clear(sheet, address, None)
        # Remove lingering note indicators (red triangles) from template/input mapping.
        for address in (
            "C13",
            "C16",
            "C17",
            "F9",
            "F10",
            "F11",
            "F12",
            "F13",
            "F14",
            "F18",
        ):
            sheet[address].comment = None



def _add_prior_actual_year_display_column(
    dcf_base: Worksheet,
    dcf_bull: Worksheet,
    dcf_bear: Worksheet,
    payload: dict[str, Any],
    timeline_years: list[int | None],
    historical_years: set[int],
    divisor: float,
) -> None:
    if not timeline_years:
        return
    first_timeline_year = next((year for year in timeline_years if year is not None), None)
    if first_timeline_year is None:
        return
    prior_candidates = [year for year in historical_years if year < first_timeline_year]
    if not prior_candidates:
        return

    prior_year = max(prior_candidates)
    prior_label = f"FY{prior_year}A"

    revenue = _historical_value_for_year(payload, prior_year, statement="income", keys=["Total Revenue", "Revenue", "Sales"])
    cost_of_revenue = _historical_value_for_year(
        payload,
        prior_year,
        statement="income",
        keys=["Cost of Revenue", "COGS", "Cost Of Revenue", "Cost of Sales", "Purchases"],
    )
    rnd = _historical_value_for_year(payload, prior_year, statement="income", keys=["Research & Development", "R&D", "Research and Development"])
    sga = _historical_value_for_year(
        payload,
        prior_year,
        statement="income",
        keys=["SG&A", "SGA", "General and Administrative", "GeneralAndAdministrative", "G&A", "GA"],
    )
    da = _historical_value_for_year(
        payload,
        prior_year,
        statement="income",
        keys=["D&A (included in Operating)", "D&A", "DA", "Depreciation & Amortization", "Depreciation"],
    )
    if da is None:
        da = _historical_value_for_year(payload, prior_year, statement="cashflow", keys=["Depreciation"])
    other_opex = _historical_value_for_year(payload, prior_year, statement="income", keys=["Other Operating Expenses", "Other"])
    ebit = _historical_value_for_year(payload, prior_year, statement="income", keys=["Operating Income (EBIT)", "EBIT", "Operating Income"])
    capex = _historical_value_for_year(payload, prior_year, statement="cashflow", keys=["Capex", "Capital Expenditures", "Capital Expenditure"])

    for sheet in (dcf_base, dcf_bull, dcf_bear):
        # Carry timeline header style to the added prior-year display column.
        for row in (18, 63, 72):
            sheet[f"G{row}"]._style = copy(sheet[f"H{row}"]._style)
            _force_set(sheet, f"G{row}", prior_label)

        # Apply consistent number/border styles from first timeline column.
        for row in (20, 24, 27, 30, 32, 33, 36, 39, 42, 45, 48, 49, 51, 52, 54, 55, 57, 60, 65, 66, 67, 68, 69, 74, 75, 76, 77, 78, 79):
            sheet[f"G{row}"]._style = copy(sheet[f"H{row}"]._style)

        if revenue is not None:
            _force_set(sheet, "G20", _scale(revenue, divisor))
            _force_set(sheet, "H21", "=IFERROR(H20/G20-1,0)")
        _force_set(sheet, "G21", "-")
        if cost_of_revenue is not None:
            _force_set(sheet, "G24", _scale(cost_of_revenue, divisor))
        _force_set(sheet, "G25", "=IFERROR(G24/G20,0)")
        _force_set(sheet, "G28", "=IFERROR(G27/G20,0)")
        _force_set(sheet, "G30", "=G24+G27")
        _force_set(sheet, "G32", "=G20-G30")
        _force_set(sheet, "G33", "=IFERROR(G32/G20,0)")
        if rnd is not None:
            _force_set(sheet, "G36", _scale(rnd, divisor))
        _force_set(sheet, "G37", "=IFERROR(G36/G20,0)")
        if sga is not None:
            _force_set(sheet, "G39", _scale(sga, divisor))
        _force_set(sheet, "G40", "=IFERROR(G39/G20,0)")
        if da is not None:
            _force_set(sheet, "G42", _scale(da, divisor))
        _force_set(sheet, "G43", "=IFERROR(G42/G20,0)")
        if other_opex is not None:
            _force_set(sheet, "G45", _scale(other_opex, divisor))
        _force_set(sheet, "G46", "=IFERROR(G45/G20,0)")
        _force_set(sheet, "G48", "=-(G36+G39+G45)")
        _force_set(sheet, "G49", "=IFERROR(-G48/G20,0)")
        if ebit is not None:
            _force_set(sheet, "G51", _scale(ebit, divisor))
        else:
            _force_set(sheet, "G51", "=G32+G48")
        _force_set(sheet, "G52", "=IFERROR(G51/G20,0)")
        _force_set(sheet, "G54", "=G51+G68")
        _force_set(sheet, "G55", "=IFERROR(G54/G20,0)")
        _force_set(sheet, "G57", "=-G51*$F$11")
        _force_set(sheet, "G58", "=$F$11")
        _force_set(sheet, "G60", "=G51+G57")
        _force_set(sheet, "G61", "=IFERROR(G60/G20,0)")

        if capex is not None:
            _force_set(sheet, "G65", -(_scale(capex, divisor) or 0.0))
        else:
            _force_set(sheet, "G65", "=H65")
        _force_set(sheet, "G66", "=IFERROR(-G65/G20,0)")
        _safe_set_or_clear(sheet, "G67", None)
        if da is not None:
            _force_set(sheet, "G68", _scale(da, divisor))
        else:
            _force_set(sheet, "G68", "=G69*G20")
        _force_set(sheet, "G69", "=$F$13")
        _force_set(sheet, "G70", "=IFERROR(-G68/G65,0)")

        _force_set(sheet, "G74", "=G60")
        _force_set(sheet, "G75", "=G68")
        _force_set(sheet, "G76", "=G65")
        _force_set(sheet, "G77", "=0")
        _force_set(sheet, "G78", "=SUM(G74:G77)")
        _force_set(sheet, "G79", "=IFERROR(G78/G20,0)")


def _map_sensitivity_blocks(
    dcf_base: Worksheet,
    dcf_bull: Worksheet,
    dcf_bear: Worksheet,
    payload: dict[str, Any],
    divisor: float,
) -> None:
    assumptions = payload.get("assumptions", {})
    assumptions = assumptions if isinstance(assumptions, dict) else {}
    terminal = assumptions.get("terminal", {})
    terminal = terminal if isinstance(terminal, dict) else {}
    sensitivities = payload.get("sensitivities", {})
    sensitivities = sensitivities if isinstance(sensitivities, dict) else {}

    base_wacc = _sanitize_wacc_rate(assumptions.get("waccRate")) or 0.10
    base_growth = _sanitize_terminal_growth_rate(terminal.get("g"), reference_wacc=base_wacc) or 0.025
    base_revenue_growth = _to_float(assumptions.get("revenueGrowth")) or 0.03
    base_ebit_margin = _to_float(assumptions.get("ebitMargin")) or _to_float(assumptions.get("ebitMarginTarget")) or 0.15

    wacc_axis = _axis_from_bounds(
        _numeric_list(sensitivities.get("waccAxis")) or _numeric_list(sensitivities.get("waccGrid")),
        center=base_wacc,
        step=0.01,
        min_value=0.01,
        max_value=0.30,
    )
    growth_axis = _axis_from_bounds(
        _numeric_list(sensitivities.get("terminalGrowthAxis")) or _numeric_list(sensitivities.get("gGrid")),
        center=base_growth,
        step=0.005,
        min_value=0.0,
        max_value=max(0.001, min(0.08, min(wacc_axis) - 0.001)),
    )
    revenue_growth_axis = _axis_from_bounds(
        _numeric_list(sensitivities.get("revenueGrowthAxis")),
        center=base_revenue_growth,
        step=0.01,
        min_value=-0.10,
        max_value=0.30,
    )
    ebit_margin_axis = _axis_from_bounds(
        _numeric_list(sensitivities.get("ebitMarginAxis")),
        center=base_ebit_margin,
        step=0.01,
        min_value=0.01,
        max_value=0.60,
    )

    for sheet in (dcf_base, dcf_bull, dcf_bear):
        # Add visual space between the two sensitivity tables.
        current_i_width = sheet.column_dimensions["I"].width
        if current_i_width is None or current_i_width < 14:
            sheet.column_dimensions["I"].width = 14

        _safe_set(sheet, "C117", "Enterprise Value - WACC x Terminal Growth")
        _safe_set(sheet, "D118", "WACC")
        _force_set(sheet, "C119", f"='{SHEET_OUTPUTS}'!$D$35")
        _safe_set(sheet, "B121", "Terminal Growth")
        _safe_set(sheet, "B122", "Rate")
        _set_percent_axis_row(sheet, cells=("D119", "E119", "F119", "G119", "H119"), values=wacc_axis)
        _set_percent_axis_column(sheet, cells=("C120", "C121", "C122", "C123", "C124"), values=growth_axis)
        _force_set(sheet, "F119", "=$F$12")
        _force_set(sheet, "C122", "=$Q$103")

        first_forecast_idx = next(
            (
                idx
                for idx, col in enumerate(DCF_TIMELINE_COLUMNS)
                if isinstance(sheet[f"{col}18"].value, str)
                and str(sheet[f"{col}18"].value).endswith("E")
            ),
            None,
        )
        if first_forecast_idx is not None:
            forecast_columns = DCF_TIMELINE_COLUMNS[first_forecast_idx:]
            last_forecast_col = forecast_columns[-1]
            for row_idx in range(5):
                growth_ref = f"$C${120 + row_idx}"
                for col_idx in range(5):
                    wacc_ref = f"${chr(ord('D') + col_idx)}$119"
                    explicit_pv = "+".join(
                        f"{col}78/(1+{wacc_ref})^{col}84"
                        for col in forecast_columns
                    )
                    terminal_pv = (
                        f"{last_forecast_col}78*(1+{growth_ref})/({wacc_ref}-{growth_ref})"
                        f"/(1+{wacc_ref})^{last_forecast_col}83"
                    )
                    _force_set(
                        sheet,
                        f"{chr(ord('D') + col_idx)}{120 + row_idx}",
                        f'=IFERROR(IF({wacc_ref}>{growth_ref},{explicit_pv}+{terminal_pv},""),"")',
                    )

        _force_set(sheet, "C126", "=MIN(D120:H124)")
        _force_set(sheet, "C127", "=PERCENTILE(D120:H124,0.25)")
        _force_set(sheet, "C128", "=MEDIAN(D120:H124)")
        _force_set(sheet, "C129", "=PERCENTILE(D120:H124,0.75)")
        _force_set(sheet, "C130", "=MAX(D120:H124)")
        _safe_set(sheet, "B126", "Min")
        _safe_set(sheet, "B127", "Q1")
        _safe_set(sheet, "B128", "Median")
        _safe_set(sheet, "B129", "Q3")
        _safe_set(sheet, "B130", "Max")

        # Ensure every right-side table cell is materialized in sheet XML.
        for row in range(117, 131):
            for col in range(ord("I"), ord("O") + 1):
                ref = f"{chr(col)}{row}"
                if sheet[ref].value is None:
                    _force_set(sheet, ref, "")

        _safe_set(sheet, "I117", "Enterprise Value - Revenue Growth x EBIT Margin")
        _safe_set(sheet, "J118", "Revenue Growth")
        _safe_set(sheet, "I121", "EBIT Margin")
        _safe_set(sheet, "I122", "Rate")
        _set_percent_axis_row(sheet, cells=("J119", "K119", "L119", "M119", "N119"), values=revenue_growth_axis)
        _set_percent_axis_column(sheet, cells=("I120", "I121", "I122", "I123", "I124"), values=ebit_margin_axis)
        _force_set(sheet, "L119", "=$F$14")
        _force_set(sheet, "I122", f"={forecast_columns[0]}52" if first_forecast_idx is not None else "=$F$14")

        if first_forecast_idx is not None:
            prior_col = DCF_TIMELINE_COLUMNS[first_forecast_idx - 1] if first_forecast_idx > 0 else forecast_columns[0]
            last_forecast_col = forecast_columns[-1]
            for row_idx in range(5):
                margin_ref = f"$I${120 + row_idx}"
                for col_idx in range(5):
                    growth_ref = f"{chr(ord('J') + col_idx)}$119"
                    explicit_pv_terms: list[str] = []
                    for year_idx, col in enumerate(forecast_columns, start=1):
                        revenue = f"{prior_col}20*(1+{growth_ref})^{year_idx}"
                        fcff = f"({revenue}*{margin_ref}*(1-$F$11)+{revenue}*$F$13-$F$9-$F$10)"
                        explicit_pv_terms.append(f"{fcff}/(1+$F$12)^{col}84")
                    terminal_revenue = f"{prior_col}20*(1+{growth_ref})^{len(forecast_columns)}"
                    terminal_ebitda = f"({terminal_revenue}*{margin_ref}+{terminal_revenue}*$F$13)"
                    terminal_pv = f"{terminal_ebitda}*$C$16/(1+$F$12)^{last_forecast_col}83"
                    formula = "+".join([*explicit_pv_terms, terminal_pv])
                    _force_set(
                        sheet,
                        f"{chr(ord('J') + col_idx)}{120 + row_idx}",
                        f'=IFERROR({formula},"")',
                    )

        _safe_set(sheet, "I126", "Min")
        _safe_set(sheet, "I127", "Q1")
        _safe_set(sheet, "I128", "Median")
        _safe_set(sheet, "I129", "Q3")
        _safe_set(sheet, "I130", "Max")
        _force_set(sheet, "J126", "=MIN(J120:N124)")
        _force_set(sheet, "J127", "=PERCENTILE(J120:N124,0.25)")
        _force_set(sheet, "J128", "=MEDIAN(J120:N124)")
        _force_set(sheet, "J129", "=PERCENTILE(J120:N124,0.75)")
        _force_set(sheet, "J130", "=MAX(J120:N124)")

        # Highlight low/median/high valuation outcomes with a standard red-yellow-green gradient.
        for target_range in ("D120:H124", "J120:N124"):
            sheet.conditional_formatting.add(
                target_range,
                ColorScaleRule(
                    start_type="min",
                    start_color="F8696B",
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="max",
                    end_color="63BE7B",
                ),
            )

    _clear_sensitivity_blocks(dcf_bull, dcf_bear)



def _replace_template_placeholders(
    *,
    company_name: str | None,
    ticker: str,
    sheets: tuple[Worksheet, ...],
) -> None:
    long_name = company_name or ticker
    replacements = {
        "ABC/SNS Health": long_name,
        "ABC Corp.": long_name,
        "ABC Corp": long_name,
        "ABC": ticker,
    }

    for sheet in sheets:
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                if cell.data_type == "f":
                    continue
                if not isinstance(cell.value, str):
                    continue

                updated = cell.value
                for old, new in replacements.items():
                    updated = updated.replace(old, new)
                if updated != cell.value:
                    cell.value = updated

def _finalize_timeline_headers(
    outputs: Worksheet,
    dcf_base: Worksheet,
    dcf_bull: Worksheet,
    dcf_bear: Worksheet,
    timeline_years: list[int | None],
    historical_years: set[int],
) -> None:
    # Final guardrail: force all timeline headers to explicit FY labels as text.
    for idx, col in enumerate(DCF_TIMELINE_COLUMNS):
        year = timeline_years[idx]
        if year is None:
            _safe_set_or_clear(outputs, f"{col}6", None)
            for sheet in (dcf_base, dcf_bull, dcf_bear):
                for row in (18, 63, 72, 89):
                    _safe_set_or_clear(sheet, f"{col}{row}", None)
            continue
        label = f"FY{year}{'A' if year in historical_years else 'E'}"
        _force_set(outputs, f"{col}6", label)
        outputs[f"{col}6"].number_format = "@"
        for sheet in (dcf_base, dcf_bull, dcf_bear):
            _force_set(sheet, f"{col}18", label)
            _force_set(sheet, f"{col}63", label)
            _force_set(sheet, f"{col}72", label)
            _force_set(sheet, f"{col}89", label)
            sheet[f"{col}89"].number_format = "@"



def _reset_dcf_sheet_view_to_top(*sheets: Worksheet) -> None:
    # Ensure each scenario tab opens at the top of the sheet instead of
    # preserving a template viewport near the bottom.
    for sheet in sheets:
        sheet.sheet_view.topLeftCell = "A1"
        if sheet.sheet_view.selection:
            sheet.sheet_view.selection[0].activeCell = "A1"
            sheet.sheet_view.selection[0].sqref = "A1"



def _remove_assumption_breakdown(workbook: Workbook, cover: Worksheet) -> None:
    if SHEET_ASSUMPTION_BREAKDOWN in workbook.sheetnames:
        workbook.remove(workbook[SHEET_ASSUMPTION_BREAKDOWN])
    if SHEET_DATA_ORIGINAL in workbook.sheetnames:
        workbook.remove(workbook[SHEET_DATA_ORIGINAL])
    if "Data ->" in workbook.sheetnames:
        workbook.remove(workbook["Data ->"])

    # Keep cover TOC coherent after removing helper tabs.
    _safe_set(cover, "F15", "Data Given (Recalculated)")
    _safe_set_or_clear(cover, "F16", None)
    _safe_set_or_clear(cover, "F17", None)
