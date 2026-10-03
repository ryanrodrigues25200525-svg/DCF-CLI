from __future__ import annotations

import math
import re
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

_INPUT_FONT = Font(name='Arial', size=10, color='0000FF')
_FORMULA_FONT = Font(name='Arial', size=10, color='000000')
_BODY_FONT = Font(name='Arial', size=10, color='000000')
_SMALL_FONT = Font(name='Arial', size=9, color='404040')
_HEADER_FONT = Font(name='Arial', size=9, bold=True, color='FFFFFF')
_SECTION_FONT = Font(name='Arial', size=10, bold=True, color='17365D')
_TITLE_FILL = PatternFill('solid', fgColor='17365D')
_SECTION_FILL = PatternFill('solid', fgColor='D9EAF7')
_INPUT_FILL = PatternFill('solid', fgColor='FFF2CC')
_FORMULA_FILL = PatternFill('solid', fgColor='E2F0D9')
_HEADER_FILL = PatternFill('solid', fgColor='365F91')
_MONEY_FORMAT = '#,##0.0;[Red](#,##0.0);-'
_PERCENT_FORMAT = '0.0%;[Red](0.0%);0.0%'
_YEAR_FORMAT = '0'
_PRICE_FORMAT = '$#,##0.00;[Red]($#,##0.00);-'


def _record(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _finite(value: Any, name: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(float(value)):
        raise ValueError(f'Mature-pharma workbook requires a finite value for {name}.')
    return float(value)


def _source_line(line: Any, name: str, year: int) -> dict[str, Any]:
    record = _record(line)
    if record.get('source') not in {'sec_native', 'derived'}:
        raise ValueError(f'FY{year} PFE workbook input {name} is missing or ambiguous.')
    value = _finite(record.get('value'), name)
    sources = record.get('sources')
    if not isinstance(sources, list) or not sources or any(not _record(source).get('accession') or not _record(source).get('filed') for source in sources):
        raise ValueError(f'FY{year} PFE workbook input {name} has incomplete SEC lineage.')
    return record


def _line_value(line: Any, name: str, year: int) -> float:
    return _finite(_source_line(line, name, year).get('value'), name)


def _write(
    sheet: Worksheet,
    coordinate: str,
    value: Any,
    *,
    formula: bool = False,
    input_cell: bool = False,
    number_format: str | None = None,
    wrap: bool = False,
) -> None:
    cell = sheet[coordinate]
    cell.value = value
    cell.font = _INPUT_FONT if input_cell else _FORMULA_FONT if formula else _BODY_FONT
    cell.fill = _INPUT_FILL if input_cell else _FORMULA_FILL if formula else PatternFill(fill_type=None)
    cell.alignment = Alignment(vertical='center', wrap_text=wrap)
    if number_format:
        cell.number_format = number_format


def _section(sheet: Worksheet, row: int, title: str, end_column: int = 17) -> None:
    for column in range(1, end_column + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = _SECTION_FILL
        cell.font = _SECTION_FONT
    sheet.cell(row=row, column=1, value=title)


def _amount_mm(value: float) -> float:
    return value / 1_000_000


def patent_names_for_model(product_name: str) -> set[str]:
    key = re.sub(r'\s*\([a-z0-9]+\)\s*$', '', product_name.lower())
    key = ''.join(character for character in key if character.isalnum())
    if key == 'prevnarfamily':
        aliases = {'prevnar13prevenar13', 'prevnar20prevenar20'}
    elif key == 'vyndaqelfamily':
        aliases = {'vyndaqelvyndamaxvynmac'}
    elif key == 'braftovimektovi':
        aliases = {'braftovi', 'mektovi'}
    else:
        aliases = {key}
    return aliases


def _validate_model(
    model: dict[str, Any],
    *,
    allow_missing_product_revenue: bool = False,
) -> tuple[list[dict[str, Any]], list[int], dict[str, Any], list[dict[str, Any]]]:
    history = _record(model.get('history'))
    annual = history.get('annual')
    years = history.get('years')
    assumptions = _record(model.get('assumptions'))
    if not isinstance(annual, list) or len(annual) != 3 or not isinstance(years, list) or len(years) != 3:
        raise ValueError('Mature-pharma workbook requires three annual source periods.')
    rows = [_record(row) for row in annual]
    if [row.get('year') for row in rows] != years or any(years[i] != years[i - 1] + 1 for i in range(1, len(years))):
        raise ValueError('PFE product history must be aligned and consecutive.')
    if _finite(assumptions.get('forecastYears'), 'forecastYears') != 5 or _finite(assumptions.get('baseYear'), 'baseYear') != years[-1]:
        raise ValueError('Mature-pharma workbook requires five forecast years and a matching base year.')
    products = assumptions.get('products')
    if not isinstance(products, list) or not products:
        raise ValueError('Mature-pharma workbook requires product-level analyst assumptions.')
    product_assumptions = [_record(product) for product in products]
    sources = _record(assumptions.get('assumptionSources'))
    if not sources or any(not isinstance(value, str) or not value.strip() for value in sources.values()):
        raise ValueError('Mature-pharma workbook requires a source or analyst-input disclosure for every global assumption.')
    wacc = _finite(assumptions.get('wacc'), 'wacc')
    growth = _finite(assumptions.get('terminalGrowthRate'), 'terminalGrowthRate')
    if wacc <= 0 or growth < 0 or growth >= wacc:
        raise ValueError('Mature-pharma terminal growth must remain below WACC.')
    latest_products = [_record(product) for product in _record(rows[-1].get('pharma')).get('products', [])]
    if len(latest_products) != len(product_assumptions):
        raise ValueError('PFE product assumptions must match the latest filed product rows.')
    expected = {str(product.get('product_name') or '') for product in latest_products}
    actual = {str(product.get('productName') or '') for product in product_assumptions}
    if expected != actual:
        raise ValueError('PFE product assumption names must match filed product rows exactly.')
    return rows, years, assumptions, product_assumptions


def _source_register(model: dict[str, Any]) -> list[list[Any]]:
    history = _record(model.get('history'))
    annual = history.get('annual') if isinstance(history.get('annual'), list) else []
    rows: list[list[Any]] = []

    def add_line(period: Any, name: str, line: Any) -> None:
        record = _record(line)
        sources = record.get('sources') if isinstance(record.get('sources'), list) else []
        if not sources:
            rows.append([period, name, record.get('value'), None, None, None, record.get('source'), record.get('method'), record.get('concept'), None, None, None, None])
            return
        for source_item in sources:
            source = _record(source_item)
            rows.append([
                period, name, record.get('value'), source.get('reported_value'), source.get('unit'), source.get('unit_scale'),
                record.get('source'), record.get('method'), source.get('concept') or record.get('concept'),
                source.get('label'), source.get('statement'), source.get('accession'), source.get('filed'),
            ])

    for item in annual:
        year = _record(item)
        pharma = _record(year.get('pharma'))
        for product_item in pharma.get('products', []) if isinstance(pharma.get('products'), list) else []:
            product = _record(product_item)
            add_line(year.get('year'), f"Product revenue — {product.get('product_name')}", product.get('revenue'))
        add_line(year.get('year'), 'Note 17 reported total revenue', pharma.get('reported_total_revenue'))
        for field, label in (
            ('revenue', 'Consolidated revenue'), ('ebit', 'EBIT'), ('interestExpense', 'Interest expense'),
            ('taxRate', 'Effective tax rate'), ('depreciation', 'Depreciation and amortization'), ('capex', 'Capital expenditures'),
            ('nwcChange', 'Working-capital change'), ('cash', 'Cash and equivalents'),
            ('marketableSecurities', 'Short-term investments'), ('debt', 'Interest-bearing debt'),
            ('nonControllingInterest', 'Noncontrolling interest'), ('preferredEquity', 'Preferred equity'),
            ('dilutedShares', 'Diluted shares'),
        ):
            add_line(year.get('year'), label, year.get(field))
        for patent_item in pharma.get('patents', []) if isinstance(pharma.get('patents'), list) else []:
            patent = _record(patent_item)
            add_line(year.get('year'), f"{patent.get('product_name')} {patent.get('region')} {patent.get('metric')} — {patent.get('reported_text')}", patent.get('year'))
    assumptions = _record(model.get('assumptions'))
    for field, note in _record(assumptions.get('assumptionSources')).items():
        rows.append(['Assumption', field, None, None, None, None, 'analyst_or_market', note, None, None, None, None, None])
    for item in assumptions.get('products', []) if isinstance(assumptions.get('products'), list) else []:
        product = _record(item)
        rows.append(['Assumption', f"{product.get('productName')} modeled LOE and growth", product.get('modeledGlobalLoeYear'), None, 'calendar year', 'analyst_input', 'analyst_input', product.get('sourceNote'), None, product.get('productName'), None, None, None])
    return rows


def _map_data_review(workbook: Workbook, model: dict[str, Any]) -> None:
    sheet = workbook.create_sheet('Data Review')
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'A2'
    headers = ['Fiscal year', 'Model field', 'Canonical value', 'Reported value', 'Unit', 'Scale', 'Source type', 'Method / analyst note', 'SEC concept', 'Filed label', 'SEC statement/table', 'Accession', 'Filed date']
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    for row in _source_register(model):
        sheet.append(row)
    widths = [14, 44, 20, 18, 18, 18, 20, 72, 42, 48, 54, 25, 16]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in range(2, sheet.max_row + 1):
        for cell in sheet[row]:
            cell.font = _BODY_FONT
            cell.alignment = Alignment(vertical='top', wrap_text=cell.column >= 8)
            if cell.column in (3, 4) and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                if cell.value == 0 or abs(cell.value) >= 0.005:
                    cell.number_format = '#,##0.00'
        source_lines = [math.ceil(len(str(sheet.cell(row=row, column=column).value or '')) / widths[column - 1]) for column in range(8, 12)]
        sheet.row_dimensions[row].height = max(15, min(60, 12 * max(source_lines or [1])))
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def apply_mature_pharma_model(workbook: Workbook, payload: dict[str, Any]) -> None:
    model = _record(payload.get('maturePharmaModel'))
    allow_missing_product_revenue = payload.get('buildStatus') == 'input_required'
    rows, years, assumptions, product_assumptions = _validate_model(
        model,
        allow_missing_product_revenue=allow_missing_product_revenue,
    )
    ticker = str(_record(payload.get('company')).get('ticker') or '').upper()
    if ticker != 'PFE':
        raise ValueError('The first mature-pharma workbook is issuer-specific to PFE.')
    for existing in list(workbook.worksheets):
        workbook.remove(existing)

    sheet = workbook.create_sheet('Mature Pharma Model')
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 75
    sheet.freeze_panes = 'F30'
    sheet.column_dimensions['A'].width = 25
    sheet.column_dimensions['B'].width = 34
    for column in 'CDEFGHIJ':
        sheet.column_dimensions[column].width = 12
    for column in 'KLMNOPQ':
        sheet.column_dimensions[column].width = 11
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.merge_cells('A1:Q1')
    sheet['A1'] = f"{_record(payload.get('company')).get('name') or ticker} ({ticker}) — Product and Patent FCFF DCF"
    sheet['A1'].font = Font(name='Arial', size=16, bold=True, color='FFFFFF')
    sheet['A1'].fill = _TITLE_FILL
    sheet.row_dimensions[1].height = 28
    sheet.merge_cells('A2:Q2')
    sheet['A2'] = 'Filed product revenue and regional patent years are separate from blue analyst assumptions. Modeled global LOE timing uses an editable U.S.-patent starting proxy; forecast formulas remain visible.'
    sheet['A2'].font = _SMALL_FONT
    sheet['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    sheet.row_dimensions[2].height = 30

    _section(sheet, 4, 'Dated market, WACC, operating, and terminal assumptions', 3)
    assumption_rows = [
        (5, 'Current share price', 'currentPrice', 'price'),
        (6, 'Current market capitalization (USD mm)', 'marketCapitalization', 'money'),
        (7, 'Common diluted shares (millions)', 'commonSharesOutstanding', 'shares'),
        (8, 'Cash and equivalents (USD mm)', 'cash', 'money'),
        (9, 'Short-term investments (USD mm)', 'marketableSecurities', 'money'),
        (10, 'Interest-bearing debt (USD mm)', 'debt', 'money'),
        (11, 'Noncontrolling interest (USD mm)', 'nonControllingInterest', 'money'),
        (12, 'Preferred equity (USD mm)', 'preferredEquity', 'money'),
        (13, 'Risk-free rate', 'riskFreeRate', 'percent'),
        (14, 'Equity risk premium', 'equityRiskPremium', 'percent'),
        (15, 'Levered beta', 'beta', 'multiple'),
        (16, 'Current cost of debt', 'costOfDebt', 'percent'),
        (17, 'Marginal tax rate', 'taxRate', 'percent'),
        (22, 'Other/alliance revenue growth', 'otherRevenueGrowth', 'percent'),
        (23, 'EBIT margin', 'ebitMargin', 'percent'),
        (24, 'D&A / revenue', 'depreciationPctRevenue', 'percent'),
        (25, 'Cash CapEx / revenue', 'capexPctRevenue', 'percent'),
        (26, 'Working-capital investment / revenue', 'workingCapitalInvestmentPctRevenue', 'percent'),
        (27, 'Terminal FCFF growth', 'terminalGrowthRate', 'percent'),
    ]
    for row, label, key, kind in assumption_rows:
        missing_other_revenue_growth = (
            payload.get('buildStatus') == 'input_required'
            and key == 'otherRevenueGrowth'
            and assumptions.get(key) is None
        )
        value = None if missing_other_revenue_growth else _finite(assumptions.get(key), key)
        if value is not None and kind in {'money', 'shares'}:
            value /= 1_000_000
        number_format = _PRICE_FORMAT if kind == 'price' else _MONEY_FORMAT if kind in {'money', 'shares'} else _PERCENT_FORMAT if kind == 'percent' else _YEAR_FORMAT
        _write(sheet, f'A{row}', label)
        _write(sheet, f'B{row}', value, input_cell=not missing_other_revenue_growth, formula=missing_other_revenue_growth, number_format=number_format)
        # Full source/method descriptions remain in Data Review; keep the input block compact.
    for row, label, formula, number_format in (
        (18, 'Debt weight', '=IFERROR(B10/(B6+B10),0)', _PERCENT_FORMAT),
        (19, 'Equity weight', '=1-B18', _PERCENT_FORMAT),
        (20, 'Cost of equity — CAPM', '=B13+B15*B14', _PERCENT_FORMAT),
        (21, 'Weighted average cost of capital', '=B19*B20+B18*B16*(1-B17)', _PERCENT_FORMAT),
    ):
        _write(sheet, f'A{row}', label)
        _write(sheet, f'B{row}', formula, formula=True, number_format=number_format)

    product_start = 31
    header_row = product_start
    _section(sheet, header_row - 1, 'Filed product sales and editable loss-of-exclusivity schedule', 17)
    headers = ['Product / revenue group', 'Primary indication', *[year for year in years],
        *[years[-1] + index for index in range(1, 6)],
        'U.S. basic patent', 'Major Europe basic patent', 'Japan basic patent', 'Modeled global LOE year',
        'Pre-LOE growth', 'First-year LOE erosion', 'Post-LOE annual erosion']
    for index, label in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=index, value=label)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        if 3 <= index <= 10:
            cell.number_format = '"FY"0"A"' if index <= 5 else '"FY"0"E"'
    sheet.row_dimensions[header_row].height = 33
    product_map = {str(item.get('product_name')): _record(item) for item in _record(rows[-1].get('pharma')).get('products', [])}
    assumptions_by_name = {str(product.get('productName')): product for product in product_assumptions}
    product_names = list(product_map)
    product_names.sort(key=lambda name: -float(_record(product_map[name].get('revenue')).get('value') or 0))
    source_patents = [_record(item) for item in _record(rows[-1].get('pharma')).get('patents', [])]
    product_rows: dict[str, int] = {}
    for offset, product_name in enumerate(product_names):
        row = product_start + 1 + offset
        product_rows[product_name] = row
        product = product_map[product_name]
        assumption = _record(assumptions_by_name.get(product_name))
        _write(sheet, f'A{row}', product_name, wrap=True)
        _write(sheet, f'B{row}', product.get('indication') or '', wrap=True)
        for index, year_row in enumerate(rows):
            year_products = {_record(item).get('product_name'): _record(item) for item in _record(year_row.get('pharma')).get('products', [])}
            line = _record(year_products.get(product_name)).get('revenue')
            sources = line.get('sources') if isinstance(line.get('sources'), list) else []
            has_filed_revenue = (
                line.get('source') in {'sec_native', 'derived'}
                and isinstance(line.get('value'), (int, float))
                and not isinstance(line.get('value'), bool)
                and math.isfinite(float(line.get('value')))
                and bool(sources)
                and all(_record(source).get('accession') and _record(source).get('filed') for source in sources)
            )
            if allow_missing_product_revenue and not has_filed_revenue:
                _write(sheet, f'{get_column_letter(3 + index)}{row}', None, number_format=_MONEY_FORMAT)
                continue
            amount = _line_value(line, f'{product_name} product revenue', int(year_row['year']))
            _write(sheet, f'{get_column_letter(3 + index)}{row}', _amount_mm(amount), number_format=_MONEY_FORMAT)
        for column, region in (('K', 'us'), ('L', 'major_europe'), ('M', 'japan')):
            values = [
                _finite(_record(patent.get('year')).get('value'), f'{product_name} {region} patent year')
                for patent in source_patents
                if patent_names_for_model(str(patent.get('product_name') or '')) & patent_names_for_model(product_name)
                and patent.get('region') == region
                and patent.get('metric') == 'basic_patent_expiration_year'
                and _record(patent.get('year')).get('value') is not None
            ]
            _write(sheet, f'{column}{row}', min(values) if values else None, number_format=_YEAR_FORMAT)
        modeled_loe = assumption.get('modeledGlobalLoeYear')
        _write(sheet, f'N{row}', modeled_loe, input_cell=True, number_format=_YEAR_FORMAT)
        if allow_missing_product_revenue and assumption.get('preLoeGrowthRate') is None:
            _write(sheet, f'O{row}', None, number_format=_PERCENT_FORMAT)
        else:
            _write(sheet, f'O{row}', _finite(assumption.get('preLoeGrowthRate'), f'{product_name} pre-LOE growth'), input_cell=True, number_format=_PERCENT_FORMAT)
        _write(sheet, f'P{row}', _finite(assumption.get('firstYearErosionRate'), f'{product_name} first-year erosion'), input_cell=True, number_format=_PERCENT_FORMAT)
        _write(sheet, f'Q{row}', _finite(assumption.get('postLoeAnnualErosionRate'), f'{product_name} ongoing erosion'), input_cell=True, number_format=_PERCENT_FORMAT)
        for period, forecast_column in enumerate('FGHIJ', start=1):
            previous = get_column_letter(4 + period)
            formula = (
                f'=IF($N{row}="",{previous}{row}*(1+$O{row}),'
                f'IF({forecast_column}${header_row}<$N{row},{previous}{row}*(1+$O{row}),'
                f'IF({forecast_column}${header_row}=$N{row},{previous}{row}*(1+$O{row})*(1-$P{row}),'
                f'{previous}{row}*(1-$Q{row}))))'
            )
            _write(sheet, f'{forecast_column}{row}', formula, formula=True, number_format=_MONEY_FORMAT)

    product_end = product_start + len(product_names)
    total_row = product_end + 2
    other_row = product_end + 1
    _write(sheet, f'A{other_row}', 'Other products / alliance / royalty revenue', wrap=True)
    _write(sheet, f'B{other_row}', 'Derived residual of filed consolidated revenue less individually disclosed product rows', wrap=True)
    for index, year_row in enumerate(rows):
        column = get_column_letter(3 + index)
        pharma = _record(year_row.get('pharma'))
        filed_total = _line_value(pharma.get('reported_total_revenue'), 'Note 17 reported total revenue', int(year_row['year']))
        _write(sheet, f'{column}{product_end+3}', _amount_mm(filed_total), number_format=_MONEY_FORMAT)
        _write(sheet, f'{column}{other_row}', f'={column}{product_end+3}-SUM({column}{product_start+1}:{column}{product_end})', formula=True, number_format=_MONEY_FORMAT)
        if payload.get('buildStatus') == 'input_required':
            sheet[f'{column}{other_row}'] = f'=IF(\'Input Required\'!$B$3<>"READY","",{column}{product_end+3}-SUM({column}{product_start+1}:{column}{product_end}))'
    for period, forecast_column in enumerate('FGHIJ', start=1):
        previous = get_column_letter(4 + period)
        _write(sheet, f'{forecast_column}{other_row}', f'={previous}{other_row}*(1+$B$22)', formula=True, number_format=_MONEY_FORMAT)
    _write(sheet, f'A{total_row}', 'Total revenue', formula=False)
    _write(sheet, f'B{total_row}', 'Sum of product schedules and other/alliance residual')
    _write(sheet, f'A{product_end+3}', 'Filed total revenue', wrap=True)
    _write(sheet, f'B{product_end+3}', 'SEC 10-K Note 17 total revenue')
    for column in 'CDEFGHIJ':
        product_sum = f'SUM({column}{product_start+1}:{column}{product_end})'
        _write(sheet, f'{column}{total_row}', f'={product_sum}+{column}{other_row}', formula=True, number_format=_MONEY_FORMAT)
        if column not in 'CDE':
            _write(sheet, f'{column}{product_end+3}', None)
    _write(sheet, f'A{product_end+4}', 'Product-table revenue reconciliation check')
    for column in 'CDE':
        _write(sheet, f'{column}{product_end+4}', f'={column}{total_row}-{column}{product_end+3}', formula=True, number_format=_MONEY_FORMAT)
        if payload.get('buildStatus') == 'input_required':
            status_ref = "'Input Required'!$B$3"
            sheet[f'{column}{total_row}'] = f'=IF({status_ref}<>"READY", "", SUM({column}{product_start+1}:{column}{product_end})+{column}{other_row})'
            sheet[f'{column}{product_end+4}'] = f'=IF({status_ref}<>"READY", "", {column}{total_row}-{column}{product_end+3})'
    for column in 'FGHIJ':
        _write(sheet, f'{column}{product_end+4}', None)

    cashflow_start = product_end + 6
    _section(sheet, cashflow_start, 'Unlevered operating cash flow and valuation', 10)
    line_rows = {
        'ebit': cashflow_start + 1,
        'cash_taxes': cashflow_start + 2,
        'nopat': cashflow_start + 3,
        'depreciation': cashflow_start + 4,
        'capex': cashflow_start + 5,
        'nwc': cashflow_start + 6,
        'fcff': cashflow_start + 7,
        'discount': cashflow_start + 8,
        'pv_fcff': cashflow_start + 9,
    }
    labels = {
        'ebit': 'EBIT = revenue × editable EBIT margin', 'cash_taxes': 'Cash taxes = EBIT × editable tax rate',
        'nopat': 'NOPAT', 'depreciation': 'D&A', 'capex': 'Cash CapEx',
        'nwc': 'Working-capital investment', 'fcff': 'Unlevered free cash flow',
        'discount': 'Discount factor', 'pv_fcff': 'Present value of FCFF',
    }
    for key, row in line_rows.items():
        _write(sheet, f'A{row}', labels[key])
    _write(sheet, f'A{total_row}', 'Total revenue')
    for period, column in enumerate('CDEFGHIJ'):
        if period < 3:
            actual = rows[period]
            actual_ebit = _amount_mm(_line_value(actual.get('ebit'), 'EBIT', int(actual['year'])))
            actual_da = _amount_mm(_line_value(actual.get('depreciation'), 'depreciation', int(actual['year'])))
            actual_capex = _amount_mm(abs(_line_value(actual.get('capex'), 'CapEx', int(actual['year']))))
            actual_nwc = _amount_mm(_line_value(actual.get('nwcChange'), 'working-capital change', int(actual['year'])))
            _write(sheet, f'{column}{line_rows["ebit"]}', actual_ebit, number_format=_MONEY_FORMAT)
            _write(sheet, f'{column}{line_rows["cash_taxes"]}', f'={column}{line_rows["ebit"]}*$B$17', formula=True, number_format=_MONEY_FORMAT)
            _write(sheet, f'{column}{line_rows["nopat"]}', f'={column}{line_rows["ebit"]}-{column}{line_rows["cash_taxes"]}', formula=True, number_format=_MONEY_FORMAT)
            _write(sheet, f'{column}{line_rows["depreciation"]}', actual_da, number_format=_MONEY_FORMAT)
            _write(sheet, f'{column}{line_rows["capex"]}', actual_capex, number_format=_MONEY_FORMAT)
            _write(sheet, f'{column}{line_rows["nwc"]}', actual_nwc, number_format=_MONEY_FORMAT)
            _write(sheet, f'{column}{line_rows["fcff"]}', f'={column}{line_rows["nopat"]}+{column}{line_rows["depreciation"]}-{column}{line_rows["capex"]}-{column}{line_rows["nwc"]}', formula=True, number_format=_MONEY_FORMAT)
            _write(sheet, f'{column}{line_rows["discount"]}', None)
            _write(sheet, f'{column}{line_rows["pv_fcff"]}', None)
            continue
        forecast_index = period - 2
        revenue_cell = f'{column}{total_row}'
        _write(sheet, f'{column}{line_rows["ebit"]}', f'={revenue_cell}*$B$23', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{column}{line_rows["cash_taxes"]}', f'={column}{line_rows["ebit"]}*$B$17', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{column}{line_rows["nopat"]}', f'={column}{line_rows["ebit"]}-{column}{line_rows["cash_taxes"]}', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{column}{line_rows["depreciation"]}', f'={revenue_cell}*$B$24', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{column}{line_rows["capex"]}', f'={revenue_cell}*$B$25', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{column}{line_rows["nwc"]}', f'={revenue_cell}*$B$26', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{column}{line_rows["fcff"]}', f'={column}{line_rows["nopat"]}+{column}{line_rows["depreciation"]}-{column}{line_rows["capex"]}-{column}{line_rows["nwc"]}', formula=True, number_format=_MONEY_FORMAT)
        _write(sheet, f'{column}{line_rows["discount"]}', f'=1/(1+$B$21)^{forecast_index}', formula=True, number_format=_PERCENT_FORMAT)
        _write(sheet, f'{column}{line_rows["pv_fcff"]}', f'={column}{line_rows["fcff"]}*{column}{line_rows["discount"]}', formula=True, number_format=_MONEY_FORMAT)

    valuation_start = cashflow_start + 12
    _section(sheet, valuation_start, 'Enterprise-to-common-equity valuation', 8)
    valuation_rows = [
        (valuation_start + 1, 'PV of five-year FCFF', f'=SUM(F{line_rows["pv_fcff"]}:J{line_rows["pv_fcff"]})', _MONEY_FORMAT),
        (valuation_start + 2, 'Terminal value (Gordon growth)', f'=J{line_rows["fcff"]}*(1+$B$27)/($B$21-$B$27)', _MONEY_FORMAT),
        (valuation_start + 3, 'PV of terminal value', f'=B{valuation_start+2}*J{line_rows["discount"]}', _MONEY_FORMAT),
        (valuation_start + 4, 'Enterprise value', f'=B{valuation_start+1}+B{valuation_start+3}', _MONEY_FORMAT),
        (valuation_start + 5, 'Common equity value', f'=B{valuation_start+4}+B8+B9-B10-B11-B12', _MONEY_FORMAT),
        (valuation_start + 6, 'Current share price', '=B5', _PRICE_FORMAT),
        (valuation_start + 7, 'Implied value per diluted share', f'=B{valuation_start+5}/B7', _PRICE_FORMAT),
        (valuation_start + 8, 'Upside / (downside)', f'=B{valuation_start+7}/B5-1', _PERCENT_FORMAT),
    ]
    for row, label, formula, number_format in valuation_rows:
        _write(sheet, f'A{row}', label)
        _write(sheet, f'B{row}', formula, formula=True, number_format=number_format)

    sensitivity_start = valuation_start + 11
    _section(sheet, sensitivity_start, 'Implied value per share sensitivity — WACC vs. terminal growth', 7)
    for index, column in enumerate('CDEFG', start=-2):
        _write(sheet, f'{column}{sensitivity_start+1}', f'=$B$27+({index})*0.005', formula=True, number_format=_PERCENT_FORMAT)
    for row_index, delta in enumerate((-0.01, -0.005, 0, 0.005, 0.01), start=2):
        row = sensitivity_start + row_index
        _write(sheet, f'B{row}', f'=$B$21+({delta})', formula=True, number_format=_PERCENT_FORMAT)
        for column in 'CDEFG':
            pv_terms = '+'.join(f'${forecastCol}${line_rows["fcff"]}/(1+$B{row})^{period}' for period, forecastCol in enumerate('FGHIJ', start=1))
            terminal = f'$J${line_rows["fcff"]}*(1+{column}${sensitivity_start+1})/($B{row}-{column}${sensitivity_start+1})/(1+$B{row})^5'
            bridge = '+$B$8+$B$9-$B$10-$B$11-$B$12'
            _write(sheet, f'{column}{row}', f'=({pv_terms}+{terminal}{bridge})/$B$7', formula=True, number_format=_PRICE_FORMAT)

    _map_data_review(workbook, model)
    sheet.freeze_panes = f'F{header_row+1}'
    sheet.print_area = f'A1:Q{sensitivity_start+6}'
    sheet.print_title_rows = f'1:{header_row}'


def apply_incomplete_mature_pharma_model(
    workbook: Workbook,
    payload: dict[str, Any],
    input_cells: dict[str, dict[str, str]],
) -> None:
    """Link missing product sales to their dynamic product/year rows and gate valuation outputs."""
    sheet = workbook['Mature Pharma Model']
    model = _record(payload.get('maturePharmaModel'))
    history = _record(model.get('history'))
    years = history.get('years') if isinstance(history.get('years'), list) else []
    if len(years) != 3:
        raise ValueError('Incomplete mature-pharma workbook requires three annual product periods.')
    requirements = payload.get('requiredInputs')
    requirements = requirements if isinstance(requirements, list) else []
    status_ref = "'Input Required'!$B$3"
    missing_product_rows: list[tuple[str, int]] = []
    for requirement in requirements:
        if not isinstance(requirement, dict) or not str(requirement.get('key') or '').startswith('product_revenue:'):
            continue
        label = str(requirement.get('label') or '')
        product_name = label.split(' — ', 1)[-1].strip()
        fiscal_year = requirement.get('fiscalYear')
        if not product_name or not isinstance(fiscal_year, int) or fiscal_year not in years:
            raise ValueError('Product revenue input must identify a product and fiscal year.')
        input_identity = f"{requirement['key']}:{fiscal_year}"
        destination = input_cells.get(input_identity)
        if not destination or destination.get('sheet') != 'Input Required':
            raise ValueError(f"Mature-pharma workbook has no editable input for {product_name} FY{fiscal_year}.")
        product_row = next((row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == product_name), None)
        if product_row is None:
            raise ValueError(f"Mature-pharma workbook product row is missing for {product_name}.")
        year_column = get_column_letter(3 + years.index(fiscal_year))
        value_ref = f"'Input Required'!{destination['cell']}"
        cell = sheet[f'{year_column}{product_row}']
        cell.value = f'=IF(AND({status_ref}="READY",ISNUMBER({value_ref}),{value_ref}>=0),{value_ref}/1000000,"")'
        cell.font = _FORMULA_FONT
        cell.fill = _FORMULA_FILL
        cell.number_format = _MONEY_FORMAT
        missing_product_rows.append((product_name, product_row))

    if not missing_product_rows:
        raise ValueError('Incomplete mature-pharma workbook requires a product_revenue input.')

    other_row = next((row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == 'Other products / alliance / royalty revenue'), None)
    total_row = next((row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == 'Total revenue'), None)
    reconciliation_row = next((row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == 'Product-table revenue reconciliation check'), None)
    if other_row is None or total_row is None or reconciliation_row is None:
        raise ValueError('Mature-pharma workbook is missing its product-revenue reconciliation rows.')
    product_start = 31
    product_end = total_row - 2
    for column in 'CDE':
        sheet[f'{column}{other_row}'] = f'=IF({status_ref}<>"READY","",{column}{other_row+2}-SUM({column}{product_start+1}:{column}{product_end}))'
        sheet[f'{column}{total_row}'] = f'=IF({status_ref}<>"READY","",SUM({column}{product_start+1}:{column}{product_end})+{column}{other_row})'
        sheet[f'{column}{reconciliation_row}'] = f'=IF({status_ref}<>"READY","",{column}{total_row}-{column}{other_row+2})'

    product_assumptions = {
        str(item.get('productName')): _record(item)
        for item in _record(model.get('assumptions')).get('products', [])
        if isinstance(item, dict)
    }
    for product_name, product_row in missing_product_rows:
        assumption = product_assumptions.get(product_name, {})
        if assumption.get('preLoeGrowthRate') is not None:
            continue
        c, d, e = (f'{column}{product_row}' for column in 'CDE')
        cagr = f'IF({c}>0,({e}/{c})^(1/2)-1,0)'
        latest_growth = f'IF({d}>0,{e}/{d}-1,0)'
        prior_growth = f'IF({c}>0,{d}/{c}-1,0)'
        current_growth = f'IF({d}>0,{e}/{d}-1,0)'
        direction_reversal = f'AND(({prior_growth})*({current_growth})<0,MAX(ABS({prior_growth}),ABS({current_growth}))>25%)'
        covid_product = 'covid' in f"{product_name} {sheet[f'B{product_row}'].value or ''}".lower()
        use_latest = f'OR({"TRUE" if covid_product else "FALSE"},{direction_reversal},{cagr}<-75%,{cagr}>150%)'
        formula = f'=IF({status_ref}<>"READY","",IF({use_latest},IF(AND({latest_growth}>=-75%,{latest_growth}<=150%),{latest_growth},0),{cagr}))'
        sheet[f'O{product_row}'] = formula
        sheet[f'O{product_row}'].font = _FORMULA_FONT
        sheet[f'O{product_row}'].fill = _FORMULA_FILL
        sheet[f'O{product_row}'].number_format = _PERCENT_FORMAT

    assumptions = _record(model.get('assumptions'))
    if assumptions.get('otherRevenueGrowth') is None:
        c, d, e = (f'{column}{other_row}' for column in 'CDE')
        cagr = f'IF({c}>0,({e}/{c})^(1/2)-1,0)'
        recent = f'IF({d}>0,{e}/{d}-1,0)'
        sheet['B22'] = f'=IF({status_ref}<>"READY","",IF(AND({c}>0,ABS({cagr})<=50%),{cagr},IF(AND({d}>0,ABS({recent})<=50%),{recent},0)))'
        sheet['B22'].font = _FORMULA_FONT
        sheet['B22'].fill = _FORMULA_FILL
        sheet['B22'].number_format = _PERCENT_FORMAT

    for product_row in range(product_start + 1, product_end + 1):
        for column in 'FGHIJ':
            cell = sheet[f'{column}{product_row}']
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.value = f'=IF({status_ref}<>"READY","",{cell.value[1:]})'
    for row in (other_row, total_row):
        for column in 'FGHIJ':
            cell = sheet[f'{column}{row}']
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.value = f'=IF({status_ref}<>"READY","",{cell.value[1:]})'

    cashflow_start = next((row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == 'EBIT = revenue × editable EBIT margin'), None)
    if cashflow_start is None:
        raise ValueError('Mature-pharma workbook is missing its FCFF schedule.')
    for row in range(cashflow_start, cashflow_start + 10):
        for column in 'FGHIJ':
            cell = sheet[f'{column}{row}']
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.value = f'=IF({status_ref}<>"READY","",{cell.value[1:]})'
    valuation_header = next((row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == 'Enterprise-to-common-equity valuation'), None)
    if valuation_header is None:
        raise ValueError('Mature-pharma workbook is missing its valuation schedule.')
    for row in range(valuation_header + 1, valuation_header + 9):
        if row == valuation_header + 6:
            continue
        cell = sheet[f'B{row}']
        if isinstance(cell.value, str) and cell.value.startswith('='):
            cell.value = f'=IF({status_ref}<>"READY","",{cell.value[1:]})'
    sensitivity_header = next((row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == 'Implied value per share sensitivity — WACC vs. terminal growth'), None)
    if sensitivity_header is not None:
        for row in range(sensitivity_header + 2, sensitivity_header + 7):
            for column in 'CDEFG':
                cell = sheet[f'{column}{row}']
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    cell.value = f'=IF({status_ref}<>"READY","",{cell.value[1:]})'


def apply_incomplete_mature_pharma_model(
    workbook: Workbook,
    payload: dict[str, Any],
    input_cells: dict[str, dict[str, str]],
) -> None:
    """Link missing product sales to their dynamic product/year rows and gate valuation outputs."""
    sheet = workbook['Mature Pharma Model']
    model = _record(payload.get('maturePharmaModel'))
    history = _record(model.get('history'))
    years = history.get('years') if isinstance(history.get('years'), list) else []
    if len(years) != 3:
        raise ValueError('Incomplete mature-pharma workbook requires three annual product periods.')
    requirements = payload.get('requiredInputs')
    requirements = requirements if isinstance(requirements, list) else []
    status_ref = "'Input Required'!$B$3"
    missing_product_rows: list[tuple[str, int]] = []
    for requirement in requirements:
        if not isinstance(requirement, dict) or not str(requirement.get('key') or '').startswith('product_revenue:'):
            continue
        label = str(requirement.get('label') or '')
        product_name = label.split(' — ', 1)[-1].strip()
        fiscal_year = requirement.get('fiscalYear')
        if not product_name or not isinstance(fiscal_year, int) or fiscal_year not in years:
            raise ValueError('Product revenue input must identify a product and fiscal year.')
        input_identity = f"{requirement['key']}:{fiscal_year}"
        destination = input_cells.get(input_identity)
        if not destination or destination.get('sheet') != 'Input Required':
            raise ValueError(f"Mature-pharma workbook has no editable input for {product_name} FY{fiscal_year}.")
        product_row = next((row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == product_name), None)
        if product_row is None:
            raise ValueError(f"Mature-pharma workbook product row is missing for {product_name}.")
        year_column = get_column_letter(3 + years.index(fiscal_year))
        value_ref = f"'Input Required'!{destination['cell']}"
        cell = sheet[f'{year_column}{product_row}']
        cell.value = f'=IF(AND({status_ref}="READY",ISNUMBER({value_ref}),{value_ref}>=0),{value_ref}/1000000,"")'
        cell.font = _FORMULA_FONT
        cell.fill = _FORMULA_FILL
        cell.number_format = _MONEY_FORMAT
        missing_product_rows.append((product_name, product_row))

    if not missing_product_rows:
        raise ValueError('Incomplete mature-pharma workbook requires a product_revenue input.')

    other_row = next((row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == 'Other products / alliance / royalty revenue'), None)
    total_row = next((row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == 'Total revenue'), None)
    reconciliation_row = next((row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == 'Product-table revenue reconciliation check'), None)
    if other_row is None or total_row is None or reconciliation_row is None:
        raise ValueError('Mature-pharma workbook is missing its product-revenue reconciliation rows.')
    for column in 'CDE':
        sheet[f'{column}{other_row}'] = f'=IF({status_ref}<>"READY","",{column}{other_row+2}-SUM({column}32:{column}{other_row-1}))'
        sheet[f'{column}{total_row}'] = f'=IF({status_ref}<>"READY","",SUM({column}32:{column}{other_row-1})+{column}{other_row})'
        sheet[f'{column}{reconciliation_row}'] = f'=IF({status_ref}<>"READY","",{column}{total_row}-{column}{other_row+2})'

    cagr_formulas: dict[str, str] = {}
    for product_name, product_row in missing_product_rows:
        c, d, e = (f'{column}{product_row}' for column in 'CDE')
        cagr = f'IF({c}>0,({e}/{c})^(1/2)-1,0)'
        latest_growth = f'IF({d}>0,{e}/{d}-1,0)'
        prior_growth = f'IF({c}>0,{d}/{c}-1,0)'
        current_growth = f'IF({d}>0,{e}/{d}-1,0)'
        direction_reversal = f'AND({prior_growth}*{current_growth}<0,MAX(ABS({prior_growth}),ABS({current_growth}))>25%)'
        covid_product = 'covid' in f'{product_name} {sheet[f"B{product_row}"].value or ""}'.lower()
        choose_latest = f'OR({"TRUE" if covid_product else "FALSE"},{direction_reversal},{cagr}<-75%,{cagr}>150%)'
        formula = f'=IF({status_ref}<>"READY","",IF({choose_latest},IF(AND({latest_growth}>=-75%,{latest_growth}<=150%),{latest_growth},0),{cagr}))'
        sheet[f'O{product_row}'] = formula
        sheet[f'O{product_row}'].font = _FORMULA_FONT
        sheet[f'O{product_row}'].fill = _FORMULA_FILL
        sheet[f'O{product_row}'].number_format = _PERCENT_FORMAT
        cagr_formulas[product_name] = formula

    if payload.get('maturePharmaModel') and _record(model.get('assumptions')).get('otherRevenueGrowth') is None:
        c_other, d_other, e_other = (f'{column}{other_row}' for column in 'CDE')
        cagr = f'IF({c_other}>0,({e_other}/{c_other})^(1/2)-1,0)'
        recent = f'IF({d_other}>0,{e_other}/{d_other}-1,0)'
        sheet['B22'] = f'=IF({status_ref}<>"READY","",IF(AND({c_other}>0,ABS({cagr})<=50%),{cagr},IF(AND({d_other}>0,ABS({recent})<=50%),{recent},0)))'
        sheet['B22'].font = _FORMULA_FONT
        sheet['B22'].fill = _FORMULA_FILL
        sheet['B22'].number_format = _PERCENT_FORMAT

    # Product revenue feeds current-year product cash flows, the residual revenue build, and FCFF.
    for product_row in range(32, total_row):
        for column in 'FGHIJ':
            cell = sheet[f'{column}{product_row}']
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.value = f'=IF({status_ref}<>"READY","",{cell.value[1:]})'
    for row in range(other_row, total_row + 1):
        for column in 'FGHIJ':
            cell = sheet[f'{column}{row}']
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.value = f'=IF({status_ref}<>"READY","",{cell.value[1:]})'

    cashflow_start = next((row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == 'EBIT = revenue × editable EBIT margin'), None)
    if cashflow_start is None:
        raise ValueError('Mature-pharma workbook is missing its FCFF schedule.')
    for row in range(cashflow_start, cashflow_start + 10):
        for column in 'FGHIJ':
            cell = sheet[f'{column}{row}']
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.value = f'=IF({status_ref}<>"READY","",{cell.value[1:]})'
    valuation_header = next((row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == 'Enterprise-to-common-equity valuation'), None)
    if valuation_header is None:
        raise ValueError('Mature-pharma workbook is missing its valuation schedule.')
    for row in range(valuation_header + 1, valuation_header + 9):
        if row == valuation_header + 6:  # Current share price remains visible.
            continue
        cell = sheet[f'B{row}']
        if isinstance(cell.value, str) and cell.value.startswith('='):
            cell.value = f'=IF({status_ref}<>"READY","",{cell.value[1:]})'
    sensitivity_header = next((row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == 'Implied value per share sensitivity — WACC vs. terminal growth'), None)
    if sensitivity_header is not None:
        for row in range(sensitivity_header + 2, sensitivity_header + 7):
            for column in 'CDEFG':
                cell = sheet[f'{column}{row}']
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    cell.value = f'=IF({status_ref}<>"READY","",{cell.value[1:]})'
