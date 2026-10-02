from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook
from openpyxl.workbook import Workbook

from ..mappers import (
    WACC_LOOP_MODE_ITERATIVE,
    apply_payload_to_workbook,
    resolve_wacc_loop_mode,
)
from ..template import load_template_artifact


def _clear_template_markers(workbook: Workbook) -> None:
    """Remove legacy template marker text from the left margin."""
    for worksheet in workbook.worksheets:
        for cell in worksheet["A"]:
            if cell.value == "x":
                cell.value = None


def export_dcf_excel(payload: dict) -> bytes:
    """Build an Excel workbook from the validated template and DCF payload."""
    incomplete_operating_dcf = (
        payload.get("buildStatus", "ready") == "input_required"
        and payload.get("valuationModel") == "unlevered_dcf"
        and isinstance(payload.get("assumptions"), dict)
    )
    if payload.get("buildStatus", "ready") == "input_required" and not incomplete_operating_dcf:
        workbook = Workbook()
        loop_mode = "current_equity"
    else:
        template = load_template_artifact()
        workbook = load_workbook(BytesIO(template.workbook_bytes), data_only=False)
        loop_mode = resolve_wacc_loop_mode(payload)
    apply_payload_to_workbook(workbook, payload)
    _clear_template_markers(workbook)

    # Excel recalculates the model when the workbook opens.
    calculation = workbook.calculation
    calculation.fullCalcOnLoad = True
    calculation.forceFullCalc = True
    calculation.iterate = loop_mode == WACC_LOOP_MODE_ITERATIVE
    if loop_mode == WACC_LOOP_MODE_ITERATIVE:
        calculation.iterateCount = 100
        calculation.iterateDelta = 0.001

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
