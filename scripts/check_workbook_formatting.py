#!/usr/bin/env python3
"""Workbook formatting guard: professional display rules for DCF exports.

Fails when any checked workbook violates:
1. A numeric or formula cell kept on General whose cached value would show
   more than 2 decimals (|value| >= 0.005 and != round(value, 2)). Tiny
   magnitudes keep full precision deliberately.
2. Any cell comment (Excel renders comment indicators on inputs).
3. Any literal "" string value (must be a true blank). [Dropped: "" renders
   blank in Excel and LibreOffice alike; recalculated files normalize to None.]
4. WACC Bull/Bear Beta not on 0.00 and Bull/Bear Cost of Equity not on a
   percent format, when those labeled rows exist.
5. Bare =MIN/=MAX/=PERCENTILE/=MEDIAN( summary formulas left on General
   (distribution summaries must render as #,##0).

Usage: check_workbook_formatting.py <workbook> [<workbook> ...]
"""
from __future__ import annotations

import sys

import openpyxl


def cached_number(cell_formula: object, cached: object) -> bool:
    return isinstance(cached, float) and not isinstance(cached, bool)


def check(path: str) -> list[str]:
    violations: list[str] = []
    book = openpyxl.load_workbook(path, data_only=False)
    cache = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        for ws in book.worksheets:
            try:
                cached_sheet = cache[ws.title]
            except KeyError:
                continue
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    is_numeric = isinstance(value, (int, float)) and not isinstance(value, bool)
                    is_formula = isinstance(value, str) and value.startswith("=")
                    try:
                        cached_value = cached_sheet[cell.coordinate].value
                    except Exception:
                        cached_value = None
                    if (is_numeric or is_formula) and cell.number_format == "General":
                        if (
                            cached_number(value, cached_value)
                            and cached_value != 0
                            and abs(cached_value) >= 0.005
                            and cached_value != round(cached_value, 2)
                        ):
                            violations.append(
                                f"{path} :: {ws.title}!{cell.coordinate} shows raw decimals "
                                f"({cached_value!r}) on General format"
                            )
            for comment in ws._comments:
                violations.append(f"{path} :: {ws.title} has a cell comment (renders an indicator)")
            if "WACC" == ws.title:
                labels = {}
                for row in ws.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.strip() in (
                            "Bull Beta", "Bear Beta", "Bull Cost of Equity", "Bear Cost of Equity",
                        ):
                            labels[cell.value.strip()] = cell.row
                for label, row in labels.items():
                    target = ws.cell(row=row, column=4)
                    if "Beta" in label and target.number_format != "0.00":
                        violations.append(f"{path} :: WACC!D{row} ({label}) is on {target.number_format!r}, want 0.00")
                    if "Cost of Equity" in label and "%" not in str(target.number_format):
                        violations.append(f"{path} :: WACC!D{row} ({label}) is on {target.number_format!r}, want a percent format")
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value.startswith(("=MIN(", "=MAX(", "=PERCENTILE(", "=MEDIAN(")):
                        if cell.number_format == "General":
                            violations.append(
                                f"{path} :: {ws.title}!{cell.coordinate} statistic is on General, want #,##0"
                            )
    finally:
        book.close()
        cache.close()
    return violations


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print("usage: check_workbook_formatting.py <workbook> [<workbook> ...]")
        return 2
    failures: list[str] = []
    for path in argv[1:]:
        failures.extend(check(path))
    for failure in failures:
        print(f"FORMAT-VIOLATION {failure}")
    print(f"{len(failures)} formatting violations in {len(argv) - 1} workbook(s)")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
